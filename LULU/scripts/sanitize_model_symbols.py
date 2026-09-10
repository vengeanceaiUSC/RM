#!/usr/bin/env python3
"""Replace em dashes, arrows, bullets, and other non-ASCII symbols in text cells.

Formulas and numeric hardcodes are never modified.

Run:  cd LULU && python3 scripts/sanitize_model_symbols.py
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "unbeiesgbar2model.xlsx"

# Order matters: longer / specific sequences first.
REPLACEMENTS: list[tuple[str, str]] = [
    ("---", "-"),
    ("--", "-"),
    ("\u2014", "-"),  # em dash —
    ("\u2013", "-"),  # en dash –
    ("\u2212", "-"),  # minus sign −
    ("\u2192", " to "),  # →
    ("\u2190", "<-"),  # ←
    ("\u00b7", " / "),  # ·
    ("\u2022", " / "),  # •
    ("\u2026", "..."),  # …
    ("\u00d7", "x"),  # ×
    ("\u00f7", "/"),  # ÷
    ("\u00b1", "+/-"),  # ±
    ("\u2019", "'"),  # '
    ("\u201c", '"'),  # "
    ("\u201d", '"'),  # "
    ("\u03b2", "beta"),  # β
    ("\u0394", "Delta "),  # Δ
    ("\u2080", "0"),
    ("\u2081", "1"),
    ("\u2082", "2"),
    ("\u2083", "3"),
    ("\u2084", "4"),
    ("\u2085", "5"),
    ("\u2086", "6"),
    ("\u2087", "7"),
    ("\u2088", "8"),
    ("\u2089", "9"),
    ("\u208a", "+"),
    ("\u208b", "-"),
    ("\u208c", "="),
    ("\u208d", "("),
    ("\u208e", ")"),
    ("\u2090", "a"),
    ("\u2091", "e"),
    ("\u2092", "o"),
    ("\u2093", "x"),
    ("\u2094", "schwa"),
    ("\u2095", "h"),
    ("\u2096", "k"),
    ("\u2097", "l"),
    ("\u2098", "m"),
    ("\u2099", "n"),
    ("\u209a", "p"),
    ("\u209b", "s"),
    ("\u209c", "t"),
]

WEIRD_CHAR = re.compile(r"[^\x00-\x7F]")


def sanitize_text(text: str) -> str:
    if not text or text.startswith("="):
        return text
    # Protect URLs from slash collapsing (https:// must not become "https: / ").
    url_slots: list[str] = []

    def _stash_url(match: re.Match[str]) -> str:
        url_slots.append(match.group(0))
        return f"__URL_{len(url_slots) - 1}__"

    out = re.sub(r"https?://[^\s\)\]\"']+", _stash_url, text)
    for old, new in REPLACEMENTS:
        out = out.replace(old, new)
    # Collapse repeated separators/spaces left by substitutions.
    out = re.sub(r"\s*/\s*/\s*", " / ", out)
    out = re.sub(r"\s{2,}", " ", out)
    out = re.sub(r"\s-\s-\s", " - ", out)
    out = re.sub(r"-{2,}", "-", out)
    out = re.sub(r"\bbetau\b", "beta-u", out, flags=re.I)
    out = re.sub(r"\bbetaL\b", "beta L", out, flags=re.I)
    out = re.sub(r"\bbeta u\b", "beta-u", out, flags=re.I)
    out = re.sub(r"^\s*-\s*", "", out)
    out = re.sub(r"\s+\.\s*$", ".", out)
    for i, url in enumerate(url_slots):
        out = out.replace(f"__URL_{i}__", url)
    return out.strip()


def sanitize_workbook(path: Path = TARGET) -> Path:
    tmp = path.with_suffix(".sanitizing.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)
    changed = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if not isinstance(val, str) or val.startswith("="):
                    continue
                new_val = sanitize_text(val)
                if new_val != val:
                    cell.value = new_val if new_val else None
                    changed += 1
    wb.save(tmp)
    tmp.replace(path)
    print(f"Sanitized {path.name} ({changed} text cells updated)")
    return path


def verify(path: Path = TARGET) -> None:
    wb = openpyxl.load_workbook(path, data_only=False)
    weird: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if isinstance(val, str) and not val.startswith("="):
                    if WEIRD_CHAR.search(val):
                        weird.append(f"{ws.title}!{cell.coordinate}: {val[:60]!r}")
    if weird:
        raise AssertionError(
            "Non-ASCII text remains:\n" + "\n".join(weird[:20])
        )
    print("Verify OK: all text cells are ASCII-only")


if __name__ == "__main__":
    sanitize_workbook()
    verify()
