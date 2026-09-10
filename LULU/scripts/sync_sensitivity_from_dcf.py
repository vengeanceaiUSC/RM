#!/usr/bin/env python3
"""Sync slide 20 WACC × g sensitivity table directly from DCF sheet cells.

Reads DCF row 98 (terminal-g headers) and rows 99–103 (WACC × price grid)
after LibreOffice recalc. Rebuilds only the sensitivity table using the
original GIS stmt_table layout (left grid + base-case callout on the right).

Run:  cd LULU && python3 scripts/sync_sensitivity_from_dcf.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

SCRIPTS = Path(__file__).resolve().parent
ROOT = SCRIPTS.parent
DECK = ROOT / "LULU_Investment_Pitch_Deck.pptx"
MODEL = ROOT / "model18_wsp_formulas.xlsx"

sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(ROOT.parent / "GIS"))

from gis_pitch import CARD, INK, PitchDeck, add_para, textbox  # noqa: E402
from recalc_model18 import read_dcf_outputs, recalc_workbook  # noqa: E402


def _fmt_g(v: float) -> str:
    pct = float(v) * 100
    if abs(pct - round(pct, 1)) < 1e-9:
        return f"{pct:.1f}%"
    return f"{pct:.2f}%"


def read_dcf_sensitivity(dcf) -> tuple[list[str], list[list[str]]]:
    """Pull headers + body straight from DCF grid (row 98 / rows 99–103)."""
    headers = ["WACC vs g"] + [_fmt_g(dcf.cell(98, c).value) for c in range(6, 11)]
    rows: list[list[str]] = []
    for r in range(99, 104):
        w = dcf.cell(r, 1).value
        if not isinstance(w, (int, float)):
            continue
        prices = [dcf.cell(r, c).value for c in range(6, 11)]
        if not all(isinstance(p, (int, float)) for p in prices):
            continue
        rows.append([f"{w * 100:.1f}%"] + [f"${round(p)}" for p in prices])
    return headers, rows


def _delete_shape(shape) -> None:
    shape.element.getparent().remove(shape.element)


def _is_sensitivity_table(shape) -> bool:
    if not shape.has_table:
        return False
    h0 = shape.table.rows[0].cells[0].text.strip().replace("\\", " vs ")
    return h0 in ("WACC vs g", "WACC  vs g")


def _center_sensitivity_grid(table) -> None:
    """Center g-rate headers and price cells for a balanced grid."""
    for ri in range(len(table.rows)):
        for ci, cell in enumerate(table.rows[ri].cells):
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def _ensure_base_case_callout(slide, px: str, wacc: str = "9.0", g: str = "2.25") -> None:
    """Restore the base-case annotation box to the right of the sensitivity grid."""
    note = f"WACC {wacc}% × g {g}% gives ${px}. Grid brackets ±100bps WACC and 1.5-3.0% g."
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if shape.text_frame.text.strip().startswith("Base-case cell"):
            shape.left = int(Inches(8.1))
            shape.top = int(Inches(5.80))
            shape.width = int(Inches(4.75))
            shape.height = int(Inches(1.08))
            tf = shape.text_frame
            tf.clear()
            add_para(tf, "Base-case cell", 9, CARD, bold=True, first=True, space_after=2)
            add_para(tf, note, 9, INK, space_after=0)
            return

    tb, tf = textbox(slide, Inches(8.1), Inches(5.80), Inches(4.75), Inches(1.08))
    add_para(tf, "Base-case cell", 9, CARD, bold=True, first=True, space_after=2)
    add_para(tf, note, 9, INK, space_after=0)


def sync(path: Path = DECK, model: Path = MODEL) -> dict:
    wb = recalc_workbook(model)
    data = read_dcf_outputs(wb)
    headers, rows = read_dcf_sensitivity(wb["DCF"])
    px = f"{data['scenarios']['base']['implied_px']:.2f}"

    prs = Presentation(str(path))
    slide = prs.slides[19]
    for shape in list(slide.shapes):
        if _is_sensitivity_table(shape):
            _delete_shape(shape)

    PitchDeck().stmt_table(
        slide,
        rows,
        headers,
        col0w=1.0,
        top=5.80,
        height=1.08,
        left=0.5,
        width=7.4,
        font_size=9,
        header_font_size=9,
    )
    for sh in slide.shapes:
        if _is_sensitivity_table(sh):
            _center_sensitivity_grid(sh.table)
            break
    _ensure_base_case_callout(slide, px)
    prs.save(str(path))

    return {"headers": headers, "rows": rows, "deck": str(path), "base_px": px}


if __name__ == "__main__":
    out = sync()
    print(f"Synced slide 20 sensitivity → {out['deck']}")
    print("headers:", out["headers"])
    for row in out["rows"]:
        print(" ", row)
