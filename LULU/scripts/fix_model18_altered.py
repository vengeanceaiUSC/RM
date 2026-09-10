#!/usr/bin/env python3
"""Repair model18altered.xlsx after polish row/column deletes broke formulas.

Restores correct Hamada beta chain (β used ≈ 0.84), rebinding WACC / Revenue
Drivers / DCF / Scenarios refs, then bakes WACC col D to hardcoded numbers so
link-sourced inputs display without recalc.

Run:  cd LULU/scripts && python3 fix_model18_altered.py
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "model18altered.xlsx"

# Hardcoded link-sourced inputs (Yahoo / FRED / 10-K / NASDAQ) — never formulas.
WACC_PLUGS: dict[str, float] = {
    "D3": 0.048,
    "D4": 0.06,
    "D5": 0.3,
    "D8": 100.0,
    "D9": 111380.0,
    "D11": 1798441.0,
    "D12": 0.0,
    "D16": 0.86,  # Yahoo levered βL (5Y monthly)
    "D17": 2140000.0,  # Yahoo total debt (mrq)
    "D18": 11140000.0,  # Yahoo market cap (mrq)
    "D20": 0.4469,  # Yahoo book D/E — reference only
    "D23": 0.95,  # Damodaran sector βu benchmark (unused)
    "D29": 0.05,  # Pre-tax lease-equivalent cost of debt
}


def _eval_wacc_col(ws, col: int = 4) -> dict[int, float]:
    vals: dict[int, float] = {}

    def get(row: int) -> float:
        if row in vals:
            return vals[row]
        addr = f"D{row}"
        if addr in WACC_PLUGS:
            vals[row] = WACC_PLUGS[addr]
            return vals[row]
        cell = ws.cell(row, col)
        raw = cell.value
        if isinstance(raw, (int, float)):
            vals[row] = float(raw)
            return vals[row]
        if isinstance(raw, str) and raw.startswith("="):
            expr = raw[1:]
            expr = re.sub(
                r"D(\d+)",
                lambda m: str(get(int(m.group(1)))),
                expr,
            )
            result = float(eval(expr))  # noqa: S307 — trusted model formulas
            vals[row] = result
            return result
        raise ValueError(f"WACC row {row}: cannot evaluate {raw!r}")

    for r in range(1, ws.max_row + 1):
        addr = f"D{r}"
        if addr in WACC_PLUGS or (
            isinstance(ws.cell(r, col).value, str)
            and str(ws.cell(r, col).value).startswith("=")
        ):
            get(r)
    return vals


def _fix_wacc(wacc) -> None:
    for addr, val in WACC_PLUGS.items():
        wacc[addr].value = val

    wacc["D10"].value = "=D8*D9"
    wacc["D13"].value = "=D11+D12"
    wacc["D19"].value = "=D17/D18"
    wacc["D21"].value = "=D16/(1+(1-D5)*D19)"
    wacc["D22"].value = "=D13/D10"
    wacc["D24"].value = "=D21*(1+(1-D5)*D22)"
    wacc["D26"].value = "=D3+D24*D4"
    wacc["D30"].value = "=D29*(1-D5)"
    wacc["D33"].value = "=D10/(D10+D13)"
    wacc["D34"].value = "=D13/(D10+D13)"
    wacc["D35"].value = "=D33*D26+D34*D30"

    computed = _eval_wacc_col(wacc)
    for row, num in computed.items():
        wacc.cell(row, 4).value = num


def _fix_revenue_drivers(rd) -> None:
    forecast_cols = ("C", "D", "E", "F", "G")
    for col in forecast_cols:
        rd[f"{col}38"].value = f"={col}35*1000000*{col}36*{col}37/1000"

    rd["C42"].value = "=B42*(1+-0.02)"
    rd["D42"].value = "=C42*(1+0.0)"
    rd["E42"].value = "=D42*(1+0.02)"
    rd["F42"].value = "=E42*(1+0.02)"
    rd["G42"].value = "=F42*(1+0.02)"

    for col in forecast_cols:
        rd[f"{col}43"].value = f"=B43*({col}31+{col}38+{col}42)/(B31+B38+B42)"
        rd[f"{col}44"].value = f"=B44*({col}31+{col}38+{col}42)/(B31+B38+B42)"
        rd[f"{col}45"].value = f"=B45*({col}31+{col}38+{col}42)/(B31+B38+B42)"
        rd[f"{col}52"].value = f"={col}31+{col}38+{col}42"


def _fix_dcf(dcf) -> None:
    for col in ("F", "G", "H", "I", "J"):
        dcf[f"{col}36"].value = f"=1/(1+Scenarios!$F$9)^{col}35"
        dcf[f"{col}37"].value = f"={col}34*{col}36"

    dcf["E41"].value = "=SUM(F37:J37)"
    dcf["E44"].value = "=J34*(1+Scenarios!$F$10)/(Scenarios!$F$9-Scenarios!$F$10)"
    dcf["E46"].value = "=E44/(1+Scenarios!$F$9)^J35"

    dcf["K34"].value = (
        '=IF(MAX(ABS(F34-Scenarios!$F$125),ABS(G34-Scenarios!$F$126),'
        'ABS(H34-Scenarios!$F$127),ABS(I34-Scenarios!$F$128),'
        'ABS(J34-Scenarios!$F$129))<0.5,"OK","CHECK")'
    )


def _fix_scenarios(scn) -> None:
    scn["F9"].value = "=WACC!D35"


def _find_direct_self_refs(wb) -> list[str]:
    issues = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not isinstance(v, str) or not v.startswith("="):
                    continue
                sheet = ws.title
                col = openpyxl.utils.get_column_letter(cell.column)
                key = (sheet, col, cell.row)
                for m in re.finditer(r"(?:'([^']+)'!)?(\$?)([A-Z]{1,3})(\$?)(\d+)", v):
                    sh = m.group(1) or sheet
                    if (sh, m.group(3), int(m.group(5))) == key:
                        issues.append(f"{sheet}!{col}{cell.row}={v}")
    return issues


def fix_workbook(path: Path = TARGET) -> Path:
    tmp = path.with_suffix(".fixing.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)

    _fix_wacc(wb["WACC"])
    _fix_revenue_drivers(wb["Revenue Drivers"])
    _fix_dcf(wb["DCF"])
    _fix_scenarios(wb["Scenarios"])

    wb.save(tmp)
    tmp.replace(path)

    wb2 = openpyxl.load_workbook(path, data_only=False)
    issues = _find_direct_self_refs(wb2)
    if issues:
        raise AssertionError("Circular refs remain:\n" + "\n".join(issues))

    wacc = wb2["WACC"]
    beta_used = wacc["D24"].value
    wacc_val = wacc["D35"].value
    if not isinstance(beta_used, (int, float)) or not (0.83 <= beta_used <= 0.86):
        raise AssertionError(f"Beta used out of range: {beta_used!r}")
    if not isinstance(wacc_val, (int, float)) or not (0.085 <= wacc_val <= 0.095):
        raise AssertionError(f"WACC out of range: {wacc_val!r}")

    print(f"Fixed {path}")
    print(f"  WACC D24 (beta used) = {beta_used:.4f}")
    print(f"  WACC D35 (WACC)      = {wacc_val:.4f}")
    print(f"  Circular refs        = 0")
    return path


if __name__ == "__main__":
    fix_workbook()
