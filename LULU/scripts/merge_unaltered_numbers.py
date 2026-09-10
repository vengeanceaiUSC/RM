#!/usr/bin/env python3
"""Rebuild model18_final.xlsx from model18unaltered.xlsx (numbers source of truth).

Copies all values/formulas from unaltered, then overlays short Notes text from
rewrite_model_notes.py and applies safe presentation cleanup (clear Ctrl+F cols,
scrub AI artifacts) without deleting columns or rows that would shift formulas.

Run:  cd LULU/scripts && python3 merge_unaltered_numbers.py
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string

from rewrite_model_notes import (
    HEADER_CELLS,
    NOTES_MAP,
    SHEETS,
    _lines_ok,
    verify as verify_notes,
)

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "model18unaltered.xlsx"
OUTPUT = ROOT / "model18_final.xlsx"

CTRLF_COL_SHEETS = ("WACC", "Scenarios", "NOPAT Bridge", "Comps")


def _clear_ctrlf_col(ws, col: str = "D") -> int:
    """Clear text in Ctrl+F column; leave formulas untouched."""
    idx = column_index_from_string(col)
    cleared = 0
    for r in range(1, ws.max_row + 1):
        cell = ws.cell(r, idx)
        if isinstance(cell.value, str) and not cell.value.startswith("="):
            if cell.value.strip():
                cell.value = None
                cleared += 1
    return cleared


def _clear_dcf_doc_cols(dcf) -> None:
    """Clear DCF justification / Ctrl+F text columns; keep formula cells."""
    for col in ("B", "D", "M"):
        idx = column_index_from_string(col)
        for r in range(1, dcf.max_row + 1):
            cell = dcf.cell(r, idx)
            if isinstance(cell.value, str) and not cell.value.startswith("="):
                cell.value = None
    if isinstance(dcf["A4"].value, str):
        dcf["A4"].value = "Assumptions Guide (PDF) — see Cover tab link."


def _scrub_text(val: str) -> str | None:
    v = val
    v = re.sub(r"Firecrawl[- ]?ingested[^.]*\.?", "", v, flags=re.I)
    v = re.sub(r"Firecrawl[:\s]*", "", v, flags=re.I)
    v = re.sub(r"agent workflow", "operating adjustments", v, flags=re.I)
    v = re.sub(r"Ctrl\+F[^.\n]*", "", v, flags=re.I)
    v = re.sub(r"\n{2,}", "\n", v).strip()
    return v or None


def _global_scrub(wb, scrub_cols: dict[str, set[int]] | None = None) -> None:
    """Scrub AI/Ctrl+F artifacts from presentation text columns only."""
    default_cols = {2, 3}  # source / narrative cols — never col A (labels)
    for ws in wb.worksheets:
        cols = (scrub_cols or {}).get(ws.title, default_cols)
        for row in ws.iter_rows():
            for cell in row:
                if cell.column not in cols:
                    continue
                if isinstance(cell.value, str) and not cell.value.startswith("="):
                    cleaned = _scrub_text(cell.value)
                    if cleaned != cell.value:
                        cell.value = cleaned


def _polish_cover(cov) -> None:
    cov["B12"].value = "Red font = analyst assumptions (cols B/C on driver tabs)"
    b17 = str(cov["B17"].value or "")
    cov["B17"].value = re.sub(r"\s*Ctrl+F:.*", "", b17, flags=re.I).strip()
    for r in (20, 21, 22, 23):
        v = cov.cell(r, 2).value
        if isinstance(v, str):
            cov.cell(r, 2).value = v.split("\nCtrl+F")[0].split("Ctrl+F")[0].strip()


def _apply_notes(wb) -> int:
    """Overlay short Notes from NOTES_MAP onto col B."""
    rewritten = 0
    for sheet_name in SHEETS:
        ws = wb[sheet_name]
        sheet_map = NOTES_MAP.get(sheet_name, {})

        for row in range(1, ws.max_row + 1):
            label = ws.cell(row, 1).value
            cell = ws.cell(row, 2)
            old_val = cell.value

            key = (sheet_name, row)
            if key in HEADER_CELLS:
                new_val = HEADER_CELLS[key]
                if old_val != new_val:
                    cell.value = new_val
                    rewritten += 1
                continue

            if label is None or not isinstance(old_val, str) or not old_val.strip():
                continue

            label_str = str(label).strip()
            new_val = sheet_map.get(label_str)
            if new_val is None:
                if re.match(r"^Justification\b", str(old_val).strip(), re.I):
                    cell.value = "Notes"
                    rewritten += 1
                elif not _lines_ok(old_val):
                    # Clear orphan long justification text (e.g. Comps Alo build rows).
                    cell.value = None
                    rewritten += 1
                continue

            if old_val != new_val:
                cell.value = new_val
                rewritten += 1

    return rewritten


def merge_workbook(source: Path = SOURCE, output: Path = OUTPUT) -> Path:
    if not source.exists():
        raise FileNotFoundError(source)
    shutil.copy2(source, output)
    wb = openpyxl.load_workbook(output)

    notes_count = _apply_notes(wb)

    for sheet_name in CTRLF_COL_SHEETS:
        _clear_ctrlf_col(wb[sheet_name], "D")

    _clear_ctrlf_col(wb["Revenue Drivers"], "L")

    _clear_dcf_doc_cols(wb["DCF"])
    _polish_cover(wb["Cover"])

    wb["Revenue Drivers"]["A2"].value = (
        "FY2025A = historical 10-K anchors (blue). "
        "FY26–30 = red operational assumptions."
    )
    npb_a4 = str(wb["NOPAT Bridge"]["A4"].value or "")
    wb["NOPAT Bridge"]["A4"].value = "Phases 1–4 clean and forecast operating profit."

    # Scrub source columns (col C) on narrative tabs; Notes col B already rewritten above.
    _global_scrub(wb, scrub_cols={name: {3} for name in SHEETS if name != "DCF"})

    wb.save(output)
    print(f"Merged {source.name} → {output.name} ({notes_count} Notes cells rewritten)")
    return output


def _cell_values_equal(a, b) -> bool:
    if a == b:
        return True
    if isinstance(a, float) and isinstance(b, float):
        return abs(a - b) < 1e-9
    return False


def verify_numbers(source: Path = SOURCE, output: Path = OUTPUT) -> None:
    """Assert all non-notes cells match unaltered (formulas + plugged values)."""
    src_wb = openpyxl.load_workbook(source, data_only=False)
    out_wb = openpyxl.load_workbook(output, data_only=False)
    diffs: list[str] = []

    notes_cols = {2}  # col B on all sheets may differ (Notes text)
    ctrl_f_cols = {4}  # col D cleared on narrative tabs

    for sheet_name in src_wb.sheetnames:
        src_ws = src_wb[sheet_name]
        out_ws = out_wb[sheet_name]
        max_row = max(src_ws.max_row, out_ws.max_row)
        max_col = max(src_ws.max_column, out_ws.max_column)

        for r in range(1, max_row + 1):
            for c in range(1, max_col + 1):
                if sheet_name in SHEETS and c in notes_cols:
                    continue
                if sheet_name in CTRLF_COL_SHEETS and c in ctrl_f_cols:
                    continue
                if sheet_name == "Revenue Drivers" and c == 12:
                    continue
                if sheet_name == "DCF" and c in (2, 3, 4, 13):
                    continue
                if sheet_name == "Cover" and c == 2 and r in (12, 17, 20, 21, 22, 23):
                    continue
                if sheet_name == "Revenue Drivers" and r == 2 and c == 1:
                    continue
                if sheet_name == "NOPAT Bridge" and r == 4 and c == 1:
                    continue
                if sheet_name == "DCF" and r == 4 and c == 1:
                    continue

                sv = src_ws.cell(r, c).value if r <= src_ws.max_row and c <= src_ws.max_column else None
                ov = out_ws.cell(r, c).value if r <= out_ws.max_row and c <= out_ws.max_column else None
                if not _cell_values_equal(sv, ov):
                    diffs.append(
                        f"{sheet_name}!{openpyxl.utils.get_column_letter(c)}{r}: "
                        f"src={sv!r} out={ov!r}"
                    )

    if diffs:
        raise AssertionError(
            f"Number reconciliation failed ({len(diffs)} diffs):\n" + "\n".join(diffs[:20])
        )
    print(f"Reconciliation: 0 numeric/formula diffs vs {source.name}")


def spot_checks(source: Path = SOURCE, output: Path = OUTPUT) -> list[tuple[str, object, object]]:
    """Return key assumption checks (label, unaltered, output)."""
    src = openpyxl.load_workbook(source, data_only=False)
    out = openpyxl.load_workbook(output, data_only=False)
    checks: list[tuple[str, object, object]] = []

    scn_s, scn_o = src["Scenarios"], out["Scenarios"]
    for r in range(4, 11):
        checks.append((f"Scenarios base F{r}", scn_s.cell(r, 6).value, scn_o.cell(r, 6).value))

    checks.append(("DCF E5", src["DCF"]["E5"].value, out["DCF"]["E5"].value))
    checks.append(("WACC E41", src["WACC"]["E41"].value, out["WACC"]["E41"].value))
    checks.append(("Buyback F21", scn_s["F21"].value, scn_o["F21"].value))

    rd_s, rd_o = src["Revenue Drivers"], out["Revenue Drivers"]
    for r in range(1, rd_s.max_row + 1):
        if "Variance (bottom-up" in str(rd_s.cell(r, 1).value or ""):
            checks.append(("RD variance B", rd_s.cell(r, 2).value, rd_o.cell(r, 2).value))
            break

    dcf_s = src["DCF"]
    for r in range(1, dcf_s.max_row + 1):
        if str(dcf_s.cell(r, 1).value or "").strip() == "Implied value per share":
            checks.append(("DCF implied price formula", dcf_s.cell(r, 5).value, out["DCF"].cell(r, 5).value))
            break

    return checks


if __name__ == "__main__":
    out = merge_workbook()
    verify_notes(out)
    verify_numbers()
    checks = spot_checks()
    print("\nSpot checks (unaltered vs output):")
    all_ok = True
    for label, src_val, out_val in checks:
        ok = _cell_values_equal(src_val, out_val)
        mark = "OK" if ok else "FAIL"
        if not ok:
            all_ok = False
        print(f"  [{mark}] {label}: {src_val!r}")
    if not all_ok:
        raise SystemExit(1)
    print(f"\nSaved {out}")
