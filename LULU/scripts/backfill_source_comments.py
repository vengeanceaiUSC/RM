#!/usr/bin/env python3
"""Backfill Analyst comments on value cells from model18unaltered (12).xlsx.

Use when Notes/Source columns are already deleted. Reads original doc cols B-C.

Run:  cd LULU && python3 scripts/backfill_source_comments.py
"""
from __future__ import annotations

import re
import shutil
import sys
from difflib import get_close_matches
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPTS))
from restore_source_columns import _clean_source_label, _truncate_to_words  # noqa: E402

ROOT = SCRIPTS.parent
ORIGINAL = ROOT / "model18unaltered (12).xlsx"
TARGET = ROOT / "unbeiesgbar_final.xlsx"

SHEETS = ("WACC", "Scenarios", "NOPAT Bridge", "Comps", "DCF")
# unaltered value col after label (E on main tabs before our inserts; final uses B)
ORIG_SRC = 3
ORIG_VAL = 5
FINAL_VAL = 2


def _norm(s):
    if not s:
        return ""
    t = str(s).lower().strip().replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", " ", t)


def _label_map(ws):
    m = {}
    for r in range(1, ws.max_row + 1):
        k = _norm(ws.cell(r, 1).value)
        if k and k not in m:
            m[k] = r
    return m


def _match(label, m):
    k = _norm(label)
    if k in m:
        return m[k]
    c = get_close_matches(k, m.keys(), n=1, cutoff=0.82)
    return m[c[0]] if c else None


def backfill(path: Path = TARGET, original: Path = ORIGINAL) -> dict[str, int]:
    tmp = path.with_suffix(".backfill.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)
    orig = openpyxl.load_workbook(original, data_only=False)
    stats = {"added": 0, "updated": 0}

    for sn in SHEETS:
        if sn not in wb.sheetnames or sn not in orig.sheetnames:
            continue
        ws, ow = wb[sn], orig[sn]
        om = _label_map(ow)
        for r in range(1, ws.max_row + 1):
            label = ws.cell(r, 1).value
            if not label:
                continue
            orow = _match(str(label), om)
            if not orow:
                continue
            oc = ow.cell(orow, ORIG_SRC)
            note = _truncate_to_words(str(ow.cell(orow, 2).value or ""))
            src = _clean_source_label(oc.value)
            url = oc.hyperlink.target if oc.hyperlink else None
            parts = []
            if note:
                parts.append(note)
            if src:
                parts.append(f"Source: {src}")
            if url:
                parts.append(url)
            if not parts:
                continue
            body = "\n".join(parts)[:32000]
            cell = ws.cell(r, FINAL_VAL)
            if cell.value is None and not isinstance(cell.value, (int, float)):
                continue
            if cell.comment:
                stats["updated"] += 1
            else:
                stats["added"] += 1
            cell.comment = Comment(body, "Analyst")

    wb.save(tmp)
    tmp.replace(path)
    return stats


if __name__ == "__main__":
    print(f"Backfilled comments: {backfill()}")
