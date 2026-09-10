#!/usr/bin/env python3
"""Move Notes (B) and Source (C) into cell comments; delete columns B and C.

Human pitch models use red-triangle comments, not dedicated doc columns.
Comments are applied AFTER column delete so they survive the shift.

Run:  cd LULU && python3 scripts/move_sources_to_comments.py
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.utils import column_index_from_string, get_column_letter

SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPTS))
from polish_model18_altered import (  # noqa: E402
    _rewrite_all_formulas,
    _rewrite_sheet_internal_formulas,
)

ROOT = SCRIPTS.parent
TARGET = ROOT / "unbeiesgbar_final.xlsx"

SHEETS_BC = ("WACC", "Scenarios", "NOPAT Bridge", "Comps", "DCF")
VALUE_COL = 4  # col D before delete → col B after
HEADER_SKIP = {2, 3, 16, 41}


def _col_shift_delete_bc(max_col: int) -> dict[str, str]:
    return {
        get_column_letter(c): get_column_letter(c - 2)
        for c in range(4, max_col + 1)
    }


def _build_comment(note, source, url: str | None) -> str | None:
    parts: list[str] = []
    if note and str(note).strip() not in ("Notes", "Source", ""):
        parts.append(str(note).strip())
    if source and str(source).strip() not in ("Notes", "Source", ""):
        parts.append(f"Source: {str(source).strip()}")
    if url and str(url).startswith("http"):
        parts.append(str(url).strip())
    if not parts:
        return None
    return "\n".join(parts)[:32000]


def move(path: Path = TARGET) -> dict[str, int]:
    tmp = path.with_suffix(".comments.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)
    stats = {"comments": 0, "sheets": 0, "skipped": 0}

    for name in SHEETS_BC:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        pending: list[tuple[int, str]] = []

        for r in range(1, ws.max_row + 1):
            if r in HEADER_SKIP:
                continue
            a = ws.cell(r, 1).value
            if isinstance(a, str) and a.strip() in ("Notes", "Source", "Driver"):
                continue
            note = ws.cell(r, 2).value
            src_cell = ws.cell(r, 3)
            url = None
            if src_cell.hyperlink:
                url = src_cell.hyperlink.target or src_cell.hyperlink.location
            body = _build_comment(note, src_cell.value, url)
            if not body:
                stats["skipped"] += 1
                continue
            pending.append((r, body))

        shift = _col_shift_delete_bc(ws.max_column + 2)
        ws.delete_cols(2, 2)
        _rewrite_all_formulas(wb, name, shift)
        _rewrite_sheet_internal_formulas(ws, shift)

        new_val_col = get_column_letter(VALUE_COL - 2)  # B
        for r, body in pending:
            cell = ws[f"{new_val_col}{r}"]
            if cell.value is None and not isinstance(cell.value, (int, float)):
                continue
            cell.comment = Comment(body, "Analyst")
            stats["comments"] += 1

        stats["sheets"] += 1

    wb.save(tmp)
    tmp.replace(path)
    return stats


if __name__ == "__main__":
    print(f"Moved sources to comments: {move()}")
