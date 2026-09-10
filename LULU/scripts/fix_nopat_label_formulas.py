#!/usr/bin/env python3
"""Convert NOPAT Bridge subtotal labels from invalid formulas to plain text.

Excel repair error "Removed Records: Formula from sheet5.xml" is caused by
column A labels stored as '= Adjusted EBIT...' (data_type formula). Strip the
leading '= ' so Excel treats them as strings.

Run:  cd LULU && python3 scripts/fix_nopat_label_formulas.py
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import Cell

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "unbeiesgbar_final.xlsx"

LABEL_FORMULA = re.compile(r"^=\s+(.+)$")


def fix(path: Path = TARGET) -> list[str]:
    tmp = path.with_suffix(".nopat_labels.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)
    ws = wb["NOPAT Bridge"]
    changes: list[str] = []

    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell, Cell):
                continue
            v = cell.value
            if not isinstance(v, str):
                continue
            m = LABEL_FORMULA.match(v)
            if not m:
                continue
            new = m.group(1).strip()
            if new != v:
                cell.value = new
                changes.append(f"{cell.coordinate}: {v!r} → {new!r}")

    wb.save(tmp)
    tmp.replace(path)
    return changes


def verify(path: Path = TARGET) -> None:
    ws = openpyxl.load_workbook(path, data_only=False)["NOPAT Bridge"]
    for coord in ("A21", "A26", "A35", "A36"):
        v = ws[coord].value
        if isinstance(v, str) and v.startswith("="):
            raise AssertionError(f"{coord} still formula-like: {v!r}")
    print("Verify OK: NOPAT subtotal labels are plain text")


if __name__ == "__main__":
    changed = fix()
    print(f"Fixed {len(changed)} NOPAT label cells")
    for c in changed:
        print(f"  {c}")
    verify()
