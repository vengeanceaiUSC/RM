#!/usr/bin/env python3
"""Fix truncated Notes ending in vs. minus. as. .. — complete sentences, 8–15 words.

Does not change hardcoded values or Source hyperlinks on assumption rows.

Run:  cd LULU && python3 scripts/fix_truncated_notes.py
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "unbeiesgbar_final.xlsx"

# Exact truncated text → complete analyst note (8–15 words).
NOTE_REPLACEMENTS: dict[str, str] = {
    "Store openings skew to China (+20 FY26) vs.": (
        "China plus twenty store openings FY26; Americas and RoW open slower."
    ),
    "Closures assume 2–3 per mature region annually as.": (
        "Assume two to three store closures per mature region annually."
    ),
    "Simple roll-forward: where you started, plus openings, minus.": (
        "Beginning stores plus regional openings minus closures each year."
    ),
    "Ending stores times avg sq ft per box..": (
        "Total sq ft equals ending store count times average box size."
    ),
    "FY25 China comp +20% (+19% CCY) per 10-K..": (
        "FY25 China comparable sales plus twenty percent per FY25 10-K."
    ),
    "FY25 RoW comp +9% (+7% CCY) per 10-K..": (
        "FY25 rest of world comp plus nine percent per FY25 10-K."
    ),
    "Not % of consolidated revenue; store EBIT ÷.": (
        "Store EBIT margin applied to store channel revenue only."
    ),
    "Not % of consolidated revenue; e-commerce EBIT ÷.": (
        "E-commerce EBIT margin applied to e-comm channel revenue only."
    ),
    "Not % of consolidated revenue; other-channel EBIT ÷.": (
        "Other-channel EBIT margin applied to other channel revenue only."
    ),
    "No page prints Alo or Color Image EBITDA..": (
        "Alo Yoga excluded from comp set; no reliable EBITDA on pages."
    ),
    "βu = 0.86 ÷ [1 + 0.70 ×.": (
        "Hamada unlever beta using Yahoo market D/E and tax rate."
    ),
    "Yahoo Market Cap = $11.14B on the same.": (
        "Yahoo market cap from same page as beta and debt."
    ),
    "Unlever D/E = Yahoo debt (mrq) ÷ Yahoo.": (
        "Yahoo total debt divided by Yahoo market cap ratio."
    ),
}

NOTE_COLS: dict[str, int] = {
    "WACC": 2,
    "Scenarios": 2,
    "NOPAT Bridge": 2,
    "Comps": 2,
    "DCF": 2,
    "Revenue Drivers": 9,
}

TRUNCATED = re.compile(
    r"( vs\.| minus\.| as\.| not\.| for\.|\.\.$|,\s*not\.$| ÷\.)$",
    re.I,
)


def _word_count(text: str) -> int:
    return len(text.split())


def fix_notes(path: Path = TARGET) -> list[str]:
    tmp = path.with_suffix(".trunc_fix.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)
    changes: list[str] = []

    for sheet, col in NOTE_COLS.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        for r in range(1, ws.max_row + 1):
            cell = ws.cell(r, col)
            old = cell.value
            if not isinstance(old, str) or old.strip() in ("Notes", "Source"):
                continue
            new = NOTE_REPLACEMENTS.get(old.strip())
            if new is None and TRUNCATED.search(old.strip()):
                continue  # flagged but no map entry
            if new and new != old:
                if not (8 <= _word_count(new) <= 15):
                    raise ValueError(f"{sheet}!{cell.coordinate}: {new!r} not 8–15 words")
                cell.value = new
                changes.append(f"{sheet}!{cell.coordinate}: {old!r} → {new!r}")

    wb.save(tmp)
    tmp.replace(path)
    return changes


def verify(path: Path = TARGET) -> None:
    wb = openpyxl.load_workbook(path, data_only=False)
    issues: list[str] = []
    for sheet, col in NOTE_COLS.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        for r in range(1, ws.max_row + 1):
            v = ws.cell(r, col).value
            if isinstance(v, str) and TRUNCATED.search(v.strip()):
                issues.append(f"{sheet}!{ws.cell(r,col).coordinate}: {v!r}")
    if issues:
        raise AssertionError("Truncated notes remain:\n" + "\n".join(issues))
    print("Verify OK: no truncated note endings")


if __name__ == "__main__":
    changed = fix_notes()
    print(f"Fixed {len(changed)} truncated notes")
    for c in changed:
        print(f"  {c}")
    verify()
