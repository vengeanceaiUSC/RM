#!/usr/bin/env python3
"""Remove NOPAT references and GIS template boilerplate from the pitch deck."""
from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "LULU_Investment_Pitch_Deck.pptx"

_REPLACEMENTS = [
    (re.compile(r"\s*/\s*NOPAT Bridge[^/\n]*", re.I), ""),
    (re.compile(r"NOPAT Bridge\s*&\s*", re.I), ""),
    (re.compile(r"NOPAT Bridge\b", re.I), "Scenarios"),
    (re.compile(r"\bNOPAT\b", re.I), "operating profit"),
    (
        re.compile(
            r"Content slides can be added / removed depending on the "
            r"investment opportunity \(PM discretion\)\.\s*",
            re.I,
        ),
        "",
    ),
]


def _scrub(text: str) -> str:
    out = text
    for pattern, repl in _REPLACEMENTS:
        out = pattern.sub(repl, out)
    return re.sub(r"  +", " ", out).strip()


def scrub(path: Path = DECK) -> int:
    prs = Presentation(str(path))
    changes = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                new = _scrub(shape.text_frame.text)
                if new != shape.text_frame.text:
                    shape.text_frame.text = new
                    changes += 1
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        new = _scrub(cell.text)
                        if new != cell.text:
                            cell.text = new
                            changes += 1
    prs.save(str(path))
    return changes


if __name__ == "__main__":
    n = scrub()
    print(f"Scrubbed {n} text region(s) → {DECK}")
