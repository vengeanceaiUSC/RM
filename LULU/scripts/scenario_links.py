"""External links from the 3-statement workbook to DCF Scenarios (base case = column G).

Build order: run build_dcf.py before build_3statement.py so Scenarios exists.
Row numbers match build_dcf.py SC{} keys (s_assum block starting at row 4).

Assumptions pull from the local "Scenarios Base" tab, which links to Scenarios!G
in the DCF workbook (live when both files are open in Excel).
"""
DCF_FILE = "LULU_DCF_Valuation_Model.xlsx"
MIRROR_SHEET = "Scenarios Base"
SCEN_SHEET = "Scenarios"
BASE_COL = "G"

# Row index on Scenarios tab (must stay in sync with build_dcf.py s_assum order)
ROWS = {
    "g1": 4,
    "gterm": 5,
    "m1": 6,
    "tariff": 7,
    "mterm": 8,
    "wacc": 9,
    "g": 10,
    "tax": 11,
    "da_pct": 12,
    "capex_pct": 13,
    "gm_pct": 14,
}

# Mirror sheet row for each key (col B holds the linked Scenarios!G value)
MIRROR_ROW = {key: i + 2 for i, key in enumerate(ROWS)}


def scen_cell(key: str, col: str = BASE_COL) -> str:
    """External cell reference (no leading =) — DCF Scenarios column G."""
    return f"'[{DCF_FILE}]{SCEN_SHEET}'!${col}${ROWS[key]}"


def scen_ref(row: int, col: str = BASE_COL) -> str:
    """Excel external reference to a Scenarios cell (base case column G by default)."""
    return f"='[{DCF_FILE}]{SCEN_SHEET}'!${col}${row}"


def scen_base(key: str) -> str:
    """External link formula to Scenarios!G for one assumption key."""
    return f"={scen_cell(key)}"


def mirror_cell(key: str) -> str:
    """Internal mirror cell reference (no leading =)."""
    return f"'{MIRROR_SHEET}'!$B${MIRROR_ROW[key]}"


def mirror_ref(key: str) -> str:
    """Internal link to the local Scenarios Base mirror (column B)."""
    return f"={mirror_cell(key)}"


def clean_ebit_margin_path(year_index: int) -> str:
    """Interpolated clean EBIT margin — mirrors Scenarios mar_rows formula (t=0..4)."""
    t = year_index
    return f"={mirror_cell('m1')}+({mirror_cell('mterm')}-{mirror_cell('m1')})*{t}/4"


def read_scenario_base(dcf_wb):
    """Read Scenarios column G (base case) from a recalculated DCF workbook."""
    sc = dcf_wb[SCEN_SHEET]
    return {key: sc.cell(ROWS[key], 7).value for key in ROWS}


def sync_mirror_sheet(model_path, dcf_wb):
    """Write live Scenarios!G values into the local mirror for headless recalc."""
    import openpyxl

    vals = read_scenario_base(dcf_wb)
    wb = openpyxl.load_workbook(model_path)
    mir = wb[MIRROR_SHEET]
    for key, row in MIRROR_ROW.items():
        mir.cell(row, 2).value = vals[key]
    wb.save(model_path)
