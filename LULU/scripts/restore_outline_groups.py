"""Restore Excel outline (+/-) controls for collapsible column/row groups."""
from __future__ import annotations

from openpyxl.utils import column_index_from_string, get_column_letter

# Column groups after humanization: hide doc columns only; keep Source (C) visible.
SHEET_COL_GROUPS: dict[str, list[tuple[str, str]]] = {
    "WACC": [("B", "B"), ("D", "D")],
    "Scenarios": [("B", "B"), ("D", "D")],
    "NOPAT Bridge": [("B", "B"), ("D", "D")],
    "Comps": [("B", "B"), ("D", "D")],
    "Revenue Drivers": [("J", "K")],
    "DCF": [("B", "B"), ("D", "D")],
}

# Row groups: (start_row, end_row, summary_below)
SHEET_ROW_GROUPS: dict[str, list[tuple[int, int, bool]]] = {
    "DCF": [(20, 32, False)],
    "NOPAT Bridge": [(15, 40, True)],
}


def enable_outline_symbols(ws) -> None:
    ws.sheet_view.showOutlineSymbols = True


def clear_sheet_outlines(ws) -> None:
    """Remove all row/column outline metadata from a worksheet."""
    for dim in ws.column_dimensions.values():
        dim.outline_level = 0
        dim.hidden = False
    for dim in ws.row_dimensions.values():
        dim.outline_level = 0
        dim.hidden = False


def group_columns(
    ws,
    start_col: str,
    end_col: str,
    *,
    hidden: bool = True,
    outline_level: int = 1,
) -> None:
    start = column_index_from_string(start_col)
    end = column_index_from_string(end_col)
    if end < start:
        return
    for i in range(start, end + 1):
        col = get_column_letter(i)
        dim = ws.column_dimensions[col]
        dim.outline_level = outline_level
        dim.hidden = hidden
    ws.sheet_properties.outlinePr.summaryRight = True
    ws.sheet_properties.outlinePr.applyStyles = True


def group_rows(
    ws,
    start_row: int,
    end_row: int,
    *,
    hidden: bool = True,
    outline_level: int = 1,
    summary_below: bool = True,
) -> None:
    if end_row < start_row:
        return
    ws.row_dimensions.group(start_row, end_row, outline_level=outline_level, hidden=hidden)
    ws.sheet_properties.outlinePr.summaryBelow = summary_below
    ws.sheet_properties.outlinePr.applyStyles = True


def remove_outline_groups(wb) -> int:
    """Strip all Excel row/column outline levels and hide +/- controls."""
    cleared = 0
    for ws in wb.worksheets:
        touched = False
        for col, dim in ws.column_dimensions.items():
            if dim.outline_level or dim.hidden:
                dim.outline_level = 0
                dim.hidden = False
                touched = True
        for row, dim in ws.row_dimensions.items():
            if dim.outline_level or dim.hidden:
                dim.outline_level = 0
                dim.hidden = False
                touched = True
        ws.sheet_view.showOutlineSymbols = False
        if ws.sheet_properties.outlinePr is not None:
            ws.sheet_properties.outlinePr.showOutlineSymbols = False
        if touched:
            cleared += 1
    return cleared


def restore_outline_groups(wb, *, col_hidden: bool = True) -> int:
    """Re-apply outline metadata and show +/- controls on every grouped sheet."""
    restored = 0
    for sheet_name, col_ranges in SHEET_COL_GROUPS.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        clear_sheet_outlines(ws)
        enable_outline_symbols(ws)
        for start, end in col_ranges:
            group_columns(ws, start, end, hidden=col_hidden)
        for start, end, summary_below in SHEET_ROW_GROUPS.get(sheet_name, []):
            group_rows(ws, start, end, hidden=col_hidden, summary_below=summary_below)
        restored += 1

    for ws in wb.worksheets:
        enable_outline_symbols(ws)
    return restored
