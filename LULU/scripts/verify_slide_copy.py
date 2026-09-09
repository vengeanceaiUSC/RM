"""Verify narrative slides 2-13 contain the user's canonical copy verbatim.

Run after build_pitch.py:  python3 verify_slide_copy.py
Exit code 0 = every title/descriptor/bullet/link fragment is present.
Exit code 1 = at least one fragment is missing (prints which).

The expectations are parsed straight out of `LULU/slide_copy_source.txt`, which
is the user's slide text pasted verbatim. Nothing is re-typed here, so the
checked copy cannot drift from what the user actually wrote: to change the
deck's wording, edit that file (and build_pitch.py) rather than this script.

Deviations tolerated while matching:
  * en/em dashes are treated as hyphens and curly quotes as straight quotes
    (the deck deliberately uses the ASCII forms),
  * a trailing sentence period on a bullet,
  * the author notes in STRAY_NOTES, which are asides rather than slide copy,
  * on the slides in SPLIT_LABEL_SLIDES, a "Label: body" line may render as a
    separate label shape and body shape.
"""
import os
import re
import sys
import unicodedata

from pptx import Presentation

HERE = os.path.dirname(__file__)
DECK = os.path.join(HERE, "..", "LULU_Investment_Pitch_Deck.pptx")
SOURCE = os.path.join(HERE, "..", "slide_copy_source.txt")

# "Slide N" heading in the source file -> slide number in the built deck.
# The user's slide 7 has a second part, so every later slide shifts by one.
USER_TO_DECK = {
    "2": 2,
    "3": 3,
    "4": 4,
    "5": 5,
    "6": 6,
    "7": 7,
    "7 Part 2": 8,
    "8": 9,
    "9": 10,
    "10": 11,
    "11": 12,
    "12": 13,
}

# Asides the user typed into the copy that are not slide text.
STRAY_NOTES = ("is that ready",)

# Deck slides that render a "Label: body" source line as two shapes.
SPLIT_LABEL_SLIDES = {13}

_HEADING = re.compile(r"^Slide (\d+(?: Part 2)?)\s*:?\s*(.*)$")
_SEPARATOR = re.compile(r"^[_\-\u2013\u2014]+$")
_CITATION_SPLIT = re.compile(r"(?=\[\d+\]\s)")


def _norm(s: str) -> str:
    s = unicodedata.normalize("NFKC", s)
    for dash in ("\u2013", "\u2014", "\u2212"):
        s = s.replace(dash, "-")
    for quote in ("\u2018", "\u2019"):
        s = s.replace(quote, "'")
    for quote in ("\u201c", "\u201d"):
        s = s.replace(quote, '"')
    return " ".join(s.split())


def _strip_stray(frag: str) -> str:
    for note in STRAY_NOTES:
        if frag.endswith(note):
            frag = frag[: -len(note)]
    return frag.strip().rstrip(".").strip()


def parse_source(path: str = SOURCE) -> dict:
    """Read the canonical copy file into {deck slide: [(kind, fragment), ...]}."""
    expected: dict[int, list[tuple[str, str]]] = {}
    current: list[tuple[str, str]] | None = None
    for line in open(path, encoding="utf-8"):
        line = line.strip()
        if not line or _SEPARATOR.match(line):
            continue
        heading = _HEADING.match(line)
        if heading:
            key, title = heading.group(1), heading.group(2).strip()
            if key not in USER_TO_DECK:
                raise SystemExit(f"unmapped source heading: Slide {key}")
            current = expected.setdefault(USER_TO_DECK[key], [])
            if title:
                current.append(("title", title))
            continue
        if current is None:
            raise SystemExit(f"copy found before the first slide heading: {line}")
        line = re.sub(r"^Links & Sources:?\s*", "", line).strip()
        if not line:
            continue
        if line.startswith("["):
            for part in _CITATION_SPLIT.split(line):
                if part.strip():
                    current.append(("link", part.strip()))
        else:
            current.append(("line", line))
    return expected


def _present(frag: str, slide_text: str, deck_n: int) -> bool:
    if frag in slide_text:
        return True
    if deck_n in SPLIT_LABEL_SLIDES and ": " in frag:
        label, _, body = frag.partition(": ")
        return label in slide_text and body in slide_text
    return False


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        if shape.has_table:
            for row in shape.table.rows:
                parts.extend(cell.text for cell in row.cells)
    return _norm(" \n ".join(parts))


def main() -> int:
    expected = parse_source()
    prs = Presentation(DECK)
    missing = 0
    checked = 0
    for deck_n, frags in sorted(expected.items()):
        slide_text = _slide_text(prs.slides[deck_n - 1])
        for kind, frag in frags:
            checked += 1
            if not _present(_strip_stray(_norm(frag)), slide_text, deck_n):
                missing += 1
                print(f"SLIDE {deck_n} MISSING {kind}: {frag}")
    if missing:
        print(f"\nFAIL: {missing} of {checked} fragment(s) missing from slides 2-13.")
        return 1
    print(f"PASS: slides 2-13 contain all {checked} canonical fragments verbatim.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
