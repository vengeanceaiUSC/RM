#!/usr/bin/env python3
"""Fix DCF terminal-value reconciliation block — E→D column refs after Notes insert.

Rows 90–93 were still pointing at empty column E; outputs live in column D.
Also restores variance formulas and row 89 headers from unaltered logic.

Run:  cd LULU && python3 scripts/fix_dcf_tv_reconciliation.py
"""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "unbeiesgbar_final.xlsx"

FIXES: dict[str, str | None] = {
    "B89": "Gordon Growth",
    "C89": "Exit Multiple",
    "D89": "Variance",
    "B90": "=D45",
    "C90": "=D83",
    "D90": "=B90-C90",
    "B91": "=D46",
    "C91": "=D82",
    "D91": "=B91-C91",
    "D92": "=(B90-C90)/B90",
    "B93": "=D56",
    "C93": "=D86",
    "D93": "=B93-C93",
    "B94": '=IF(ABS(D91)<=1.5,"PASS","REVIEW")',
    "C94": '=TEXT(D91,"0.0")&"x spread vs Gordon-implied exit"',
    "D94": "Selected exit is the Gordon identity, so spread should be 0",
}


def fix(path: Path = TARGET) -> list[str]:
    tmp = path.with_suffix(".tv_fix.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)
    ws = wb["DCF"]
    changes: list[str] = []

    for coord, val in FIXES.items():
        cell = ws[coord]
        if cell.value != val:
            changes.append(f"{coord}: {cell.value!r} → {val!r}")
            cell.value = val

    wb.save(tmp)
    tmp.replace(path)
    return changes


def verify(path: Path = TARGET) -> None:
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["DCF"]
    issues: list[str] = []

    for coord, expected in [
        ("B90", "=D45"),
        ("C90", "=D83"),
        ("D90", "=B90-C90"),
        ("B93", "=D56"),
        ("C93", "=D86"),
    ]:
        if ws[coord].value != expected:
            issues.append(f"{coord} expected {expected!r}, got {ws[coord].value!r}")

    # E45 must stay empty (wrong column)
    if ws["E45"].value is not None:
        issues.append(f"E45 should be empty, got {ws['E45'].value!r}")

    if issues:
        raise AssertionError("Verify failed:\n" + "\n".join(issues))
    print("Verify OK: TV reconciliation refs point to column D")


if __name__ == "__main__":
    changed = fix()
    print(f"Fixed {len(changed)} cells in {TARGET.name}")
    for c in changed:
        print(f"  {c}")
    verify()
