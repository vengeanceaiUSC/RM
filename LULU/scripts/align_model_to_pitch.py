#!/usr/bin/env python3
"""Align unbeiesgbar_final.xlsx with pitch deck (pitch_values.json / slides).

- $750M/yr base buybacks (Scenarios + DCF)
- Pitch repurchase price path ($100 → $108 → $115 → $122 → $130)
- DCF repurchase block pulls NI / shares / EPS from Scenarios pitch bridge (col D)
- Clean human-readable decimals (DSO 6.3, DIO 128.8, no 6.26788364887504)

Run:  cd LULU && python3 scripts/align_model_to_pitch.py
"""
from __future__ import annotations

import math
import shutil
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "unbeiesgbar_final.xlsx"

FCOLS = ["C", "D", "E", "F", "G"]
BASE = "D"  # Scenarios base-case column (was G pre column-delete)

# Pitch deck repurchase avg price path ($/sh) — FY26–FY30
REPURCHASE_PRICES = [100, 108, 115, 122, 130]

# Human-rounded WC / ratio display anchors (Scenarios cols C–E + DCF col B links)
SCN_ROUNDS: dict[int, float] = {
    15: 6.3,
    16: 128.8,
    18: 25.1,
    19: 0.051,
    20: 0.060,
}

NUM = "#,##0;(#,##0)"
PCT = "0.0%"
DAY = "0.0"
MONEY = "$#,##0.00"

# Pitch bridge row map (Scenarios tab)
PITCH = {
    "ni": 143,       # Net income year 1 = row 143
    "shares": 183,   # Diluted shares year 1
    "eps": 188,      # Diluted EPS year 1
}


def _round_if_float(v, places: int = 4):
    if isinstance(v, float) and not isinstance(v, bool):
        if math.isfinite(v) and abs(v - round(v)) > 10 ** (-places - 1):
            return round(v, places)
    return v


def _clean_hardcode_decimals(wb) -> int:
    """Round ugly float literals and link FY25 anchors to formulas where possible."""
    n = 0
    dcf = wb["DCF"]
    nb = wb["NOPAT Bridge"]
    scn = wb["Scenarios"]

    # Scenarios WC assumptions — human rounded
    for row, val in SCN_ROUNDS.items():
        for col in (3, 4, 5):
            cell = scn.cell(row, col)
            if cell.value != val:
                cell.value = val
                n += 1
            cell.number_format = DAY if row in (15, 16, 18) else PCT

    # DCF FY25 WC/ratio inputs — link to Scenarios base (no long literals)
    links: list[tuple[str, str, str]] = [
        ("B7", "=(B6-B8)/B6", PCT),
        ("B12", "=B11/B6", PCT),
        ("B21", f"=Scenarios!${BASE}$15", DAY),
        ("B23", f"=Scenarios!${BASE}$16", DAY),
        ("B25", f"=Scenarios!${BASE}$18", DAY),
        ("B27", f"=Scenarios!${BASE}$19", PCT),
        ("B29", f"=Scenarios!${BASE}$20", PCT),
    ]
    for addr, formula, fmt in links:
        if dcf[addr].value != formula:
            dcf[addr] = formula
            n += 1
        dcf[addr].number_format = fmt
    # guard: accrued % row must not inherit comma format
    dcf["B29"].number_format = PCT

    # NOPAT FY25 channel EBIT — mirror col C formulas (not stale hardcodes)
    for r in (31, 32, 33, 35):
        want = f"=C{r}"
        if nb[f"B{r}"].value != want:
            nb[f"B{r}"] = want
            n += 1
        nb[f"B{r}"].number_format = NUM

    # Tax rate FY25 — round stored value if hardcoded
    b39 = nb["B39"]
    if isinstance(b39.value, float):
        rounded = round(b39.value, 3)
        if b39.value != rounded:
            b39.value = rounded
            n += 1
    b39.number_format = PCT

    # FY25 accretive EPS display
    if isinstance(dcf["B79"].value, float):
        dcf["B79"].value = round(dcf["B79"].value, 2)
        dcf["B79"].number_format = MONEY
        n += 1

    # Sweep remaining long floats in value columns
    for ws in (dcf, nb, scn, wb["WACC"], wb["Comps"]):
        for row in ws.iter_rows():
            for cell in row:
                if cell.column == 1:
                    continue
                if isinstance(cell.value, float) and not isinstance(cell.value, bool):
                    s = f"{cell.value:.12f}".rstrip("0")
                    if "." in s and len(s.split(".")[1]) > 4:
                        label = str(ws.cell(cell.row, 1).value or "").lower()
                        if any(k in label for k in ("dso", "dio", "dpo", "days")):
                            cell.value = round(cell.value, 1)
                            cell.number_format = DAY
                        elif "%" in str(cell.number_format) or any(
                            k in label for k in ("margin", "rate", "growth", "tax", "weight", "mix", "erp")
                        ):
                            cell.value = round(cell.value, 3)
                            cell.number_format = PCT
                        elif abs(cell.value) >= 1000:
                            cell.value = round(cell.value)
                            cell.number_format = NUM
                        else:
                            cell.value = round(cell.value, 2)
                        n += 1

    return n


def _align_buybacks(wb) -> int:
    n = 0
    scn = wb["Scenarios"]
    dcf = wb["DCF"]
    for col in ("C", "D"):
        if scn[f"{col}21"].value != 750000:
            scn[f"{col}21"] = 750000
            n += 1
    if dcf["B66"].value != f"=Scenarios!${BASE}$21":
        dcf["B66"] = f"=Scenarios!${BASE}$21"
        n += 1
    dcf["B66"].number_format = NUM
    return n


def _align_repurchase_prices(wb) -> int:
    """Pitch price path in assumption row + DCF projected share price row."""
    n = 0
    scn = wb["Scenarios"]
    dcf = wb["DCF"]

    # Assumption labels (col A/B) for audit trail
    labels = [
        (198, "Repurchase avg price Y1 ($/sh)", 100),
        (199, "Repurchase avg price Y2 ($/sh)", 108),
        (200, "Repurchase avg price Y3 ($/sh)", 115),
        (201, "Repurchase avg price Y4 ($/sh)", 122),
        (202, "Repurchase avg price Y5 ($/sh)", 130),
    ]
    for row, label, price in labels:
        if scn[f"A{row}"].value != label:
            scn[f"A{row}"] = label
            n += 1
        if scn[f"B{row}"].value != price:
            scn[f"B{row}"] = price
            scn[f"B{row}"].number_format = MONEY
            n += 1

    # DCF spot + forecast price path
    if dcf["B73"].value != 100:
        dcf["B73"] = 100
        n += 1
    dcf["B73"].number_format = MONEY
    for i, col in enumerate(FCOLS):
        ref = f"Scenarios!$B${198 + i}"
        if dcf[f"{col}73"].value != f"={ref}":
            dcf[f"{col}73"] = f"={ref}"
            dcf[f"{col}73"].number_format = MONEY
            n += 1

    # Scenarios diluted share path — $750M/yr at pitch price assumptions
    for col in ("C", "D"):
        bb = f"{col}21"
        r0 = 183
        if scn[f"{col}{r0}"].value != f"=$B$197-{bb}/$B$198":
            scn[f"{col}{r0}"] = f"=$B$197-{bb}/$B$198"
            n += 1
        for i in range(1, 5):
            row = r0 + i
            prev = row - 1
            if scn[f"{col}{row}"].value != f"={col}{prev}-{bb}/$B${198 + i}":
                scn[f"{col}{row}"] = f"={col}{prev}-{bb}/$B${198 + i}"
                n += 1

    # Bull case: % of FCF buybacks at same price path as base
    for i, row in enumerate(range(183, 188)):
        price = f"$B${198 + i}"
        fcf_row = 153 + i
        if i == 0:
            want = f"=$B$197-ROUND(E22*E{fcf_row},0)/{price}"
        else:
            want = f"=E{row - 1}-ROUND(E22*E{fcf_row},0)/{price}"
        if scn[f"E{row}"].value != want:
            scn[f"E{row}"] = want
            n += 1

    return n


def _align_dcf_repurchase_to_pitch_bridge(wb) -> int:
    """Wire DCF repurchase schedule to Scenarios pitch bridge (base col D)."""
    n = 0
    dcf = wb["DCF"]
    scn = wb["Scenarios"]

    # Budget links to Scenarios base assumption
    for col in FCOLS:
        if dcf[f"{col}66"].value != f"=Scenarios!${BASE}$21":
            dcf[f"{col}66"] = f"=Scenarios!${BASE}$21"
            n += 1
        if dcf[f"{col}70"].value != f"=Scenarios!${BASE}$21":
            dcf[f"{col}70"] = f"=Scenarios!${BASE}$21"
            n += 1

    # Share price already set in _align_repurchase_prices
    for col in FCOLS:
        if dcf[f"{col}74"].value != f"={col}70/{col}73":
            dcf[f"{col}74"] = f"={col}70/{col}73"
            n += 1

    # Beginning / ending shares from pitch bridge
    dcf["C75"] = f"=Scenarios!${BASE}$197"
    for i, col in enumerate(FCOLS):
        sh_row = PITCH["shares"] + i
        if dcf[f"{col}76"].value != f"=Scenarios!${BASE}${sh_row}":
            dcf[f"{col}76"] = f"=Scenarios!${BASE}${sh_row}"
            n += 1
        if i > 0:
            prev = FCOLS[i - 1]
            if dcf[f"{col}75"].value != f"={prev}76":
                dcf[f"{col}75"] = f"={prev}76"
                n += 1

    # Net income + EPS from pitch bridge (matches slides)
    for i, col in enumerate(FCOLS):
        ni_row = PITCH["ni"] + i
        eps_row = PITCH["eps"] + i
        if dcf[f"{col}78"].value != f"=Scenarios!${BASE}${ni_row}":
            dcf[f"{col}78"] = f"=Scenarios!${BASE}${ni_row}"
            n += 1
        if dcf[f"{col}79"].value != f"=Scenarios!${BASE}${eps_row}":
            dcf[f"{col}79"] = f"=Scenarios!${BASE}${eps_row}"
            n += 1
        dcf[f"{col}78"].number_format = NUM
        dcf[f"{col}79"].number_format = MONEY

    # Share repurchases on pitch bridge pull fixed budget
    for i in range(5):
        r = 158 + i
        for col in ("C", "D"):
            if scn[f"{col}{r}"].value != f"=-{col}21":
                scn[f"{col}{r}"] = f"=-{col}21"
                n += 1

    dcf["B78"].number_format = NUM
    dcf["B79"].number_format = MONEY
    dcf["B79"].comment = Comment("FY25 reported diluted EPS (10-K).", "Analyst")
    return n


def _fix_wacc_formats(wb) -> int:
    n = 0
    w = wb["WACC"]
    fixes = {
        "B4": PCT,
        "B16": "0.00",
        "B19": "0.00",
        "B20": "0.00",
        "B21": "0.00",
        "B22": "0.00",
        "B24": "0.00",
        "B27": PCT,
        "B30": PCT,
        "B31": PCT,
        "B34": PCT,
        "B35": PCT,
        "B36": PCT,
    }
    for addr, fmt in fixes.items():
        if w[addr].number_format != fmt:
            w[addr].number_format = fmt
            n += 1
    return n


def align(path: Path = TARGET) -> dict:
    tmp = path.with_suffix(".pitch.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)

    stats = {
        "decimals": _clean_hardcode_decimals(wb),
        "buybacks": _align_buybacks(wb),
        "prices": _align_repurchase_prices(wb),
        "repurchase": _align_dcf_repurchase_to_pitch_bridge(wb),
        "wacc_fmt": _fix_wacc_formats(wb),
    }

    wb.save(tmp)
    tmp.replace(path)
    return stats


if __name__ == "__main__":
    print(f"Align model to pitch: {align()}")
