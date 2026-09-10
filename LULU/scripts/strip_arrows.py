#!/usr/bin/env python3
"""Remove Unicode arrow (→) from text cells and Analyst comments in unbeiesgbar_final.xlsx.

Run:  cd LULU && python3 scripts/strip_arrows.py
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "unbeiesgbar_final.xlsx"
ARROW = "\u2192"


def _strip_arrows(text: str) -> str:
    out = text.replace(ARROW, " to ")
    out = re.sub(r"\s{2,}", " ", out)
    return out.strip()


def strip_workbook(path: Path = TARGET) -> tuple[Path, list[tuple[str, str, str]]]:
    tmp = path.with_suffix(".stripping-arrows.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)
    changed: list[tuple[str, str, str]] = []

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if isinstance(val, str) and ARROW in val and not val.startswith("="):
                    new_val = _strip_arrows(val)
                    if new_val != val:
                        changed.append((ws.title, cell.coordinate, val[:80]))
                        cell.value = new_val if new_val else None
                if cell.comment and ARROW in (cell.comment.text or ""):
                    old = cell.comment.text
                    new = _strip_arrows(old)
                    if new != old:
                        changed.append((ws.title, f"{cell.coordinate} (comment)", old[:80]))
                        cell.comment.text = new

    wb.save(tmp)
    tmp.replace(path)
    print(f"Stripped arrows from {path.name}: {len(changed)} updates")
    return path, changed


def verify(path: Path = TARGET) -> None:
    wb = openpyxl.load_workbook(path, data_only=False)
    hits: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and ARROW in cell.value:
                    hits.append(f"{ws.title}!{cell.coordinate}")
                if cell.comment and ARROW in (cell.comment.text or ""):
                    hits.append(f"{ws.title}!{cell.coordinate} (comment)")
    if hits:
        raise AssertionError("Arrow remains:\n" + "\n".join(hits))
    print("Verify OK: no → in cells or comments")


if __name__ == "__main__":
    _, samples = strip_workbook()
    verify()
    if samples:
        print("\nSample changes:")
        for sheet, coord, before in samples[:8]:
            print(f"  [{sheet} {coord}] {before!r}")
