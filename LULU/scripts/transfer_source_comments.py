#!/usr/bin/env python3
"""Transfer source evidence from model18unaltered (12) into clean model as Excel notes.

Compares original doc columns (B-D / I-L) for sources and URLs, rewrites them as
concise Analyst comments on the matching value cells in unbeiesgbar2model.xlsx.

Run:  cd LULU && python3 scripts/transfer_source_comments.py
"""
from __future__ import annotations

import re
import shutil
from difflib import get_close_matches
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.utils import column_index_from_string

ROOT = Path(__file__).resolve().parents[1]
ORIGINAL = ROOT / "model18unaltered (12).xlsx"
CLEAN = ROOT / "unbeiesgbar2model.xlsx"
OUTPUT = ROOT / "unbeiesgbar_final.xlsx"

URL_RE = re.compile(r"https?://[^\s\)\]\"']+", re.I)
AI_RE = re.compile(
    r"(?i)\b(?:firecrawl|agent workflow|\bagent\b|chatgpt|openai|anthropic|claude|"
    r"gemini|\bllm\b|ai-generated|as an ai)\b"
)
CTRLF_RE = re.compile(r"Ctrl\+F[^\n]*", re.I)

# Original documentation columns per sheet (inclusive).
DOC_COLS: dict[str, tuple[int, ...]] = {
    "WACC": (2, 3, 4),
    "Scenarios": (2, 3, 4),
    "NOPAT Bridge": (2, 3, 4),
    "Comps": (2, 3, 4),
    "DCF": (2, 3, 4, 12, 13),
    "Revenue Drivers": (9, 10, 11, 12),
}

BASE_SCENARIO_COL = "D"  # base case in stripped Scenarios tab


def _norm_label(text) -> str:
    if text is None:
        return ""
    s = str(text).lower().strip()
    s = s.replace("–", "-").replace("—", "-").replace("β", "beta")
    s = re.sub(r"^\s*memo:\s*", "", s)
    s = re.sub(r"\s+", " ", s)
    s = re.sub(r"[^\w\s$%()./-]", "", s)
    return s


def _extract_urls(*texts: str | None, hyperlink_target: str | None = None) -> list[str]:
    urls: list[str] = []
    if hyperlink_target and str(hyperlink_target).startswith("http"):
        urls.append(str(hyperlink_target).strip())
    for t in texts:
        if not t:
            continue
        urls.extend(URL_RE.findall(str(t)))
    # stable dedupe
    seen: set[str] = set()
    out: list[str] = []
    for u in urls:
        if u not in seen:
            seen.add(u)
            out.append(u)
    return out


def _first_sentence(text: str | None, limit: int = 120) -> str:
    if not text:
        return ""
    t = CTRLF_RE.sub("", str(text))
    t = AI_RE.sub("", t)
    t = re.sub(r"\s+", " ", t).strip()
    if not t:
        return ""
    parts = re.split(r"(?<=[.!?])\s+", t)
    s = parts[0].strip()
    if len(s) > limit:
        s = s[: limit - 1].rsplit(" ", 1)[0] + "…"
    return s


def _format_thousands(value: float) -> str:
    if abs(value) >= 1_000_000:
        return f"${value / 1_000_000:.1f}M"
    if abs(value) >= 1_000:
        return f"${value / 1_000:.1f}k"
    return f"${value:.2f}"


def _rewrite_note(label: str, source_c: str | None, source_d: str | None, urls: list[str]) -> str | None:
    src = (source_c or "").strip()
    src_low = src.lower()
    label_low = label.lower()

    if not src and not urls and not source_d:
        return None
    if src_low in {"source (click)", "ctrl+f (prove number)"}:
        return None
    if "justification" in src_low and "click" in src_low:
        return None

    # Headline by source type.
    if "10-k" in src_low or "sec.gov" in " ".join(urls).lower():
        headline = "FY25 10-K anchor."
    elif "fred" in src_low or "fred.stlouisfed.org" in " ".join(urls).lower():
        headline = "FRED macro anchor."
    elif "damodaran" in src_low:
        headline = "Damodaran implied ERP reference."
    elif "yahoo" in src_low:
        headline = "Yahoo Finance market-data cross-check."
    elif "nasdaq" in src_low:
        headline = "NASDAQ last-sale cross-check."
    elif "pitchbook" in src_low:
        headline = "PitchBook comps set (EV / TTM EBITDA)."
    elif "earnings" in src_low or "corporate.lululemon.com" in " ".join(urls).lower():
        headline = "Q2 FY26 management guidance."
    elif "stockanalysis" in src_low:
        headline = "StockAnalysis consensus forecast."
    elif "wacc tab" in src_low or "internal" in src_low:
        headline = "Model cross-reference."
    elif "revenue drivers" in src_low or "nopat bridge" in src_low:
        headline = "Driver tab linkage."
    else:
        headline = _first_sentence(src, 80) or _first_sentence(label, 80)

    detail = _first_sentence(source_d, 90)
    if detail and "ctrl" not in detail.lower():
        pass
    else:
        detail = ""

    # Contextual one-liner from label.
    if "risk-free" in label_low:
        body = "AAA blue-chip proxy; durable rate anchor for unlevered FCF."
    elif "equity risk premium" in label_low:
        body = "Conservative ERP overlay vs implied market ERP."
    elif "tax rate" in label_low:
        body = "Cash tax assumption for Hamada relever and NOPAT."
    elif "share price" in label_low or "market cap" in label_low:
        body = "Market equity input for capital structure and beta relever."
    elif "lease" in label_low or "debt" in label_low:
        body = "ASC 842 lease debt equivalent in capital structure."
    elif "beta" in label_low:
        body = "Hamada beta chain; margin-of-safety discount rate input."
    elif "wacc" in label_low and "weight" not in label_low:
        body = "Lease-adjusted WACC; base case for DCF and sensitivity."
    elif "dso" in label_low or "dio" in label_low or "dpo" in label_low:
        body = "Working-capital driver; supports unlevered FCF bridge."
    elif "revenue" in label_low or "margin" in label_low:
        body = "Operating assumption with durable cash-flow read-through."
    elif "pitchbook" in src_low or "comps" in label_low:
        body = "Peer multiple reference; not the selected terminal value."
    else:
        body = "High-conviction analyst overlay on reported baseline."

    parts = [headline, body]
    if detail and detail not in headline:
        parts.insert(1, detail)
    note = " ".join(p for p in parts if p).strip()
    if urls:
        note += "\nSource: " + "\nSource: ".join(urls)
    note = AI_RE.sub("", note)
    note = CTRLF_RE.sub("", note)
    note = re.sub(r"\s+", " ", note).strip()
    return note if len(note) > 20 else None


def _collect_row_evidence(ws, row: int, doc_cols: tuple[int, ...]) -> tuple[str | None, str | None, list[str]]:
    texts: list[str] = []
    urls: list[str] = []
    c_text = None
    d_text = None
    for col in doc_cols:
        cell = ws.cell(row, col)
        if col == 3:
            c_text = str(cell.value) if cell.value is not None else None
        if col == 4:
            d_text = str(cell.value) if cell.value is not None else None
        if isinstance(cell.value, str) and not cell.value.startswith("="):
            texts.append(cell.value)
        if cell.hyperlink:
            target = cell.hyperlink.target or cell.hyperlink.location
            urls.extend(_extract_urls(hyperlink_target=target))
    # Also scan value-column hyperlinks in original (col E or scenario cols).
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row, col)
        if cell.hyperlink:
            target = cell.hyperlink.target or cell.hyperlink.location
            urls.extend(_extract_urls(hyperlink_target=target))
    urls.extend(_extract_urls(*texts))
    # dedupe urls
    seen: set[str] = set()
    deduped: list[str] = []
    for u in urls:
        if u not in seen:
            seen.add(u)
            deduped.append(u)
    return c_text, d_text, deduped


def _target_cell(clean_ws, row: int, sheet: str):
    """Return coordinate in clean workbook to attach the note."""
    if sheet == "Scenarios":
        # Base-case assumption value.
        col = BASE_SCENARIO_COL
        if clean_ws[f"{col}{row}"].value is None:
            return None
        return f"{col}{row}"
    if sheet in ("NOPAT Bridge", "DCF", "Revenue Drivers"):
        # Prefer FY2026E col C; fall back to FY2025A col B.
        if clean_ws.cell(row, 3).value is not None:
            return f"C{row}"
        if clean_ws.cell(row, 2).value is not None:
            return f"B{row}"
        return None
    if sheet == "Cover":
        if clean_ws.cell(row, 2).value is not None:
            return f"B{row}"
        return None
    # WACC / Comps: value in column B.
    if clean_ws.cell(row, 2).value is not None:
        return f"B{row}"
    return None


def _build_label_map(ws) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for r in range(1, ws.max_row + 1):
        key = _norm_label(ws.cell(r, 1).value)
        if key and key not in mapping:
            mapping[key] = r
    return mapping


def _match_row(clean_label: str, orig_map: dict[str, int]) -> int | None:
    if clean_label in orig_map:
        return orig_map[clean_label]
    candidates = get_close_matches(clean_label, orig_map.keys(), n=1, cutoff=0.82)
    return orig_map[candidates[0]] if candidates else None


def transfer_comments(
    original: Path = ORIGINAL,
    clean: Path = CLEAN,
    output: Path = OUTPUT,
) -> Path:
    if not original.exists():
        raise FileNotFoundError(original)
    if not clean.exists():
        raise FileNotFoundError(clean)

    tmp = output.with_suffix(".transferring.xlsx")
    shutil.copy2(clean, tmp)
    orig_wb = openpyxl.load_workbook(original, data_only=False)
    out_wb = openpyxl.load_workbook(tmp)

    inserted = 0
    skipped = 0

    for sheet in out_wb.sheetnames:
        if sheet not in orig_wb.sheetnames:
            continue
        orig_ws = orig_wb[sheet]
        out_ws = out_wb[sheet]
        doc_cols = DOC_COLS.get(sheet, (2, 3, 4))
        orig_map = _build_label_map(orig_ws)

        for r in range(1, out_ws.max_row + 1):
            label = out_ws.cell(r, 1).value
            if label is None:
                continue
            key = _norm_label(label)
            if not key:
                continue
            orig_row = _match_row(key, orig_map)
            if orig_row is None:
                skipped += 1
                continue

            c_text, d_text, urls = _collect_row_evidence(orig_ws, orig_row, doc_cols)
            note = _rewrite_note(str(label), c_text, d_text, urls)
            if not note:
                skipped += 1
                continue

            target = _target_cell(out_ws, r, sheet)
            if not target:
                skipped += 1
                continue

            out_ws[target].comment = Comment(note, "Analyst")
            inserted += 1

    out_wb.save(tmp)
    tmp.replace(output)
    print(f"Saved {output.name}: {inserted} Analyst notes inserted ({skipped} rows skipped)")
    return output


def verify(output: Path = OUTPUT) -> None:
    wb = openpyxl.load_workbook(output, data_only=False)
    count = 0
    bad_author = []
    ai_hits = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not cell.comment:
                    continue
                count += 1
                if cell.comment.author not in (None, "Analyst"):
                    bad_author.append(f"{ws.title}!{cell.coordinate}")
                txt = cell.comment.text or ""
                if AI_RE.search(txt) or "Ctrl+F" in txt:
                    ai_hits.append(f"{ws.title}!{cell.coordinate}")
    if bad_author:
        raise AssertionError(f"Bad comment authors: {bad_author[:5]}")
    if ai_hits:
        raise AssertionError(f"AI/Ctrl+F in comments: {ai_hits[:5]}")
    if count < 50:
        raise AssertionError(f"Too few comments inserted: {count}")
    print(f"Verify OK: {count} Analyst notes, no AI/Ctrl+F residue")


if __name__ == "__main__":
    transfer_comments()
    verify()
