#!/usr/bin/env python3
"""Fix Table of Contents formatting on slide 1 (narrow # column wraps digits).

Run:  cd LULU && python3 scripts/fix_pitch_toc_formatting.py
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "LULU_Investment_Pitch_Deck.pptx"

COL0_W = Inches(0.58)
ROW_AREA_H = Inches(5.20)
FONT_PT = 9


def _fix_toc_table(table) -> int:
    if table.rows[0].cells[0].text.strip() != "#":
        return 0
    total_w = table.columns[0].width + table.columns[1].width
    table.columns[0].width = int(COL0_W)
    table.columns[1].width = int(total_w - COL0_W)
    row_h = int(ROW_AREA_H / len(table.rows))
    for ri, row in enumerate(table.rows):
        row.height = row_h
        c0 = row.cells[0]
        c0.margin_left = Inches(0.05)
        c0.margin_right = Inches(0.05)
        c0.margin_top = Inches(0.02)
        c0.margin_bottom = Inches(0.02)
        c0.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf0 = c0.text_frame
        tf0.word_wrap = False
        for para in tf0.paragraphs:
            para.alignment = PP_ALIGN.CENTER
            for run in para.runs:
                run.font.size = Pt(FONT_PT)
        c1 = row.cells[1]
        c1.vertical_anchor = MSO_ANCHOR.MIDDLE
        c1.margin_left = Inches(0.06)
        c1.margin_top = Inches(0.02)
        c1.margin_bottom = Inches(0.02)
        if ri == 0:
            for para in c1.text_frame.paragraphs:
                for run in para.runs:
                    run.font.bold = True
    return 1


def main() -> None:
    prs = Presentation(str(DECK))
    fixed = 0
    for shape in prs.slides[0].shapes:
        if shape.has_table:
            fixed += _fix_toc_table(shape.table)
    prs.save(str(DECK))
    print(f"Fixed {fixed} TOC table(s) → {DECK}")


if __name__ == "__main__":
    main()
