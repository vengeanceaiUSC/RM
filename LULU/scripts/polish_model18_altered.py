#!/usr/bin/env python3
"""Polish model18altered.xlsx for submission — strip AI artifacts, delete doc columns, tighten copy.

Run:  cd LULU/scripts && python3 polish_model18_altered.py
Reads/writes: ../model18altered.xlsx (model18unaltered.xlsx is never touched).
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "model18altered.xlsx"

SCEN_SHORT = {
    "FY2026E revenue growth": "Q2 FY26 Guidance Midpoint",
    "FY2027–FY2030E revenue growth": "StockAnalysis 3Y Revenue Forecast",
    "Run-rate EBIT margin": "Q2 FY26 Run-Rate OM (ex-tariff)",
    "IEEPA tariff refunds": "FY26 One-Time Tariff Refund",
    "Terminal (FY2030E) EBIT margin": "Partial Recovery vs FY25 10-K OM",
    "WACC": "Lease-Adjusted CAPM (WACC Tab)",
    "Terminal growth": "FRED GDPC1 Anchor (~2.1%)",
    "Tax rate": "FY26 Guidance (~30%)",
    "Cash tax rate": "FY26 Guidance (~30%)",
    "D&A % of revenue": "FY25 CF Statement (D&A / Sales)",
    "Capex % of revenue": "FY26 Guide Fade to 5.5% Blend",
    "Gross margin %": "FY25 Reported GM (~56.6%)",
    "DSO": "Flat vs FY25 (~6.3 days)",
    "DIO": "FY25 Anchor; −1 Day / Yr",
    "DPO": "Flat vs FY25 (~25 days)",
    "Prepaid expenses (%": "FY25 OCA % of Revenue",
    "Accrued liabilities (%": "FY25 Accrued % of Revenue",
    "fixed buyback": "Pitch CFF — Fixed Buyback",
    "% of FCF to buybacks": "Pitch CFF — % FCF to Buybacks",
}

WACC_SHORT = {
    "Risk-free rate": "FRED DGS10 (4.77% → 4.8%)",
    "Equity risk premium": "Damodaran implied ERP + 177bps overlay",
    "Tax rate": "FY26 Guidance (~30%)",
    "Share price": "NASDAQ Last Sale (~$100)",
    "Shares outstanding": "FY25 10-K Share Count",
    "Market value of equity": "Price × Shares",
    "Operating lease liabilities": "ASC 842 Lease Debt Equiv.",
    "Funded debt": "No Term Debt (FY25 10-K)",
    "Observed Beta": "Yahoo βL (5Y Monthly)",
    "Yahoo Total Debt": "Yahoo MRQ Total Debt",
    "Yahoo Market Cap": "Yahoo MRQ Market Cap",
    "D/E for unlever": "Yahoo MRQ D/E (Hamada)",
    "Yahoo book D/E": "Reference Only — Not in Hamada",
    "Unlevered βu": "Hamada Unlever (Yahoo D/E)",
    "D/E for relever": "FY25 Lease Debt / Mkt Cap",
    "Sector βu benchmark": "Damodaran Benchmark (unused)",
    "Beta used": "Relevered β for CAPM",
    "Cost of equity": "CAPM: rf + β × ERP",
    "Pre-tax cost of debt": "Lease-Equivalent Borrowing Cost",
    "After-tax cost of debt": "kd × (1 − T)",
    "Equity weight": "Mkt Cap / Total Capital",
    "Debt weight": "Lease Debt / Total Capital",
    "WACC": "Weighted Avg Cost of Capital",
}

NOPAT_SHORT = {
    "Add back SBC": "SBC Flag (Convention A = 0)",
    "Store-channel EBIT margin": "Store EBIT % of Store Rev",
    "E-commerce EBIT margin": "E-comm EBIT % of E-comm Rev",
    "Other-channels EBIT margin": "Other EBIT % of Other Rev",
    "Terminal marginal tax rate": "Statutory Marginal (30%)",
    "R&D / software amortization": "R&D Capitalization Period",
    "Reported operating income": "10-K EBIT / Scenarios Forecast",
    "Impairment / intangible amortization": "Adds back run-rate intangible amortization",
    "Restructuring costs": "Non-Recurring Restructuring Add-Back",
    "Legal / M&A one-offs": "One-Off Legal / M&A Strip",
    "Stock-based compensation": "SBC Add-Back (if flag = 1)",
    "Implied lease interest": "Lease Interest Reclass (memo)",
    "Capitalize R&D": "R&D Capitalization (immaterial)",
    "Amortization of capitalized": "Capitalized R&D Amortization",
}


def _rewrite_formula_refs(val: str, sheet: str, col_shift: dict[str, str]) -> str:
    """Shift column letters in cross-sheet refs, e.g. Scenarios!G25 -> Scenarios!F25."""
    if not isinstance(val, str) or not val.startswith("="):
        return val

    def repl(m: re.Match) -> str:
        sh_quoted, sh_plain = m.group(1), m.group(2)
        sh = sh_quoted or sh_plain
        if sh != sheet:
            return m.group(0)
        d1, col, d2, row = m.group(3), m.group(4), m.group(5), m.group(6)
        new_col = col_shift.get(col, col)
        prefix = f"'{sh}'" if sh_quoted else sh
        return f"{prefix}!{d1}{new_col}{d2}{row}"

    pat = re.compile(
        r"(?:'([^']+)'|([A-Za-z][A-Za-z0-9_]*))!(\$?)([A-Z]{1,3})(\$?)(\d+)"
    )
    return pat.sub(repl, val)


def _rewrite_all_formulas(wb, sheet: str, col_shift: dict[str, str]) -> None:
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.value = _rewrite_formula_refs(cell.value, sheet, col_shift)


def _delete_col_and_fix_refs(wb, sheet_name: str, col_letter: str) -> None:
    idx = column_index_from_string(col_letter)
    ws = wb[sheet_name]
    col_shift = {}
    for c in range(idx + 1, ws.max_column + 2):
        old = get_column_letter(c)
        new = get_column_letter(c - 1)
        col_shift[old] = new
    ws.delete_cols(idx, 1)
    _rewrite_all_formulas(wb, sheet_name, col_shift)


def _clear_col(ws, col: str) -> None:
    idx = column_index_from_string(col)
    for row in ws.iter_rows():
        row[idx - 1].value = None


def _delete_rows_matching(ws, *patterns: str, col: int = 1) -> None:
    to_del = []
    for r in range(1, ws.max_row + 1):
        txt = str(ws.cell(r, col).value or "")
        if any(p in txt for p in patterns):
            to_del.append(r)
    for r in sorted(to_del, reverse=True):
        ws.delete_rows(r, 1)


def _delete_nopat_phase5_block(ws) -> None:
    start = None
    for r in range(1, ws.max_row + 1):
        a = str(ws.cell(r, 1).value or "")
        if "Phase 5" in a or "Reconciliation workflow" in a:
            start = r
            break
    if not start:
        return
    end = start
    for r in range(start, ws.max_row + 1):
        combined = str(ws.cell(r, 1).value or "") + str(ws.cell(r, 2).value or "")
        if any(
            p in combined
            for p in (
                "Scenarios base-case NOPAT",
                "Bear/bull apply",
                "Without applied",
            )
        ):
            end = r
    for r in range(end, start - 1, -1):
        ws.delete_rows(r, 1)


def _shorten_col_b(ws, mapping: dict[str, str]) -> None:
    for r in range(1, ws.max_row + 1):
        label = str(ws.cell(r, 1).value or "")
        cell = ws.cell(r, 2)
        if not isinstance(cell.value, str):
            continue
        for key, short in mapping.items():
            if key.lower() in label.lower() or key.lower() in cell.value.lower():
                cell.value = short
                break
        else:
            txt = cell.value.split("\n")[0].strip()
            if len(txt) > 60 and "." in txt:
                txt = txt.split(".")[0].strip()
            txt = re.sub(r"Ctrl+F.*", "", txt, flags=re.I).strip()
            if txt:
                cell.value = txt[:80]


def _clean_source_col(ws, col: int = 3) -> None:
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, col).value
        if not isinstance(v, str):
            continue
        v = re.sub(r"\s*Ctrl+F.*", "", v, flags=re.I)
        v = re.sub(r"Same page as beta[^.]*\.?", "", v, flags=re.I)
        v = re.sub(r"reference only, not in Hamada\.?", "", v, flags=re.I)
        v = re.sub(r"\s+", " ", v).strip(" ;,")
        ws.cell(r, col).value = v or None


def _standardize_dcf_source(val: str | None, label: str = "") -> str | None:
    if not isinstance(val, str):
        return None
    v = re.sub(r"Ctrl\+F.*", "", val, flags=re.I | re.DOTALL).strip()
    v = v.split("\n")[0].strip()
    if not v:
        return None
    low = (label + " " + v).lower()
    if "10-k" in low or "sec.gov" in low:
        return "10-K"
    if "q2 fy2026" in low or "earnings release" in low or "lululemon.com" in low:
        return "Q2 FY26 Release"
    if "stockanalysis" in low:
        return "StockAnalysis"
    if "fred" in low or "gdpc1" in low or "dgs10" in low:
        return "FRED"
    if "damodaran" in low:
        return "Damodaran"
    if "yahoo" in low:
        return "Yahoo Finance"
    if "nasdaq" in low:
        return "NASDAQ"
    if "pitchbook" in low:
        return "PitchBook"
    if "nopat bridge" in low:
        return "NOPAT Bridge"
    if "revenue drivers" in low:
        return "Revenue Drivers"
    if "wacc" in low:
        return "WACC Tab"
    if "scenarios" in low:
        return "Scenarios"
    if len(v) > 40:
        return v[:40].rstrip(" ,;:")
    return v


def _clean_dcf_tab(dcf) -> None:
    _clear_col(dcf, "B")
    _clear_col(dcf, "D")
    _clear_col(dcf, "M")
    for r in range(1, dcf.max_row + 1):
        label = str(dcf.cell(r, 1).value or "")
        src = _standardize_dcf_source(dcf.cell(r, 3).value, label)
        dcf.cell(r, 3).value = src
    if isinstance(dcf["A4"].value, str):
        dcf["A4"].value = "Assumptions Guide (PDF) — see Cover tab link."
    dcf["E5"].value = "=Scenarios!F25/(1+Scenarios!$F$4)"
    _delete_rows_matching(dcf, "memo:")


def _round_scenario_matrix(ws, bear_col: int = 5, base_col: int = 6, bull_col: int = 7) -> None:
    for r in range(4, 11):
        for c in (bear_col, base_col, bull_col):
            v = ws.cell(r, c).value
            if isinstance(v, float) and abs(v) < 1:
                ws.cell(r, c).value = round(v, 4)
            elif isinstance(v, float):
                ws.cell(r, c).value = round(v, 3)


def _global_scrub(wb) -> None:
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    v = cell.value
                    v = re.sub(r"Firecrawl[- ]?ingested[^.]*\.?", "", v, flags=re.I)
                    v = re.sub(r"Firecrawl[:\s]*", "", v, flags=re.I)
                    v = re.sub(r"agent workflow", "operating adjustments", v, flags=re.I)
                    v = re.sub(r"Ctrl\+F[^.\n]*", "", v, flags=re.I)
                    v = re.sub(r"\n{2,}", "\n", v).strip()
                    cell.value = v or None


def polish_workbook(path: Path = TARGET) -> Path:
    tmp = path.with_suffix(".polishing.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)

    # --- 1. Cover ---
    cov = wb["Cover"]
    cov["B12"].value = "Red font = analyst assumptions (cols B/C on driver tabs)"
    b17 = str(cov["B17"].value or "")
    cov["B17"].value = re.sub(r"\s*Ctrl+F:.*", "", b17, flags=re.I).strip()
    for r in (20, 21, 22, 23):
        v = cov.cell(r, 2).value
        if isinstance(v, str):
            cov.cell(r, 2).value = v.split("\nCtrl+F")[0].split("Ctrl+F")[0].strip()

    # --- 2. WACC ---
    wacc = wb["WACC"]
    _delete_col_and_fix_refs(wb, "WACC", "D")
    _shorten_col_b(wacc, WACC_SHORT)
    _clean_source_col(wacc, 3)
    _delete_rows_matching(
        wacc,
        "β walkthrough",
        "Yahoo βL",
        "Unlever at Yahoo",
        "Relever at WACC",
        "Book D/E 44.69%",
    )
    _delete_rows_matching(wacc, "memo:")

    # --- 3. Scenarios ---
    scn = wb["Scenarios"]
    _delete_col_and_fix_refs(wb, "Scenarios", "D")
    _shorten_col_b(scn, SCEN_SHORT)
    _clean_source_col(scn, 3)
    _round_scenario_matrix(scn, 5, 6, 7)
    _delete_rows_matching(scn, "memo:")

    # --- 4. Revenue Drivers ---
    rd = wb["Revenue Drivers"]
    rd["A2"].value = (
        "FY2025A = historical 10-K anchors (blue). "
        "FY26–30 = red operational assumptions."
    )
    for col in ("I", "J", "L"):
        _delete_col_and_fix_refs(wb, "Revenue Drivers", col)
    _delete_rows_matching(rd, "memo:", "Note: DCF / Scenarios")
    for r in range(1, rd.max_row + 1):
        for c in range(1, rd.max_column + 1):
            v = rd.cell(r, c).value
            if isinstance(v, str) and "Justification [cols" in v:
                rd.cell(r, c).value = None

    # --- 5. NOPAT Bridge ---
    npb = wb["NOPAT Bridge"]
    npb["A4"].value = "Phases 1–4 clean and forecast operating profit."
    _delete_col_and_fix_refs(wb, "NOPAT Bridge", "D")
    _shorten_col_b(npb, NOPAT_SHORT)
    _delete_nopat_phase5_block(npb)

    # --- 6. DCF (clear B/D/M — keep column C as standardized sources) ---
    dcf = wb["DCF"]
    _clean_dcf_tab(dcf)

    # --- 7. Comps ---
    comps = wb["Comps"]
    _delete_col_and_fix_refs(wb, "Comps", "D")
    _delete_rows_matching(
        comps,
        "memo: current EV",
        "memo: current P / E",
        "ALO YOGA BUILD",
        "Column E on the next two rows",
        "Alo / Color Image",
        "Color Image parent sales",
        "Alo implied EV/Sales",
        "Alo EV/EBITDA",
        "Rationale for selected exit",
        "• Selected exit",
        "• At base WACC",
        "• PitchBook pubcomps",
        "• That tape is a check",
        "• Alo has no EV/EBITDA",
        "• PitchBook LULU",
    )
    comps.cell(4, 5).value = "=DCF!E5"
    comps.cell(7, 5).value = "=Scenarios!F188"

    _global_scrub(wb)

    wb.save(tmp)
    tmp.replace(path)
    return path


def _verify(path: Path) -> None:
    wb = openpyxl.load_workbook(path, data_only=False)
    issues = []

    if "Ctrl+F" in str(wb["Cover"].cell(20, 2).value or ""):
        issues.append("Cover row 20 still has Ctrl+F")

    wacc = wb["WACC"]
    if wacc.max_column >= 4 and wacc.cell(2, 4).value and "Ctrl+F" in str(wacc.cell(2, 4).value):
        issues.append("WACC col D not removed")

    dcf = wb["DCF"]
    assert dcf["E5"].value == "=Scenarios!F25/(1+Scenarios!$F$4)", dcf["E5"].value
    if dcf["B5"].value:
        issues.append("DCF B5 not cleared")
    if dcf["D5"].value:
        issues.append("DCF D5 not cleared")

    npb = wb["NOPAT Bridge"]
    for r in range(1, npb.max_row + 1):
        txt = str(npb.cell(r, 1).value or "")
        if "Phase 5" in txt or "Reconciliation workflow" in txt:
            issues.append(f"NOPAT row {r} still has Phase 5 block")

    if issues:
        raise AssertionError("\n".join(issues))

    print(f"Polished {path}")
    print(f"  DCF E5 = {dcf['E5'].value}")
    print(f"  Scenarios matrix E4:G4 = {[wb['Scenarios'].cell(4, c).value for c in range(5, 8)]}")


if __name__ == "__main__":
    out = polish_workbook()
    _verify(out)
