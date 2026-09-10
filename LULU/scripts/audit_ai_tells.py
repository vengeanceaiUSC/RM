#!/usr/bin/env python3
"""Read-only scan for AI/template tells that automated humanizers may miss.

Does NOT modify the workbook. Does NOT flag hardcoded values or Source links as errors —
only reports phrasing, structure, and presentation issues for a human or AI to fix manually.

Run:  cd LULU && python3 scripts/audit_ai_tells.py
Exit code is always 0 (report only).
"""
from __future__ import annotations

import re
import sys
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "unbeiesgbar_final.xlsx"

# Patterns that suggest AI/template origin (visible text + comments).
PHRASE_PATTERNS: list[tuple[str, re.Pattern[str]]] = [
    ("prompt debris", re.compile(r"justification|~20\s*word|click\s*\+|\[cols\s|ctrl\+f|source\s*\(click\)", re.I)),
    ("pipeline jargon", re.compile(r"convention a|phase [1-5]|5-phase pipeline|agent workflow", re.I)),
    ("comment boilerplate", re.compile(
        r"high-conviction|read-through|margin-of-safety|operating assumption with|"
        r"model cross-reference|on reported baseline|durable .{0,20} anchor",
        re.I,
    )),
    ("consulting tone", re.compile(
        r"underwrite|hangs together|valuation anchor|sanity check|not the exit|pitch deck reads|"
        r"bottom-up path|stays viable",
        re.I,
    )),
    ("template anchor", re.compile(r"\banchor\.|\boverlay\b", re.I)),
    ("in-cell URL", re.compile(r"https?\s*:", re.I)),
    ("openpyxl metadata", re.compile(r"openpyxl", re.I)),
]

VALUE_COLS: dict[str, tuple[int, ...]] = {
    "WACC": (4,),
    "Scenarios": (5, 6, 7),
    "NOPAT Bridge": (4, 5, 6, 7, 8, 9),
    "DCF": tuple(range(4, 11)),
    "Comps": tuple(range(4, 11)),
}

SOURCE_COL: dict[str, int] = {
    "WACC": 3,
    "Scenarios": 2,
    "NOPAT Bridge": 8,
    "Comps": 8,
    "DCF": 8,
    "Revenue Drivers": 10,
}

NOTE_COL: dict[str, int] = {
    "WACC": 2,
    "Scenarios": 2,
    "NOPAT Bridge": 2,
    "Comps": 2,
    "Revenue Drivers": 9,
}

BLUE_RGB = {"000000CC", "0000CC", "FF0000CC"}


@dataclass
class Finding:
    category: str
    location: str
    detail: str
    severity: str  # high | medium | low


def _is_formula(val) -> bool:
    return isinstance(val, str) and val.startswith("=")


def _long_float(val) -> bool:
    if not isinstance(val, float):
        return False
    s = f"{val:.12f}".rstrip("0")
    return "." in s and len(s.split(".")[1]) > 3


def _has_human_display_format(cell, label: str) -> bool:
    """True when number_format makes the sheet look analyst-entered (not General)."""
    fmt = str(cell.number_format or "General")
    if fmt == "General":
        return False
    lab = label.lower()
    if any(k in lab for k in ("dso", "dio", "dpo", "days")):
        return fmt.startswith("0.0")
    if any(k in lab for k in ("%", "margin", "growth", "rate", "erp", "weight", "wacc")):
        return "%" in fmt
    if abs(float(cell.value or 0)) >= 1000:
        return "#" in fmt
    return fmt not in ("General", "0")


def _cell_texts(ws, cell) -> list[tuple[str, str]]:
    out: list[tuple[str, str]] = []
    if isinstance(cell.value, str) and not _is_formula(cell.value):
        out.append(("cell", cell.value))
    if cell.comment and cell.comment.text:
        out.append(("comment", cell.comment.text))
    return out


def audit_workbook(path: Path = TARGET) -> list[Finding]:
    if not path.exists():
        raise FileNotFoundError(path)

    wb = openpyxl.load_workbook(path, data_only=False)
    findings: list[Finding] = []
    comment_prefixes: dict[str, list[str]] = defaultdict(list)

    # Metadata
    creator = str(wb.properties.creator or "")
    if "openpyxl" in creator.lower():
        findings.append(Finding("metadata", "Document properties", f"Creator={creator!r}", "high"))

    for ws in wb.worksheets:
        sheet = ws.title

        # Collapsible outline groups (+/- on row headers)
        if any(
            (rd.outlineLevel or 0) > 0 for rd in ws.row_dimensions.values()
        ) or any(
            (cd.outlineLevel or 0) > 0 for cd in ws.column_dimensions.values()
        ):
            findings.append(Finding("structure", sheet, "Outline levels present (+/- groups)", "high"))

        for row in ws.iter_rows():
            for cell in row:
                loc = f"{sheet}!{cell.coordinate}"

                for label, pat in PHRASE_PATTERNS:
                    if label == "openpyxl metadata":
                        continue
                    for kind, text in _cell_texts(ws, cell):
                        if pat.search(text):
                            # in-cell URL in value cols is high; in comments medium
                            sev = "high" if label == "in-cell URL" and kind == "cell" else "medium"
                            if label == "prompt debris":
                                sev = "high"
                            findings.append(
                                Finding(label, loc, f"({kind}) {text[:90]}", sev)
                            )

                if cell.comment and cell.comment.text:
                    prefix = cell.comment.text[:70].strip()
                    comment_prefixes[prefix].append(loc)

        # Per-sheet assumption / source checks
        if sheet not in VALUE_COLS:
            continue

        src_col = SOURCE_COL.get(sheet, 3)
        note_col = NOTE_COL.get(sheet)

        for r in range(1, ws.max_row + 1):
            label = str(ws.cell(r, 1).value or "").strip()
            if not label or label.lower() in {"notes", "source", "driver", "company"}:
                continue

            # Hardcode missing Source (flag only — do not imply deletion of value)
            for vcol in VALUE_COLS[sheet]:
                vcell = ws.cell(r, vcol)
                if vcell.value is None or _is_formula(vcell.value):
                    continue
                scell = ws.cell(r, src_col)
                if not scell.value and not scell.hyperlink:
                    findings.append(
                        Finding(
                            "missing source",
                            f"{sheet}!{scell.coordinate}",
                            f"Hardcode at {vcell.coordinate} ({label[:40]}) has no Source",
                            "medium",
                        )
                    )
                # Input not blue
                if isinstance(vcell.value, (int, float)) and vcell.font and vcell.font.color:
                    rgb = str(vcell.font.color.rgb or "").upper()
                    if rgb and rgb not in BLUE_RGB:
                        findings.append(
                            Finding(
                                "formatting",
                                f"{sheet}!{vcell.coordinate}",
                                f"Hardcode not blue ({rgb}) — {label[:35]}",
                                "low",
                            )
                        )
                if _long_float(vcell.value):
                    if _has_human_display_format(vcell, label):
                        continue  # sheet displays rounded; stored value kept for DCF
                    findings.append(
                        Finding(
                            "precision",
                            f"{sheet}!{vcell.coordinate}",
                            f"{vcell.value} — {label[:35]}",
                            "medium",
                        )
                    )

            if note_col:
                ncell = ws.cell(r, note_col)
                if isinstance(ncell.value, str) and ncell.value not in ("Notes", "Source"):
                    if len(ncell.value.split()) > 8 and "\n" not in ncell.value:
                        findings.append(
                            Finding(
                                "notes length",
                                f"{sheet}!{ncell.coordinate}",
                                ncell.value[:80],
                                "low",
                            )
                        )

            # Formula color
            for vcol in VALUE_COLS[sheet]:
                vcell = ws.cell(r, vcol)
                if not _is_formula(vcell.value) or not vcell.font or not vcell.font.color:
                    continue
                rgb = str(vcell.font.color.rgb or "").upper()
                val = vcell.value
                expected = "GREEN" if "!" in val else "BLACK"
                ok = (
                    (expected == "GREEN" and rgb in {"00006100", "FF006100", "006100"})
                    or (expected == "BLACK" and rgb in {"000000", "FF000000", "00000000"})
                )
                if rgb and not ok:
                    findings.append(
                        Finding(
                            "formatting",
                            f"{sheet}!{vcell.coordinate}",
                            f"Formula color {rgb} (expect {expected})",
                            "low",
                        )
                    )

    for prefix, locs in comment_prefixes.items():
        if len(locs) >= 4:
            findings.append(
                Finding(
                    "duplicate comments",
                    f"{len(locs)} cells",
                    f"{prefix[:65]}… → {', '.join(locs[:3])}",
                    "medium",
                )
            )

    return findings


def _print_report(findings: list[Finding], path: Path) -> None:
    by_cat: dict[str, list[Finding]] = defaultdict(list)
    for f in findings:
        by_cat[f.category].append(f)

    print(f"AI tell audit (read-only): {path.name}")
    print(f"Total findings: {len(findings)}\n")

    order = [
        "metadata",
        "prompt debris",
        "structure",
        "in-cell URL",
        "pipeline jargon",
        "comment boilerplate",
        "consulting tone",
        "template anchor",
        "missing source",
        "precision",
        "duplicate comments",
        "notes length",
        "formatting",
    ]
    for cat in order:
        items = by_cat.get(cat, [])
        if not items:
            continue
        high = sum(1 for i in items if i.severity == "high")
        print(f"## {cat} ({len(items)}; {high} high)")
        for item in items[:12]:
            print(f"  [{item.severity}] {item.location}: {item.detail}")
        if len(items) > 12:
            print(f"  … +{len(items) - 12} more")
        print()

    print("---")
    print("Protected (do NOT auto-delete): hardcoded values, Source hyperlinks, formulas.")
    print("Precision: cells with 0.0 / 0.0% formats display cleanly; formula bar may still")
    print("show full floats — re-type rounded inputs in Excel only after F9 confirms ~$133.64.")
    print("Fix flagged items by rewriting text or formatting — see LULU/AI_SELF_AUDIT.md")


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else TARGET
    findings = audit_workbook(target)
    _print_report(findings, target)
