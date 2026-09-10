#!/usr/bin/env python3
"""Replace DCF Scenarios row-pull formulas with local rolling P&L math.

Forecast cols C–G (FY26–30); FY25 anchor col B. Assumption links stay on Scenarios.

Run after move_sources_to_comments.py (value col B = FY25, C = FY26, …).

Run:  cd LULU && python3 scripts/rebuild_dcf_local_formulas.py
"""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "unbeiesgbar_final.xlsx"

SCN = "Scenarios"
FCOLS = ["C", "D", "E", "F", "G"]  # FY26–FY30 after B/C delete
FY25 = "B"


def _row(dcf, *prefixes: str) -> int | None:
    for r in range(1, dcf.max_row + 1):
        lab = str(dcf.cell(r, 1).value or "").strip()
        for p in prefixes:
            if lab.lower().startswith(p.lower()) or lab.lower() == p.lower():
                return r
    return None


def rebuild(path: Path = TARGET) -> dict[str, int]:
    tmp = path.with_suffix(".local_dcf.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)
    dcf = wb["DCF"]
    stats = {"rows": 0}

    r_rev = _row(dcf, "Net revenue")
    r_g = _row(dcf, "Revenue growth")
    r_gm = _row(dcf, "Gross margin")
    r_cogs = _row(dcf, "Cost of goods sold", "COGS")
    r_em = _row(dcf, "Clean / run-rate EBIT margin", "EBIT margin")
    r_tar = _row(dcf, "IEEPA tariff", "tariff refund")
    r_ebit = _row(dcf, "EBIT (incl", "EBIT")
    r_nopat = _row(dcf, "NOPAT")
    r_da = _row(dcf, "depreciation", "D&A")
    r_dap = _row(dcf, "D&A %")
    r_cx = _row(dcf, "capital expenditures", "Capex")
    r_cxp = _row(dcf, "Capex %")

    if not all([r_rev, r_gm, r_cogs, r_em, r_ebit]):
        raise ValueError("Missing core DCF P&L rows")

    # Revenue roll-forward
    dcf[f"C{r_rev}"] = f"={FY25}{r_rev}*(1+{SCN}!$D$4)"
    for i, col in enumerate(FCOLS[1:], start=1):
        prev = FCOLS[i - 1]
        dcf[f"{col}{r_rev}"] = f"={prev}{r_rev}*(1+{SCN}!$D$5)"
    stats["rows"] += 1

    # Gross margin % — assumption link (one cell, not row-pull)
    for col in FCOLS:
        dcf[f"{col}{r_gm}"] = f"={SCN}!$D$14"
    stats["rows"] += 1

    # COGS = Rev × (1 − GM%)
    for col in FCOLS:
        dcf[f"{col}{r_cogs}"] = f"={col}{r_rev}*(1-{col}{r_gm})"
    stats["rows"] += 1

    # EBIT margin ramp (5-year path)
    for i, col in enumerate(FCOLS):
        dcf[f"{col}{r_em}"] = (
            f"={SCN}!$D$6+({SCN}!$D$8-{SCN}!$D$6)*{i}/4"
        )
    stats["rows"] += 1

    # Tariff refund — FY26 only
    if r_tar:
        dcf[f"C{r_tar}"] = f"={SCN}!$D$7"
        for col in FCOLS[1:]:
            dcf[f"{col}{r_tar}"] = 0
        stats["rows"] += 1

    # EBIT = Rev × margin (+ tariff in FY26)
    dcf[f"C{r_ebit}"] = f"=C{r_rev}*C{r_em}+C{r_tar}" if r_tar else f"=C{r_rev}*C{r_em}"
    for col in FCOLS[1:]:
        dcf[f"{col}{r_ebit}"] = f"={col}{r_rev}*{col}{r_em}"
    stats["rows"] += 1

    # NOPAT — link to NOPAT Bridge forecast row
    if r_nopat:
        for col in FCOLS:
            dcf[f"{col}{r_nopat}"] = f"='NOPAT Bridge'!{col}41"
        stats["rows"] += 1

    # D&A and Capex as % of revenue
    if r_da and r_dap:
        for col in FCOLS:
            dcf[f"{col}{r_da}"] = f"={col}{r_rev}*{SCN}!$D$12"
    if r_cx and r_cxp:
        for col in FCOLS:
            dcf[f"{col}{r_cx}"] = f"=-{col}{r_rev}*{SCN}!$D$13"
    stats["rows"] += 1

    # Revenue growth % row (derived, stays local)
    if r_g:
        dcf[f"C{r_g}"] = f"=C{r_rev}/{FY25}{r_rev}-1"
        for i, col in enumerate(FCOLS[1:], start=1):
            dcf[f"{col}{r_g}"] = f"={col}{r_rev}/{FCOLS[i-1]}{r_rev}-1"

    wb.save(tmp)
    tmp.replace(path)
    return stats


if __name__ == "__main__":
    print(f"Rebuilt local DCF P&L: {rebuild()}")
