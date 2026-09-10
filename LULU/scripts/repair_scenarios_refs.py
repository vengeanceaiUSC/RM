#!/usr/bin/env python3
"""Fix Scenarios column refs after Notes+Source restore.

After restore_source_columns.py inserts cols B+C, base-case locked assumptions
live in Scenarios col F (was D in stripped model). Some formulas incorrectly
reference col H (empty) instead of F — breaking D&A %, Capex %, WACC, terminal
growth, and Gordon growth TV.

Run:  cd LULU && python3 scripts/repair_scenarios_refs.py
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "unbeiesgbar_final.xlsx"

# Base-case assumption rows on Scenarios (locked refs should point to col F).
ASSUMPTION_ROWS = frozenset(range(4, 23))

# Wrong col H → correct col F for absolute locked refs to assumption rows.
_H_TO_F = re.compile(
    r"(Scenarios!\$)H(\$(?:" + "|".join(str(r) for r in ASSUMPTION_ROWS) + r"))",
    re.I,
)


def repair_formula(formula: str) -> tuple[str, int]:
    """Return (new_formula, replacement_count)."""
    new, n = _H_TO_F.subn(r"\1F\2", formula)
    return new, n


def repair_workbook(path: Path = TARGET) -> tuple[int, list[str]]:
    if not path.exists():
        raise FileNotFoundError(path)

    tmp = path.with_suffix(".repairing.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)
    total = 0
    changes: list[str] = []

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if not isinstance(val, str) or not val.startswith("="):
                    continue
                if "Scenarios!" not in val or "$H$" not in val:
                    continue
                new_val, n = repair_formula(val)
                if n:
                    cell.value = new_val
                    total += n
                    if len(changes) < 30:
                        changes.append(f"{ws.title}!{cell.coordinate}: {val} → {new_val}")

    wb.save(tmp)
    tmp.replace(path)
    return total, changes


def verify(path: Path = TARGET) -> None:
    wb = openpyxl.load_workbook(path, data_only=False)
    issues: list[str] = []

    remaining = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if isinstance(val, str) and "Scenarios!$H$" in val:
                    remaining += 1
                    if remaining <= 5:
                        issues.append(f"Still broken: {ws.title}!{cell.coordinate}: {val}")

    if remaining:
        raise AssertionError(
            f"{remaining} Scenarios!$H$ refs remain:\n" + "\n".join(issues)
        )

    dcf = wb["DCF"]
    scn = wb["Scenarios"]

    checks = {
        "DCF E15 (D&A %)": (str(dcf["E15"].value), "=Scenarios!$F$12"),
        "DCF E17 (Capex %)": (str(dcf["E17"].value), "=Scenarios!$F$13"),
        "DCF D43 (terminal g)": (str(dcf["D43"].value), "=Scenarios!$F$10"),
        "DCF D45 (Gordon TV)": ("$F$9" in str(dcf["D45"].value or ""), True),
        "Scenarios F12 (D&A %)": (scn["F12"].value, 0.045),
        "Scenarios F13 (Capex %)": (scn["F13"].value, 0.055),
        "Scenarios F55 (D&A yr1)": (str(scn["F55"].value), "=F25*F12"),
        "Scenarios F60 (Capex yr1)": (str(scn["F60"].value), "=F25*F13"),
    }
    for name, (got, want) in checks.items():
        if got != want:
            issues.append(f"{name}: got {got!r}, want {want!r}")

    if issues:
        raise AssertionError("Verification failed:\n" + "\n".join(issues))

    print("Verify OK: Scenarios base-case refs point to col F; D&A/Capex intact")


if __name__ == "__main__":
    n, samples = repair_workbook()
    print(f"Repaired {n} Scenarios!$H$ → $F$ refs in {TARGET.name}")
    for line in samples[:15]:
        print(f"  {line}")
    if len(samples) > 15:
        print(f"  ... and {len(samples) - 15} more")
    verify()
