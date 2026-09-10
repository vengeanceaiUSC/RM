#!/usr/bin/env python3
"""Apply IB presentation defaults without changing values.

- Hide gridlines on model tabs
- Strip '(Phase N)' from NOPAT Bridge subtotal labels (pipeline jargon tell)

Run:  cd LULU && python3 scripts/fix_ib_formatting.py
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "unbeiesgbar_final.xlsx"

MODEL_SHEETS = (
    "Cover",
    "WACC",
    "Scenarios",
    "Revenue Drivers",
    "NOPAT Bridge",
    "DCF",
    "Comps",
    "Scratch",
)

PHASE_LABEL = re.compile(r"\s*\(Phase \d+\)\s*$")


def fix(path: Path = TARGET) -> dict[str, int]:
    tmp = path.with_suffix(".ib_fmt.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)
    stats = {"gridlines_off": 0, "labels_cleaned": 0}

    for name in MODEL_SHEETS:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        if ws.sheet_view.showGridLines is not False:
            ws.sheet_view.showGridLines = False
            stats["gridlines_off"] += 1

    ws = wb["NOPAT Bridge"]
    for r in range(1, ws.max_row + 1):
        cell = ws.cell(r, 1)
        if not isinstance(cell.value, str):
            continue
        new = PHASE_LABEL.sub("", cell.value).strip()
        if new != cell.value:
            cell.value = new
            stats["labels_cleaned"] += 1

    wb.save(tmp)
    tmp.replace(path)
    return stats


if __name__ == "__main__":
    result = fix()
    print(f"IB formatting: {result}")
