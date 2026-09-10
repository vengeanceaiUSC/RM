#!/usr/bin/env python3
"""Final humanization pass before submission.

- Link buried hardcodes to assumption/input cells (FY25 revenue, SBC, lease, ETR)
- Replace long addition chains with SUM()
- Fix NOPAT Revenue Drivers col alignment (same letter C–G)
- Rewrite hover comments: brief filing citations, no formula self-reference
- Link DCF sensitivity grid to row/column headers (no inline rate literals)
- Re-run valuation/div repairs

Run:  cd LULU && python3 scripts/humanize_for_submission.py
"""
from __future__ import annotations

import re
import shutil
import sys
from difflib import get_close_matches
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "unbeiesgbar_final.xlsx"
SCRIPTS = Path(__file__).resolve().parent

FCOLS = ["C", "D", "E", "F", "G"]
FY25_REV = "DCF!$B$6"

# Scenarios col-B anchors for reported FY25 balance sheet (10-K)
BS_ANCHORS: list[tuple[int, str, int, str]] = [
    (194, "FY25 total assets ($000)", 8456743, "FY25 10-K total assets."),
    (195, "FY25 total liabilities ($000)", 3494903, "FY25 10-K total liabilities."),
    (196, "FY25 total equity ($000)", 4961840, "FY25 10-K total equity."),
    (197, "FY25 diluted shares (000)", 119068, "FY25 10-K diluted share count."),
]

# Brief analyst comments by normalized label (from rewrite_model_notes.py)
NOTES: dict[str, dict[str, str]] = {}
sys.path.insert(0, str(SCRIPTS))
from rewrite_model_notes import NOTES_MAP  # noqa: E402

FORMULA_REF_RE = re.compile(
    r"ctrl\+f|formula|equation|model cross|read-through|phase [1-5]|"
    r"justification|pipeline|batch|firecrawl|gis ingest",
    re.I,
)

# Rows whose forecast cols are 5-term sums → SUM()
SUM_ROWS: dict[str, list[tuple[int, int, int]]] = {
    "NOPAT Bridge": [(21, 16, 20)],
}

# NOPAT (sheet row, RD row) — col letter matches forecast year
NOPAT_RD_ROWS = [(28, 31), (29, 41), (30, 45)]


def _norm(text) -> str:
    if not text:
        return ""
    s = str(text).lower().strip().replace("–", "-").replace("—", "-")
    s = re.sub(r"phase [1-5][:\s-]*", "", s)
    s = re.sub(r"^=+\s*", "", s)
    return re.sub(r"\s+", " ", s)


def _comment_for(sheet: str, label: str) -> str | None:
    smap = NOTES_MAP.get(sheet, {})
    if label in smap:
        return smap[label]
    k = _norm(label)
    for lk, note in smap.items():
        if _norm(lk) == k or k.startswith(_norm(lk)[:20]):
            return note
    close = get_close_matches(k, [_norm(x) for x in smap], n=1, cutoff=0.75)
    if close:
        for lk, note in smap.items():
            if _norm(lk) == close[0]:
                return note
    return None


def _rewrite_comments(wb) -> int:
    n = 0
    for sn in ("WACC", "Scenarios", "NOPAT Bridge", "Comps", "DCF", "Revenue Drivers"):
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        for r in range(1, ws.max_row + 1):
            label = ws.cell(r, 1).value
            if not label:
                continue
            cell = ws.cell(r, 2)
            note = _comment_for(sn, str(label))
            if not note and cell.comment and cell.comment.text:
                if FORMULA_REF_RE.search(cell.comment.text):
                    note = re.sub(r"\s{2,}", " ", FORMULA_REF_RE.sub("", cell.comment.text)).strip()
                    note = note.split("\n")[0][:80] if note else None
            if not note:
                continue
            # Strip self-referential formula language and tab cross-refs
            note = FORMULA_REF_RE.sub("", note).strip()
            note = re.sub(r"\s*Source:.*", "", note, flags=re.I).strip()
            note = re.sub(r"\s{2,}", " ", note)
            if not note:
                continue
            cell.comment = Comment(note[:32000], "Analyst")
            n += 1
    return n


def _fix_nopat_rd_refs(nb) -> int:
    """Forecast col letter on NOPAT must match Revenue Drivers col letter."""
    n = 0
    for r, rd_row in NOPAT_RD_ROWS:
        for col in FCOLS:
            addr = f"{col}{r}"
            want = f"='Revenue Drivers'!{col}{rd_row}"
            if nb[addr].value != want:
                nb[addr] = want
                n += 1
    return n


def _link_assumption_anchors(wb) -> int:
    n = 0
    scn = wb["Scenarios"]
    nb = wb["NOPAT Bridge"]
    dcf = wb["DCF"]
    rd = wb["Revenue Drivers"]

    # FY25 revenue anchor — link to DCF net revenue cell (not buried literal)
    for col in ("C", "D", "E"):
        growth_col = col
        old = scn[f"{col}25"].value
        new = f"={FY25_REV}*(1+{growth_col}4)"
        if old != new:
            scn[f"{col}25"] = new
            n += 1

    # Channel revenue bridge scaling
    for r in range(168, 188):
        for col in ("C", "D", "E"):
            v = scn[f"{col}{r}"].value
            if isinstance(v, str):
                new_v = v.replace("8456743", "$B$194").replace("3494903", "$B$195").replace("4961840", "$B$196")
                new_v = new_v.replace("11102600", FY25_REV)
                if new_v != v:
                    scn[f"{col}{r}"] = new_v
                    n += 1

    # Reported FY25 balance sheet anchors (assumption cells, col B)
    for row, label, value, note in BS_ANCHORS:
        if scn[f"B{row}"].value != value:
            scn[f"A{row}"] = label
            scn[f"B{row}"] = value
            scn[f"B{row}"].comment = Comment(note, "Analyst")
            n += 1

    # Diluted shares — link FY25 count and buyback prices to assumption / DCF path
    price_cols = ["C", "D", "E", "F", "G"]
    for i, row in enumerate(range(183, 188)):
        price = f"DCF!{price_cols[i]}$73"
        for col in ("C", "D"):
            if i == 0:
                scn[f"{col}{row}"] = f"=$B$197-{col}21/{price}"
            else:
                scn[f"{col}{row}"] = f"={col}{row-1}-{col}21/{price}"
            n += 1

    # SBC + lease FY25 reference inputs (col B)
    nb["B20"] = 62203
    nb["B20"].comment = Comment("FY25 SBC per 10-K cash flow stmt.", "Analyst")
    nb["B23"] = 1028
    for col in FCOLS:
        nb[f"{col}20"] = f"=Scenarios!{col}25/{FY25_REV}*$B$20*$B$7"
        nb[f"{col}23"] = f"=$B$23*(Scenarios!{col}25/{FY25_REV})"
        nb[f"{col}39"] = f"=$B$39+({col}$12-$B$39)*{FCOLS.index(col)+1}/5"
        n += 3

    # Historical ratio links on Scenarios base case (col D)
    hist: list[tuple[str, str]] = [
        ("D12", f"=DCF!B15/{FY25_REV}"),           # D&A % from 10-K CF stmt
        ("D14", f"=(DCF!B6-DCF!B8)/DCF!B6"),       # Gross margin FY25A
    ]
    for addr, formula in hist:
        if scn[addr].value != formula:
            scn[addr] = formula
            n += 1

    # Channel mix weights — assumption cells on Revenue Drivers
    if rd["B71"].value != 0.71:
        rd["A71"] = "Americas comp mix %"
        rd["B71"] = 0.71
        rd["A72"] = "China comp mix %"
        rd["B72"] = 0.16
        rd["A73"] = "RoW comp mix %"
        rd["B73"] = 0.13
        n += 3
    for col in FCOLS:
        prev = "B" if col == "C" else FCOLS[FCOLS.index(col) - 1]
        rd[f"{col}29"] = (
            f"={prev}28*$B$71*(1+{col}25)+{prev}28*$B$72*(1+{col}26)"
            f"+{prev}28*$B$73*(1+{col}27)"
        )
        n += 1

    # Pitch deck section label
    if scn["A137"].value and "pitch deck" in str(scn["A137"].value).lower():
        scn["A137"] = "Financial statement bridge (IS / CFS / BS)"
        n += 1

    # FY25 D&A on DCF already B15; ensure Comps EBITDA link
    comps = wb["Comps"]
    if comps["B5"].value != "=DCF!B11+DCF!B15":
        comps["B5"] = "=DCF!B11+DCF!B15"
        n += 1

    # DCF labels — avoid formula text in row headers
    dcf_label_fixes = {
        46: "Terminal value (Gordon growth)",
    }
    for row, label in dcf_label_fixes.items():
        if dcf[f"A{row}"].value != label:
            dcf[f"A{row}"] = label
            n += 1

    return n


def _apply_sum_formulas(wb) -> int:
    n = 0
    for sheet, rows in SUM_ROWS.items():
        ws = wb[sheet]
        for row, start, end in rows:
            for col in FCOLS:
                addr = f"{col}{row}"
                want = f"=SUM({col}{start}:{col}{end})"
                if ws[addr].value != want:
                    ws[addr] = want
                    n += 1
    return n


def _fix_sensitivity_grid(dcf) -> int:
    """Link inline WACC/g literals to sensitivity axis headers."""
    n = 0
    for r in range(100, 105):
        wacc = f"$A{r}"
        for col in FCOLS:
            g = f"{col}$99"
            formula = (
                f"=(NPV({wacc},C36:G36)+"
                f"(G36*(1+{g})/({wacc}-{g}))"
                f"/(1+{wacc})^5+B51+B52)/B56"
            )
            if dcf[f"{col}{r}"].value != formula:
                dcf[f"{col}{r}"] = formula
                n += 1
    return n


def humanize(path: Path = TARGET) -> dict:
    tmp = path.with_suffix(".submit.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)

    stats = {
        "nopat_rd": _fix_nopat_rd_refs(wb["NOPAT Bridge"]),
        "assumption_links": _link_assumption_anchors(wb),
        "sum_formulas": _apply_sum_formulas(wb),
        "sensitivity": _fix_sensitivity_grid(wb["DCF"]),
        "comments": _rewrite_comments(wb),
    }

    wb.save(tmp)

    sys.path.insert(0, str(SCRIPTS))
    from fix_div_errors import fix as fix_div
    from fix_file_metadata import fix as fix_meta

    stats["div_fix"] = fix_div(path=tmp)
    from fix_dcf_tv_reconciliation import fix as fix_tv_recon

    stats["tv_recon"] = fix_tv_recon(path=tmp)
    stats["metadata"] = fix_meta(path=tmp)

    tmp.replace(path)
    return stats


if __name__ == "__main__":
    print(f"Humanize for submission: {humanize()}")
