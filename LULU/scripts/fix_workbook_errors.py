#!/usr/bin/env python3
"""Fix Excel formula-error popups and outline interaction blockers.

Problems addressed:
- NOPAT Bridge subtotal labels stored as invalid formulas (= Adjusted EBIT...)
- DCF D4 / Revenue Drivers K56 left as sheet-ref strings that Excel parses as formulas
- Non-contiguous B + D column outline groups (Source col C in between) freeze +/- UI

Run:  cd LULU && python3 scripts/fix_workbook_errors.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.hyperlink import Hyperlink

SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPTS))
from restore_outline_groups import restore_outline_groups  # noqa: E402

ROOT = SCRIPTS.parent
TARGETS = [
    ROOT / "model18_humanized.xlsx",
    ROOT / "model18_humanized (1).xlsx",
    ROOT / "model18_wsp_formulas.xlsx",
]

LABEL_FORMULA = re.compile(r"^=\s+(.+)$")
SHEET_REF_TEXT = re.compile(r"^'?.+?'?!")

LINK_FIXES: dict[tuple[str, str], tuple[str, str]] = {
    ("DCF", "C4"): ("NOPAT Bridge tab", "'NOPAT Bridge'!A1"),
    ("DCF", "D4"): ("Revenue Drivers tab", "'Revenue Drivers'!A1"),
    ("Revenue Drivers", "K56"): ("Scenarios tab: base revenue", "'Scenarios'!G25"),
}


def _set_internal_link(cell: Cell, text: str, location: str) -> None:
    cell.value = text
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=location)


def fix_workbook(path: Path) -> list[str]:
    wb = openpyxl.load_workbook(path)
    changes: list[str] = []

    ws = wb["NOPAT Bridge"]
    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell, Cell):
                continue
            v = cell.value
            if not isinstance(v, str):
                continue
            m = LABEL_FORMULA.match(v)
            if not m:
                continue
            new = m.group(1).strip()
            cell.value = new
            changes.append(f"NOPAT Bridge {cell.coordinate}: invalid formula label -> {new!r}")

    for (sheet, coord), (text, location) in LINK_FIXES.items():
        if sheet not in wb.sheetnames:
            continue
        cell = wb[sheet][coord]
        current = cell.value
        if isinstance(current, str) and SHEET_REF_TEXT.match(current.strip()):
            _set_internal_link(cell, text, location)
            changes.append(f"{sheet} {coord}: sheet-ref text -> hyperlink {text!r}")
        elif cell.hyperlink is None and text:
            _set_internal_link(cell, text, location)
            changes.append(f"{sheet} {coord}: added hyperlink {text!r}")

    restore_outline_groups(wb)
    changes.append("restored row-only outline groups")

    wb.save(path)
    return changes


def verify(path: Path) -> None:
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["NOPAT Bridge"]
    for coord in ("A21", "A26", "A35", "A36"):
        v = ws[coord].value
        if isinstance(v, str) and v.startswith("="):
            raise AssertionError(f"{coord} still formula-like: {v!r}")

    d4 = wb["DCF"]["D4"]
    if isinstance(d4.value, str) and "!" in d4.value:
        raise AssertionError(f"D4 still sheet-ref text: {d4.value!r}")

    ws = wb["Scenarios"]
    for col in "BD":
        dim = ws.column_dimensions.get(col)
        if dim and (dim.outline_level or dim.hidden):
            raise AssertionError(f"Scenarios col {col} still outline-hidden")


if __name__ == "__main__":
    for path in TARGETS:
        if not path.exists():
            print(f"Skip missing: {path.name}")
            continue
        changes = fix_workbook(path)
        verify(path)
        print(f"Fixed {path.name} ({len(changes)} changes)")
        for c in changes:
            print(f"  {c}")
