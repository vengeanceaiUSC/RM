#!/usr/bin/env python3
"""Remove AI leftover 'spot check' column from Scenarios tab (col I rows 3–4).

Run:  cd LULU && python3 scripts/remove_spot_check.py
"""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "unbeiesgbar_final.xlsx"


def remove(path: Path = TARGET) -> list[str]:
    tmp = path.with_suffix(".no_spot.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)
    ws = wb["Scenarios"]
    changes: list[str] = []
    for coord in ("I3", "I4"):
        if ws[coord].value is not None:
            changes.append(f"{coord}: {ws[coord].value!r} → cleared")
            ws[coord].value = None
            ws[coord].fill = openpyxl.styles.PatternFill()  # reset yellow/default
    wb.save(tmp)
    tmp.replace(path)
    return changes


if __name__ == "__main__":
    cleared = remove()
    print(f"Cleared {len(cleared)} spot-check cells")
    for c in cleared:
        print(f"  {c}")
