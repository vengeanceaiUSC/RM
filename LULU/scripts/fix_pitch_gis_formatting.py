#!/usr/bin/env python3
"""Apply GIS formatting rules to LULU_Investment_Pitch_Deck.pptx.

Fixes:
- Slide 1 TOC: rebuild # / Slide tables with Garamond, wide # column, no wrap
- Deck-wide: set Garamond on every text run (GIS rule #2)

Run:  cd LULU && python3 scripts/fix_pitch_gis_formatting.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "LULU_Investment_Pitch_Deck.pptx"

sys.path.insert(0, str(ROOT.parent / "GIS"))
from gis_pitch import FONT, INK, LGREY, NAVY, WHITE, _set_font  # noqa: E402

TOC_COL0_W = 1.08  # inches — wide enough for two-digit # in Garamond 9pt (no wrap)
TOC_FONT = 9.0
TOC_HDR_FONT = 9.5
TOC_TOP = 1.32
TOC_HEIGHT = 5.16
TOC_WIDTH = 6.02
TOC_LEFTS = (0.50, 6.80)


def _delete_shape(shape) -> None:
    el = shape.element
    el.getparent().remove(el)


def _read_toc_entries(slide) -> list[tuple[str, str]]:
    entries: list[tuple[str, str]] = []
    for shape in slide.shapes:
        if not shape.has_table:
            continue
        t = shape.table
        if t.rows[0].cells[0].text.strip() != "#":
            continue
        for ri in range(1, len(t.rows)):
            num = t.rows[ri].cells[0].text.strip()
            title = t.rows[ri].cells[1].text.strip()
            if num and title:
                entries.append((num, title))
    return entries


def _style_cell(cell, text: str, size: float, color, bold: bool = False, align=PP_ALIGN.LEFT) -> None:
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Pt(4)
    cell.margin_right = Pt(4)
    cell.margin_top = Pt(1)
    cell.margin_bottom = Pt(1)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    _set_font(r, size, color, bold=bold)


def _add_toc_table(slide, entries: list[tuple[str, str]], left: float) -> None:
    rows = [[num, title] for num, title in entries]
    nrows = len(rows) + 1
    shape = slide.shapes.add_table(
        nrows, 2, Inches(left), Inches(TOC_TOP), Inches(TOC_WIDTH), Inches(TOC_HEIGHT)
    )
    table = shape.table
    col0 = Inches(TOC_COL0_W)
    table.columns[0].width = int(col0)
    table.columns[1].width = int(Inches(TOC_WIDTH) - col0)

    row_h = int(Inches(TOC_HEIGHT) / nrows)
    for ri in range(nrows):
        table.rows[ri].height = row_h

    # header
    for ci, hdr in enumerate(("#", "Slide")):
        cell = table.cell(0, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        _style_cell(
            cell, hdr, TOC_HDR_FONT, WHITE, bold=True,
            align=PP_ALIGN.CENTER if ci == 0 else PP_ALIGN.LEFT,
        )

    for ri, (num, title) in enumerate(rows, start=1):
        for ci, val in enumerate((num, title)):
            cell = table.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if ri % 2 else LGREY
            _style_cell(
                cell, val, TOC_FONT, INK,
                align=PP_ALIGN.CENTER if ci == 0 else PP_ALIGN.LEFT,
            )


def _rebuild_toc(slide) -> None:
    entries = _read_toc_entries(slide)
    if not entries:
        return
    split = (len(entries) + 1) // 2
    left_entries = entries[:split]
    right_entries = entries[split:]
    for shape in list(slide.shapes):
        if shape.has_table and shape.table.rows[0].cells[0].text.strip() == "#":
            _delete_shape(shape)
    _add_toc_table(slide, left_entries, TOC_LEFTS[0])
    _add_toc_table(slide, right_entries, TOC_LEFTS[1])


def _run_color(run) -> object:
    try:
        if run.font.color and run.font.color.rgb:
            return run.font.color.rgb
    except AttributeError:
        pass
    return INK


def _apply_garamond_deck(prs: Presentation) -> int:
    n = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if run.text:
                            _set_font(
                                run,
                                run.font.size.pt if run.font.size else 11,
                                _run_color(run),
                                bold=bool(run.font.bold),
                                italic=bool(run.font.italic),
                            )
                            n += 1
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for para in cell.text_frame.paragraphs:
                            for run in para.runs:
                                if run.text:
                                    _set_font(
                                        run,
                                        run.font.size.pt if run.font.size else 10,
                                        _run_color(run),
                                        bold=bool(run.font.bold),
                                        italic=bool(run.font.italic),
                                    )
                                    n += 1
    return n


def main() -> None:
    from add_gis_cover_slide import cover_offset

    prs = Presentation(str(DECK))
    _rebuild_toc(prs.slides[cover_offset(prs)])
    runs = _apply_garamond_deck(prs)
    prs.save(str(DECK))
    print(f"GIS formatting applied → {DECK}")
    print(f"  TOC # column: {TOC_COL0_W}\" | Garamond runs set: {runs}")


if __name__ == "__main__":
    main()
