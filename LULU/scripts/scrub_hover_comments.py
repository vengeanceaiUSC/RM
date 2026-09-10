#!/usr/bin/env python3
"""Shorten Analyst hover comments — remove AI boilerplate, keep URLs.

Does not remove comments where Source col C already has a link (those were
stripped by remove_comment_artifacts.py). Cleans remaining hover text on DCF
sensitivity rows and any stale 10-K batch comments.

Run:  cd LULU && python3 scripts/scrub_hover_comments.py
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "unbeiesgbar_final.xlsx"

AI_PHRASES = re.compile(
    r"high-conviction|read-through|margin-of-safety|operating assumption with|"
    r"model cross-reference|on reported baseline|durable .{0,24} anchor|"
    r"FY25 10-K anchor\.|FRED macro anchor\.|cash-flow read-through\.",
    re.I,
)

URL_RE = re.compile(r"(https?://[^\s]+)", re.I)

# (sheet, row) → one-line comment when label known
ROW_COMMENT: dict[tuple[str, int], str] = {
    ("DCF", 5): "Net revenue FY25 10-K.",
    ("DCF", 8): "COGS FY25 10-K.",
    ("DCF", 11): "EBIT incl. FY26 refund FY25 10-K.",
    ("DCF", 12): "Reported EBIT margin FY25 10-K.",
    ("DCF", 100): "WACC sensitivity −100 bps vs base.",
    ("DCF", 101): "WACC sensitivity −50 bps vs base.",
    ("DCF", 102): "WACC sensitivity +50 bps vs base.",
    ("DCF", 103): "WACC sensitivity +100 bps vs base.",
    ("Comps", 18): "Peer EV/EBITDA from PitchBook.",
    ("Comps", 56): "Football-field range from comp set.",
}


def _tailor_comment(sheet: str, row: int, label: str, old: str) -> str:
    key = (sheet, row)
    if key in ROW_COMMENT:
        base = ROW_COMMENT[key]
    elif "10-k" in old.lower() or "sec.gov" in old.lower():
        short_label = (label or "Line item").strip()[:40]
        base = f"{short_label} — FY25 10-K."
    elif "wacc tab" in old.lower():
        base = re.sub(r"\s*on reported baseline\.?\s*", " ", old, flags=re.I).strip()
        base = re.sub(r"^\d+\.\d+\s*", "", base)
        base = base.replace("WACC tab to green WACC cell ~9.0%.", "WACC sensitivity vs base.")
    else:
        base = AI_PHRASES.sub("", old)
        base = re.sub(r"\s{2,}", " ", base).strip()
        base = re.sub(r"\.\s*\.", ".", base)

    urls = URL_RE.findall(old)
    if urls and urls[0] not in base:
        base = f"{base} Source: {urls[0]}"
    return base.strip()


def scrub(path: Path = TARGET) -> list[str]:
    tmp = path.with_suffix(".comments_scrub.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)
    changes: list[str] = []

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not cell.comment or not cell.comment.text:
                    continue
                old = cell.comment.text.strip()
                if not AI_PHRASES.search(old) and (ws.title, cell.row) not in ROW_COMMENT:
                    continue
                label = str(ws.cell(cell.row, 1).value or "")
                new = _tailor_comment(ws.title, cell.row, label, old)
                if new and new != old:
                    cell.comment.text = new
                    changes.append(f"{ws.title}!{cell.coordinate}")

    wb.save(tmp)
    tmp.replace(path)
    return changes


if __name__ == "__main__":
    changed = scrub()
    print(f"Scrubbed {len(changed)} hover comments")
    for c in changed:
        print(f"  {c}")
