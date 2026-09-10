#!/usr/bin/env python3
"""Compact Scenarios layout: label | Bear | Base | Bull | Source | notes.

Removes the wide empty gap (old cols D–E) that split assumptions from values.

Run:  cd LULU && python3 scripts/fix_scenarios_layout.py
"""
from __future__ import annotations

import re
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "model18_wsp_formulas.xlsx"

ANCHORS = {194: 8_456_743, 195: 3_494_903, 196: 4_961_840, 197: 119_068, 198: 11_102_600, 202: 62_203, 203: 1_028}
SCENARIO_REMAP = {"F": "B", "G": "C", "H": "D"}


def _remap_formula(formula: str) -> str:
    if not isinstance(formula, str) or not formula.startswith("="):
        return formula

    def repl(m: re.Match[str]) -> str:
        col = m.group(2)
        return f"{m.group(1)}{SCENARIO_REMAP.get(col, col)}{m.group(3)}{m.group(4)}"

    return re.sub(r"(?<![A-Z])(\$?)([FGH])(\$?)(\d+)", repl, formula)


def _fix_anchors(formula: str) -> str:
    if not isinstance(formula, str) or not formula.startswith("="):
        return formula
    for row in ANCHORS:
        formula = formula.replace(f"$F${row}", f"$B${row}")
        formula = re.sub(rf"(?<![A-Z])F{row}\b", f"B{row}", formula)
    return formula


def fix_workbook(path: Path = TARGET) -> None:
    wb = openpyxl.load_workbook(path)
    scn = wb["Scenarios"]

    snap: dict[tuple[int, int], object] = {}
    for r in range(1, 204):
        for c in range(1, 9):
            cell = scn.cell(r, c)
            snap[(r, c)] = {
                "v": cell.value,
                "nf": cell.number_format,
                "font": copy(cell.font),
                "fill": copy(cell.fill),
                "align": copy(cell.alignment),
            }

    for r in range(1, 204):
        for c in range(1, 7):
            scn.cell(r, c).value = None

    place = {"A": 1, "B": 6, "C": 7, "D": 8, "E": 3, "F": 2}
    for new_c, old_c in place.items():
        nc = ord(new_c) - ord("A") + 1
        for r in range(1, 204):
            src = snap[(r, old_c)]
            v = src["v"]
            if isinstance(v, str) and v.startswith("="):
                v = _fix_anchors(_remap_formula(v))
            dst = scn.cell(r, nc)
            dst.value = v
            dst.number_format = src["nf"]
            dst.font = src["font"]
            dst.fill = src["fill"]
            dst.alignment = src["align"]
            dst.border = Border()

    for r, val in ANCHORS.items():
        scn.cell(r, 2).value = val
        scn.cell(r, 2).number_format = "#,##0"

    scn["E3"] = "Source (click)"
    scn["F3"] = "How it moves"
    for ref in ("E3", "F3"):
        scn[ref].font = Font(bold=True, size=9)
        scn[ref].alignment = Alignment(horizontal="left", wrap_text=True)

    hdr = {
        "B2": ("C00000", "Bear"),
        "C2": ("548235", "Base"),
        "D2": ("1F3864", "Bull"),
    }
    for ref, (color, text) in hdr.items():
        scn[ref].value = text
        scn[ref].font = Font(bold=True, color="FFFFFF", size=10)
        scn[ref].fill = PatternFill("solid", fgColor=color)
        scn[ref].alignment = Alignment(horizontal="center")

    scn.column_dimensions["A"].width = 42
    scn.column_dimensions["B"].width = 11
    scn.column_dimensions["C"].width = 11
    scn.column_dimensions["D"].width = 11
    scn.column_dimensions["E"].width = 22
    scn.column_dimensions["F"].width = 28
    for col in "GHIJK":
        scn.column_dimensions[col].width = 2

    for ws in wb.worksheets:
        if ws.title == "Scenarios":
            continue
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "Scenarios!" in cell.value:
                    cell.value = re.sub(
                        r"Scenarios!\$?([FGH])\$?(\d+)",
                        lambda m: f"Scenarios!${SCENARIO_REMAP[m.group(1)]}${m.group(2)}",
                        cell.value,
                    )

    wb.save(path)
    print(f"Scenarios layout fixed → {path}")


if __name__ == "__main__":
    fix_workbook()
