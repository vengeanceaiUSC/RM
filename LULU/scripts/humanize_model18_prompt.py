#!/usr/bin/env python3
"""Humanize model18unaltered workbook per analyst submission spec.

- Clear all Justification and Ctrl+F columns (preserve Source columns).
- Scrub Firecrawl / Agent AI jargon across remaining cells.
- Delete NOPAT Bridge rows 43–49 and Cover pipeline instruction rows.
- Save as model18_humanized.xlsx without altering formulas or numeric values.
"""
from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "model18unaltered (12).xlsx"
OUTPUT = ROOT / "model18_humanized.xlsx"

FIRECRAWL_RE = re.compile(r"firecrawl", re.I)
NOT_PCT_PREFIX = re.compile(r"^Not % of revenue;\s*", re.I)
AGENT_WORKFLOW_RE = re.compile(
    r"agent\s+(?:workflow|execution)|phase\s*5\s*[—–-]\s*agent",
    re.I,
)

COVER_PIPELINE_ROW_RE = re.compile(
    r"justification\s*\|\s*source|"
    r"cols\s+b/?c/?d|"
    r"ctrl\+f\s*(?:\(|:|\"|\→|->)|"
    r"firecrawl|"
    r"data pipeline|"
    r"every red assumption.*ctrl\+f|"
    r"every red assumption on wacc",
    re.I,
)


def resolve_input(arg: str | None) -> Path:
    if arg:
        p = ROOT / arg if not Path(arg).is_absolute() else Path(arg)
        if not p.exists():
            raise FileNotFoundError(f"Input workbook not found: {p}")
        return p
    for name in (
        "model18unaltered (13).xlsx",
        "model18unaltered (12).xlsx",
        "model18unaltered.xlsx",
    ):
        p = ROOT / name
        if p.exists():
            return p
    raise FileNotFoundError("No model18unaltered workbook found in LULU/")


def _header_text(cell_value) -> str:
    return str(cell_value or "").strip()


def _is_justification_header(text: str) -> bool:
    low = text.lower().strip()
    if "justification" not in low or len(low) > 90:
        return False
    return any(token in low for token in ("expand", "~20", "cols", "justification  ["))


def _is_source_header(text: str) -> bool:
    low = text.lower().strip()
    if "source" not in low or "ctrl+f" in low or len(low) > 60:
        return False
    return "click" in low or low in {"source", "alt. source"}


def _is_ctrlf_header(text: str) -> bool:
    low = text.lower().strip()
    if "ctrl+f" not in low or len(low) > 60:
        return False
    if "→" in low or "->" in low or '"' in low:
        return False
    return low.startswith("ctrl+f") or low.startswith("alt. ctrl+f")


def find_doc_columns(ws, max_scan_rows: int = 6) -> tuple[set[int], set[int], set[int]]:
    """Return (justification_cols, source_cols, ctrlf_cols) by header text."""
    justification: set[int] = set()
    sources: set[int] = set()
    ctrlf: set[int] = set()

    for row in range(1, min(max_scan_rows, ws.max_row) + 1):
        for col in range(1, ws.max_column + 1):
            header = _header_text(ws.cell(row, col).value)
            if not header:
                continue
            if _is_justification_header(header):
                justification.add(col)
            elif _is_source_header(header):
                sources.add(col)
            elif _is_ctrlf_header(header):
                ctrlf.add(col)
    return justification, sources, ctrlf


def clear_doc_column(ws, col: int) -> int:
    """Clear analyst prose in doc columns; keep formulas and numeric values."""
    cleared = 0
    for row in range(1, ws.max_row + 1):
        cell = ws.cell(row, col)
        val = cell.value
        if val is None:
            continue
        if isinstance(val, str) and val.startswith("="):
            continue
        if isinstance(val, (int, float)):
            continue
        cell.value = None
        cleared += 1
    return cleared


CTRLF_LINE = re.compile(
    r"^\s*Ctrl\+F\s*(?:\([^)]*\))?\s*[\"']?[^\"'\n]*[\"']?\s*(?:→|->|-\>)\s*.*$",
    re.I | re.M,
)
CTRLF_PREFIX = re.compile(r"Ctrl\+F\s*(?:\([^)]*\))?\s*[:\"]\s*", re.I)
CTRLF_ANY = re.compile(r"Ctrl\+F", re.I)


def scrub_cell_text(text: str) -> str | None:
    if text.startswith("="):
        return text

    out = text
    if AGENT_WORKFLOW_RE.search(out):
        return None

    out = FIRECRAWL_RE.sub("Historical data", out)
    out = NOT_PCT_PREFIX.sub("", out)

    # Drop standalone Ctrl+F proof lines and inline Ctrl+F prefixes in prose cells.
    kept: list[str] = []
    for line in out.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if CTRLF_LINE.match(stripped) or stripped.lower().startswith("ctrl+f"):
            continue
        line = CTRLF_PREFIX.sub("", line)
        line = CTRLF_ANY.sub("", line)
        if line.strip():
            kept.append(line.strip())
    out = "\n".join(kept)

    out = re.sub(r"\(\s*Historical data\s*:\s*[^)]*\)", "", out, flags=re.I)
    out = re.sub(r"Historical data[:\s-]*", "Historical data ", out, flags=re.I)
    out = re.sub(r"\s{2,}", " ", out).strip()

    if not out:
        return None
    return out


def cover_row_is_pipeline_instruction(ws, row: int) -> bool:
    parts: list[str] = []
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row, col).value
        if isinstance(val, str) and val.strip():
            parts.append(val.strip())
    if not parts:
        return False
    blob = "\n".join(parts)
    return bool(COVER_PIPELINE_ROW_RE.search(blob))


def delete_rows(ws, rows: list[int]) -> None:
    for row in sorted(rows, reverse=True):
        ws.delete_rows(row, 1)


def humanize_workbook(src: Path, dst: Path = OUTPUT) -> Path:
    tmp = dst.with_suffix(".tmp.xlsx")
    shutil.copy2(src, tmp)
    wb = openpyxl.load_workbook(tmp)

    stats = {
        "justification_cells_cleared": 0,
        "ctrlf_cells_cleared": 0,
        "text_cells_scrubbed": 0,
        "cover_rows_deleted": 0,
        "nopat_rows_deleted": 0,
    }

    if "Cover" in wb.sheetnames:
        ws = wb["Cover"]
        rows_to_delete = [
            row
            for row in range(1, ws.max_row + 1)
            if cover_row_is_pipeline_instruction(ws, row)
        ]
        delete_rows(ws, rows_to_delete)
        stats["cover_rows_deleted"] = len(rows_to_delete)

    for ws in wb.worksheets:
        justification_cols, source_cols, ctrlf_cols = find_doc_columns(ws)
        for col in justification_cols:
            stats["justification_cells_cleared"] += clear_doc_column(ws, col)
        for col in ctrlf_cols:
            stats["ctrlf_cells_cleared"] += clear_doc_column(ws, col)

        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str) or cell.value.startswith("="):
                    continue
                if (
                    cell.column in justification_cols
                    or cell.column in ctrlf_cols
                    or cell.column in source_cols
                ):
                    continue
                cleaned = scrub_cell_text(cell.value)
                if cleaned != cell.value:
                    cell.value = cleaned
                    stats["text_cells_scrubbed"] += 1

    if "NOPAT Bridge" in wb.sheetnames:
        ws = wb["NOPAT Bridge"]
        start, end = 43, 49
        if ws.max_row >= start:
            delete_count = min(end, ws.max_row) - start + 1
            delete_rows(ws, list(range(start, start + delete_count)))
            stats["nopat_rows_deleted"] = delete_count

    wb.save(tmp)
    tmp.replace(dst)

    print(f"Saved {dst}")
    for key, val in stats.items():
        print(f"  {key}: {val}")
    return dst


def verify(src: Path, dst: Path) -> None:
    src_wb = openpyxl.load_workbook(src, data_only=False)
    dst_wb = openpyxl.load_workbook(dst, data_only=False)

    banned = re.compile(
        r"firecrawl|agent\s+workflow|agent\s+execution|ctrl\+f|not % of revenue;",
        re.I,
    )
    skip_sheets = {"Cover", "NOPAT Bridge"}

    for sheet in src_wb.sheetnames:
        if sheet in skip_sheets:
            continue
        s_ws, d_ws = src_wb[sheet], dst_wb[sheet]
        for row in range(1, s_ws.max_row + 1):
            for col in range(1, s_ws.max_column + 1):
                sv = s_ws.cell(row, col).value
                dv = d_ws.cell(row, col).value
                if isinstance(sv, str) and sv.startswith("="):
                    if sv != dv:
                        raise AssertionError(
                            f"Formula changed: {sheet}!{get_column_letter(col)}{row}"
                        )
                elif isinstance(sv, (int, float)):
                    if sv != dv:
                        raise AssertionError(
                            f"Number changed: {sheet}!{get_column_letter(col)}{row}"
                        )

    for ws in dst_wb.worksheets:
        _, source_cols, _ = find_doc_columns(ws)
        for col in source_cols:
            nonempty = sum(
                1
                for row in range(1, ws.max_row + 1)
                if ws.cell(row, col).value not in (None, "")
            )
            if nonempty == 0:
                raise AssertionError(
                    f"Source column {get_column_letter(col)} empty on {ws.title}"
                )

        for row in ws.iter_rows():
            for cell in row:
                if cell.column in source_cols:
                    continue
                if isinstance(cell.value, str) and banned.search(cell.value):
                    raise AssertionError(
                        f"Residual banned text: {ws.title}!{cell.coordinate}: {cell.value[:80]}"
                    )

    if "NOPAT Bridge" in dst_wb.sheetnames:
        ws = dst_wb["NOPAT Bridge"]
        for row in range(43, 50):
            if row <= ws.max_row:
                vals = [ws.cell(row, c).value for c in range(1, 5)]
                if any(
                    isinstance(v, str) and "agent" in v.lower() for v in vals if v
                ):
                    raise AssertionError(f"Agent rows still present at row {row}")

    dcf = dst_wb["DCF"]
    assert str(dcf["E56"].value).startswith("=")
    assert dcf["E5"].value == 11102600
    print("Verification passed.")


if __name__ == "__main__":
    input_path = resolve_input(sys.argv[1] if len(sys.argv) > 1 else None)
    print(f"Input: {input_path.name}")
    out = humanize_workbook(input_path)
    verify(input_path, out)
