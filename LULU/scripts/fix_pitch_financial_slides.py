#!/usr/bin/env python3
"""Fix financial slide layout (14–16): footer overflow into source table.

GIS footer is 0.30" tall — long source paragraphs bleed upward over the
per-row source table. Shorten footer to one line and give the source
table enough height.

Run:  cd LULU && python3 scripts/fix_pitch_financial_slides.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "LULU_Investment_Pitch_Deck.pptx"

sys.path.insert(0, str(ROOT.parent / "GIS"))
from gis_pitch import FONT, GREY, INK, LGREY, NAVY, WHITE, _set_font  # noqa: E402

FIN_SLIDES = (14, 15, 16)
FOOTER_TEXT = (
    "Source: LULU_DCF_Valuation_Model.xlsx · Hist: FY22–25 10-K · "
    "Forecast: Scenarios col C (base) · Row sources in table below"
)
SRC_TOP = 5.68
SRC_HEIGHT = 1.18
SRC_COL0 = 1.85


def _is_source_table(shape) -> bool:
    if not shape.has_table:
        return False
    return shape.table.rows[0].cells[0].text.strip() == "Line item"


def _is_footer(shape) -> bool:
    if not shape.has_text_frame:
        return False
    t = shape.text_frame.text
    return t.startswith("Forecast:") or t.startswith("Source:") or "Per-row forecast sources" in t


def _style_source_table(table) -> None:
    nrows = len(table.rows)
    row_h = int(Inches(SRC_HEIGHT) / nrows)
    for ri in range(nrows):
        table.rows[ri].height = row_h
    total_w = table.columns[0].width + table.columns[1].width + table.columns[2].width
    table.columns[0].width = int(Inches(SRC_COL0))
    rest = int(total_w - Inches(SRC_COL0))
    table.columns[1].width = rest // 2
    table.columns[2].width = rest - rest // 2
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Pt(4)
            cell.margin_right = Pt(4)
            cell.margin_top = Pt(1)
            cell.margin_bottom = Pt(1)
            tf = cell.text_frame
            tf.word_wrap = True
            for para in tf.paragraphs:
                para.alignment = PP_ALIGN.LEFT
                for run in para.runs:
                    if not run.text:
                        continue
                    color = WHITE if ri == 0 else INK
                    if ri == 0:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = NAVY
                    else:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = WHITE if ri % 2 else LGREY
                    _set_font(run, 7 if ri > 0 else 7.5, color, bold=(ri == 0))


def _fix_slide(slide) -> None:
    for shape in slide.shapes:
        if _is_footer(shape):
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = FOOTER_TEXT
            _set_font(r, 8, GREY)
            tf.word_wrap = False
        if _is_source_table(shape):
            shape.top = int(Inches(SRC_TOP))
            shape.height = int(Inches(SRC_HEIGHT))
            _style_source_table(shape.table)


def main() -> None:
    prs = Presentation(str(DECK))
    for n in FIN_SLIDES:
        _fix_slide(prs.slides[n - 1])
    prs.save(str(DECK))
    print(f"Fixed financial slides {FIN_SLIDES} → {DECK}")


if __name__ == "__main__":
    main()
