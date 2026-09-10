#!/usr/bin/env python3
"""Fix structural layout, internal consistency, and AI-tell sourcing in the model.

Run:  cd LULU && python3 scripts/fix_model_structure.py
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "unbeiesgbar_final.xlsx"

SCN = "Scenarios"
FCOLS = ["C", "D", "E", "F", "G"]
FY25 = "B"
BLUE = "0563C1"

SCN_D_TO_DCF: dict[int, tuple[int, str]] = {}
for i, scn_row in enumerate(range(25, 30)):
    SCN_D_TO_DCF[scn_row] = (6, FCOLS[i])
for i, scn_row in enumerate(range(35, 40)):
    SCN_D_TO_DCF[scn_row] = (8, FCOLS[i])
for i, scn_row in enumerate(range(40, 45)):
    SCN_D_TO_DCF[scn_row] = (9, FCOLS[i])
for i, scn_row in enumerate(range(45, 50)):
    SCN_D_TO_DCF[scn_row] = (11, FCOLS[i])
for i, scn_row in enumerate(range(50, 55)):
    SCN_D_TO_DCF[scn_row] = (14, FCOLS[i])
for i, scn_row in enumerate(range(55, 60)):
    SCN_D_TO_DCF[scn_row] = (17, FCOLS[i])
for i, scn_row in enumerate(range(60, 65)):
    SCN_D_TO_DCF[scn_row] = (19, FCOLS[i])

_WC_BLOCKS = [
    (65, 21), (70, 22), (75, 23), (80, 24), (85, 25), (90, 26),
    (95, 28), (100, 30), (105, 31), (110, 32), (115, 33), (120, 34), (125, 36),
]
for scn_start, dcf_r in _WC_BLOCKS:
    for j in range(5):
        SCN_D_TO_DCF[scn_start + j] = (dcf_r, FCOLS[j])

HORIZ: list[tuple[str, int]] = [
    ("Net revenue ($000)", 6),
    ("COGS ($000)", 8),
    ("Gross margin %", 7),
    ("EBIT margin %", 9),
    ("EBIT ($000)", 11),
    ("NOPAT ($000)", 14),
    ("D&A ($000)", 17),
    ("Capex ($000)", 19),
    ("Unlevered FCF ($000)", 36),
]
YEAR_HDRS = ["FY26E", "FY27E", "FY28E", "FY29E", "FY30E"]

RD_SOURCE_FIXES: dict[str, str] = {
    "beginning-stores formula": "LULU FY2025 10-K: store roll-forward (Item 2)",
    "ending-stores formula": "LULU FY2025 10-K: company-operated store count",
    "total-stores formula": "LULU FY2025 10-K: total company-operated stores",
    "total-sqft formula": "LULU FY2025 10-K: sales per square foot & store count",
}

RD_NOTE_FIXES: dict[int, tuple[str, str]] = {
    7: ("Americas - openings", "Americas: ~6–8 net new stores/yr (slower than China)."),
    11: ("China Mainland - openings", "China: ~20 store openings in FY26 (expansion plan)."),
    15: ("Rest of World - openings", "RoW: ~4–6 net new stores/yr."),
}

COMPS_TICKER_SOURCES: dict[str, tuple[str, str]] = {
    "lululemon": ("Yahoo Finance: LULU statistics", "https://finance.yahoo.com/quote/LULU/key-statistics/"),
    "under armour": ("Yahoo Finance: UAA statistics", "https://finance.yahoo.com/quote/UAA/key-statistics/"),
    "adidas": ("Yahoo Finance: ADS.DE statistics", "https://finance.yahoo.com/quote/ADS.DE/key-statistics/"),
    "nike": ("Yahoo Finance: NKE statistics", "https://finance.yahoo.com/quote/NKE/key-statistics/"),
    "deckers": ("Yahoo Finance: DECK statistics", "https://finance.yahoo.com/quote/DECK/key-statistics/"),
    "williams-sonoma": ("Yahoo Finance: WSM statistics", "https://finance.yahoo.com/quote/WSM/key-statistics/"),
    "crocs": ("Yahoo Finance: CROX statistics", "https://finance.yahoo.com/quote/CROX/key-statistics/"),
    "movado": ("Yahoo Finance: MOV statistics", "https://finance.yahoo.com/quote/MOV/key-statistics/"),
    "levi": ("Yahoo Finance: LEVI statistics", "https://finance.yahoo.com/quote/LEVI/key-statistics/"),
    "la-z-boy": ("Yahoo Finance: LZB statistics", "https://finance.yahoo.com/quote/LZB/key-statistics/"),
    "kontoor": ("Yahoo Finance: KTB statistics", "https://finance.yahoo.com/quote/KTB/key-statistics/"),
}


def _row(ws, *prefixes: str) -> int | None:
    for r in range(1, ws.max_row + 1):
        lab = str(ws.cell(r, 1).value or "").strip().lower()
        for p in prefixes:
            if lab.startswith(p.lower()) or lab == p.lower():
                return r
    return None


def _rewrite_scenarios_ref(formula: str) -> str:
    if not isinstance(formula, str) or not formula.startswith("="):
        return formula

    def repl(m: re.Match) -> str:
        col = m.group(1)
        row = int(m.group(2))
        if col.upper() == "D" and row in SCN_D_TO_DCF:
            dcf_r, fc = SCN_D_TO_DCF[row]
            return f"DCF!{fc}{dcf_r}"
        return m.group(0)

    return re.sub(r"Scenarios!\$?([A-Z])\$?(\d+)", repl, formula, flags=re.I)


def fix_wacc(wb) -> int:
    ws = wb["WACC"]
    n = 0
    ws["A17"] = "Total debt equivalents (for beta unlever)"
    ws["A18"] = "Market value of equity (for beta unlever)"
    for addr, val in [("B17", "=B13"), ("B18", "=B10"), ("B19", "=B17/B18")]:
        if ws[addr].value != val:
            ws[addr] = val
            n += 1
    ws["C17"] = "Model: lease debt + funded debt (row 13)"
    ws["C18"] = "Model: share price × shares (row 10)"
    return n


def fix_dcf_native(wb) -> int:
    dcf = wb["DCF"]
    n = 0
    r_rev = _row(dcf, "net revenue") or 6
    r_g = _row(dcf, "revenue growth") or 5
    r_gm = _row(dcf, "gross margin") or 7
    r_cogs = _row(dcf, "cogs") or 8
    r_em = _row(dcf, "ebit margin") or 9
    r_tar = _row(dcf, "tariff") or 10
    r_ebit = _row(dcf, exact="EBIT") or 11
    r_ebit_m = _row(dcf, "reported ebit margin") or 12
    r_tax = _row(dcf, "less: taxes") or 13
    r_nopat = _row(dcf, "nopat") or 14
    r_dap = _row(dcf, "d&a %") or 16
    r_da = _row(dcf, "plus: d&a") or 17
    r_cxp = _row(dcf, "capex %") or 18
    r_cx = _row(dcf, "less: capex") or 19
    r_dso = _row(dcf, "dso") or 21
    r_ar = _row(dcf, "accounts receivable") or 22
    r_dio = _row(dcf, "dio") or 23
    r_inv = _row(dcf, "inventories") or 24
    r_dpo = _row(dcf, "dpo") or 25
    r_ap = _row(dcf, "accounts payable") or 26
    r_ppct = _row(dcf, "prepaid expenses (%)") or 27
    r_pp = _row(dcf, "prepaid expenses (other") or 28
    r_accpct = _row(dcf, "accrued liabilities (%)") or 29
    r_acc = _row(dcf, "accrued liabilities and") or 30
    r_coa = _row(dcf, "current operating assets") or 31
    r_col = _row(dcf, "current operating liabilities") or 32
    r_nwc = _row(dcf, "net working capital") or 33
    r_dnwc = _row(dcf, "delta nwc") or 34
    r_ufcf = _row(dcf, "unlevered free cash") or 36

    dcf[f"C{r_rev}"] = f"={FY25}{r_rev}*(1+{SCN}!$D$4)"
    prev = FY25
    for col in FCOLS:
        dcf[f"{col}{r_rev}"] = f"={prev}{r_rev}*(1+{col}{r_g})"
        prev = col
        n += 1

    dcf[f"C{r_g}"] = f"={SCN}!$D$4"
    for col in FCOLS[1:]:
        dcf[f"{col}{r_g}"] = f"={SCN}!$D$5"

    for col in FCOLS:
        dcf[f"{col}{r_cogs}"] = f"={col}{r_rev}*(1-{SCN}!$D$14)"
        dcf[f"{col}{r_gm}"] = f"=({col}{r_rev}-{col}{r_cogs})/{col}{r_rev}"
        n += 2

    dcf[f"C{r_em}"] = f"={SCN}!$D$6"
    step = f"({SCN}!$D$8-{SCN}!$D$6)/4"
    for i, col in enumerate(FCOLS[1:], start=1):
        dcf[f"{col}{r_em}"] = f"={FCOLS[i - 1]}{r_em}+{step}"
        n += 1

    dcf[f"C{r_tar}"] = f"={SCN}!$D$7"
    for col in FCOLS[1:]:
        dcf[f"{col}{r_tar}"] = 0

    dcf[f"C{r_ebit}"] = f"=C{r_rev}*C{r_em}+C{r_tar}"
    for col in FCOLS[1:]:
        dcf[f"{col}{r_ebit}"] = f"={col}{r_rev}*{col}{r_em}"
        n += 1

    for col in FCOLS:
        dcf[f"{col}{r_ebit_m}"] = f"={col}{r_ebit}/{col}{r_rev}"

    for col in FCOLS:
        dcf[f"{col}{r_tax}"] = f"=-'NOPAT Bridge'!{col}40"
        dcf[f"{col}{r_nopat}"] = f"='NOPAT Bridge'!{col}41"
        n += 2

    for col in FCOLS:
        dcf[f"{col}{r_dap}"] = f"={SCN}!$D$12"
        dcf[f"{col}{r_da}"] = f"={col}{r_rev}*{SCN}!$D$12"
        dcf[f"{col}{r_cxp}"] = f"={SCN}!$D$13"
        dcf[f"{col}{r_cx}"] = f"=-{col}{r_rev}*{SCN}!$D$13"
        n += 4

    for i, col in enumerate(FCOLS):
        dcf[f"{col}{r_dso}"] = f"={SCN}!$D$15"
        dcf[f"{col}{r_dio}"] = f"={SCN}!$D$16-{i}*{SCN}!$D$17"
        dcf[f"{col}{r_dpo}"] = f"={SCN}!$D$18"
        dcf[f"{col}{r_ar}"] = f"={col}{r_rev}*{col}{r_dso}/365"
        dcf[f"{col}{r_inv}"] = f"={col}{r_cogs}*{col}{r_dio}/365"
        dcf[f"{col}{r_ap}"] = f"={col}{r_cogs}*{col}{r_dpo}/365"
        dcf[f"{col}{r_ppct}"] = f"={SCN}!$D$19"
        dcf[f"{col}{r_pp}"] = f"={col}{r_rev}*{col}{r_ppct}"
        dcf[f"{col}{r_accpct}"] = f"={SCN}!$D$20"
        dcf[f"{col}{r_acc}"] = f"={col}{r_rev}*{col}{r_accpct}"
        dcf[f"{col}{r_coa}"] = f"={col}{r_ar}+{col}{r_inv}+{col}{r_pp}"
        dcf[f"{col}{r_col}"] = f"={col}{r_ap}+{col}{r_acc}"
        dcf[f"{col}{r_nwc}"] = f"={col}{r_coa}-{col}{r_col}"
        n += 12

    dcf[f"C{r_dnwc}"] = f"={FY25}{r_nwc}-C{r_nwc}"
    for i, col in enumerate(FCOLS[1:], start=1):
        dcf[f"{col}{r_dnwc}"] = f"={FCOLS[i - 1]}{r_nwc}-{col}{r_nwc}"
        n += 1
    for col in FCOLS:
        dcf[f"{col}{r_ufcf}"] = f"={col}{r_nopat}+{col}{r_da}+{col}{r_cx}+{col}{r_dnwc}"
        n += 1

    return n


def fix_scenarios_horizontal(wb) -> int:
    scn = wb[SCN]
    n = 0
    scn["A24"] = "Base case forecast ($000)"
    for i, hdr in enumerate(YEAR_HDRS):
        scn.cell(24, 4 + i).value = hdr
        n += 1
    start = 25
    for i, (label, dcf_row) in enumerate(HORIZ):
        r = start + i
        scn.cell(r, 1).value = label
        for j, fc in enumerate(FCOLS):
            want = f"=DCF!{fc}{dcf_row}"
            if scn.cell(r, 4 + j).value != want:
                scn.cell(r, 4 + j).value = want
                n += 1
    for r in range(start + len(HORIZ), 65):
        label = str(scn.cell(r, 1).value or "")
        if re.search(r"year \d|rev y|gm% y|ebit margin y|ebit y|nopat y|d&a y|capex -", label, re.I):
            scn.cell(r, 1).value = None
            for c in range(3, 9):
                scn.cell(r, c).value = None
            n += 1
    return n


def rewrite_workbook_refs(wb) -> int:
    n = 0
    for ws in wb.worksheets:
        if ws.title == SCN:
            continue
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if not isinstance(val, str) or not val.startswith("="):
                    continue
                new_val = _rewrite_scenarios_ref(val)
                if new_val != val:
                    cell.value = new_val
                    n += 1
    return n


def fix_nopat(wb) -> int:
    nb = wb["NOPAT Bridge"]
    n = 0
    sbc_pct = "$B$20/DCF!$B$6"
    for col in FCOLS:
        want = f"=IF($B$7=1,DCF!{col}6*{sbc_pct},0)"
        if nb[f"{col}20"].value != want:
            nb[f"{col}20"] = want
            n += 1
        want_l = f"=$B$23*(DCF!{col}6/DCF!$B$6)"
        if nb[f"{col}23"].value != want_l:
            nb[f"{col}23"] = want_l
            n += 1
        for r, dcf_r in ((16, 11), (36, 11)):
            want_e = f"=DCF!{col}{dcf_r}"
            if nb[f"{col}{r}"].value != want_e:
                nb[f"{col}{r}"] = want_e
                n += 1
    return n


def fix_revenue_drivers(wb) -> int:
    rd = wb["Revenue Drivers"]
    n = 0
    for r in range(1, rd.max_row + 1):
        label = str(rd.cell(r, 1).value or "")
        src = rd.cell(r, 10)
        if isinstance(src.value, str):
            for key, repl in RD_SOURCE_FIXES.items():
                if key in src.value.lower():
                    src.value = repl
                    src.font = Font(color=BLUE, italic=True, size=9)
                    n += 1
                    break
        for _, (match, note) in RD_NOTE_FIXES.items():
            if match.lower() in label.lower():
                if rd.cell(r, 9).value != note:
                    rd.cell(r, 9).value = note
                    n += 1
    for fc in FCOLS:
        if rd[f"{fc}56"].value != f"=DCF!{fc}6":
            rd[f"{fc}56"] = f"=DCF!{fc}6"
            n += 1
    return n


def fix_comps_sources(wb) -> int:
    ws = wb["Comps"]
    n = 0
    for r in range(1, ws.max_row + 1):
        label = str(ws.cell(r, 1).value or "").lower()
        src_cell = ws.cell(r, 8) if ws.cell(r, 8).value else ws.cell(r, 3)
        if not isinstance(src_cell.value, str):
            continue
        if "pitchbook comps set" in src_cell.value.lower():
            for key, (text, url) in COMPS_TICKER_SOURCES.items():
                if key in label:
                    src_cell.value = text
                    src_cell.hyperlink = url
                    src_cell.font = Font(color=BLUE, underline="single", italic=True, size=9)
                    n += 1
                    break
    return n


def fix_dcf_text(wb) -> int:
    dcf = wb["DCF"]
    n = 0
    if dcf["D95"].value and "Gordon identity" in str(dcf["D95"].value):
        dcf["D95"] = "Exit multiple = Gordon-implied (FY30 EBITDA)"
        n += 1
    return n


def fix(path: Path = TARGET) -> dict:
    tmp = path.with_suffix(".struct.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)
    stats = {
        "wacc": fix_wacc(wb),
        "dcf_native": fix_dcf_native(wb),
        "scenarios_horiz": fix_scenarios_horizontal(wb),
        "ref_rewrite": rewrite_workbook_refs(wb),
        "nopat": fix_nopat(wb),
        "revenue_drivers": fix_revenue_drivers(wb),
        "comps_sources": fix_comps_sources(wb),
        "dcf_text": fix_dcf_text(wb),
    }
    wb.save(tmp)
    tmp.replace(path)
    return stats


if __name__ == "__main__":
    print(f"Model structure fix: {fix()}")
