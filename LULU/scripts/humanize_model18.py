#!/usr/bin/env python3
"""Humanize model18unaltered (12).xlsx → unbeiesgbar2model.xlsx.

Strips AI/scraper tells (Ctrl+F logs, Firecrawl, agent refs) and tightens copy to
institutional analyst shorthand. Formulas and numeric hardcodes are never touched.

Run:  cd LULU && python3 scripts/humanize_model18.py
"""
from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

import openpyxl

SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPTS))
from restore_outline_groups import restore_outline_groups  # noqa: E402

ROOT = SCRIPTS.parent
INPUT = ROOT / "model18unaltered (12).xlsx"
OUTPUT = ROOT / "unbeiesgbar2model.xlsx"

AI_PHRASES = [
    r"firecrawl[- ]?ingested",
    r"firecrawl[:\s]*",
    r"firecrawl",
    r"agent workflow",
    r"\bagent\b",
    r"chatgpt",
    r"openai",
    r"anthropic",
    r"claude",
    r"gemini",
    r"\bllm\b",
    r"large language model",
    r"ai-generated",
    r"as an ai",
    r"please note that",
    r"it is important to consider",
    r"here is the",
    r"i hope this helps",
    r"in conclusion",
    r"pitch_values\.py",
    r"build_pitch\.py",
    r"ingest_kpis\.py",
]

AI_PATTERN = re.compile(r"(?i)(?:" + "|".join(AI_PHRASES) + r")")

CTRLF_LINE = re.compile(
    r"^\s*Ctrl\+F\s*(?:\([^)]*\))?\s*[\"']?[^\"'\n]*[\"']?\s*(?:→|->|-\>)\s*.*$",
    re.I | re.M,
)
CTRLF_INLINE = re.compile(
    r"Ctrl\+F\s*(?:\([^)]*\))?\s*[\"'][^\"']+[\"']\s*(?:→|->)\s*[^.\n]+",
    re.I,
)
CTRLF_PREFIX = re.compile(r"Ctrl\+F:\s*", re.I)

NOT_PCT_PREFIX = re.compile(r"^Not % of revenue;\s*", re.I)

ALSO_PREFIX = re.compile(r"(^|\n)\s*Also:\s*", re.I)

YAHOO_WALKTHROUGH = re.compile(
    r"YAHOO KEY STATISTICS;.*?(?=\n[A-Z]|\Z)",
    re.I | re.S,
)

# Map robotic DSO/DIO/DPO openers to one-line analyst notes.
WORKING_CAPITAL_SHORT = {
    "DSO": "DSO flat ~6.3 days — DTC/store cash sales dominate; AR is wholesale residual.",
    "DIO": "DIO: FY25 anchor ~129 days; −1 day/yr through FY30 (inventory discipline).",
    "DPO": "DPO flat ~25 days; AP scales with COGS (~3% of sales).",
    "SBC": "SBC stays in EBIT (Convention A); DCF uses basic shares.",
    "prepaid": "Prepaids ~5.1% of revenue (FY25 10-K OCA residual).",
    "accrued": "Accrued liabilities ~6.0% of revenue (FY25 10-K).",
}


def _collapse_whitespace(text: str) -> str:
    lines = [re.sub(r"\s+", " ", ln).strip() for ln in text.splitlines()]
    lines = [ln for ln in lines if ln]
    return "\n".join(lines)


def _humanize_ctrlf_lines(text: str) -> str:
    """Drop Ctrl+F proof lines; keep analyst justification prose."""
    kept: list[str] = []
    for line in text.splitlines():
        if CTRLF_LINE.match(line.strip()) or line.strip().lower().startswith("ctrl+f"):
            continue
        line = CTRLF_INLINE.sub("", line)
        line = CTRLF_PREFIX.sub("", line)
        if line.strip():
            kept.append(line.strip())
    return _collapse_whitespace("\n".join(kept))


def _shorten_working_capital(text: str) -> str:
    if NOT_PCT_PREFIX.search(text):
        body = NOT_PCT_PREFIX.sub("", text)
        low = body.lower()
        for key, short in WORKING_CAPITAL_SHORT.items():
            if key.lower() in low[:40]:
                return short
        # Generic fallback: drop prefix, keep first sentence only.
        first = re.split(r"(?<=[.!?])\s+", body.strip())[0]
        if len(first) > 120:
            first = first[:117].rsplit(" ", 1)[0] + "…"
        return first or body[:120]
    return text


def _trim_over_explanation(text: str) -> str:
    """Remove step-by-step arithmetic proofs; keep assumption headline."""
    if len(text) < 180:
        return text
    # Drop lines that are mostly formulas with = signs in prose.
    lines = []
    for line in text.splitlines():
        if re.search(r"\d+\s*[÷/]\s*\d+", line) and "Ctrl" not in line:
            if line.count("=") >= 2 or "→" in line:
                continue
        if re.match(r"^\s*STEP \d", line, re.I):
            continue
        if re.match(r"^\s*\(\d+\)", line):
            continue
        lines.append(line)
    out = _collapse_whitespace("\n".join(lines))
    # Cap very long cells at ~3 sentences.
    sentences = re.split(r"(?<=[.!?])\s+", out)
    if len(sentences) > 4:
        out = " ".join(sentences[:3])
    return out


def _apply_phrase_replacements(text: str) -> str:
    replacements = [
        (re.compile(r"Justification \| Source \| Ctrl\+F", re.I), "Justification | Source | Proof"),
        (re.compile(r"Alt\.\s*Ctrl\+F", re.I), "Alt. source"),
        (re.compile(r"\(unique Ctrl\+F\)", re.I), "(per FY25 10-K)"),
        (re.compile(r"No peer Ctrl\+F\.", re.I), "No peer print."),
        (re.compile(r"No Ctrl\+F for EV/EBITDA — ", re.I), ""),
        (re.compile(r"No Ctrl\+F\. ", re.I), ""),
        (re.compile(r"Ctrl\+F \"WACC\" and \"Terminal growth\"", re.I), "WACC and terminal growth rows"),
        (re.compile(r"Firecrawl-ingested\s*10-K anchors", re.I), "FY25 10-K historical anchors"),
        (re.compile(r"Firecrawl-ingested", re.I), "FY25 10-K"),
        (re.compile(r"Firecrawl found none[^.]*\.?", re.I), ""),
        (re.compile(r"Phase 5 is the agent workflow summary below", re.I),
         "Phase 5 summarizes the reconciliation workflow."),
        (re.compile(r"Phase 5 — Agent workflow", re.I), "Phase 5 — Reconciliation"),
        (re.compile(r"pulls cols G \(base\) & H \(bull\) from this block via pitch_values\.py", re.I),
         "Pitch deck reads base/bull scenario columns."),
        (re.compile(r"Nothing fancy, just how many doors you end with\.?", re.I), ""),
        (re.compile(r"That's where the year starts\.?", re.I), ""),
        (re.compile(r"Same number, just carried forward into the new fiscal year\.?", re.I),
         "Prior-year ending count carried forward."),
    ]
    out = text
    for pat, repl in replacements:
        out = pat.sub(repl, out)
    out = AI_PATTERN.sub("", out)
    out = ALSO_PREFIX.sub(r"\1", out)
    out = YAHOO_WALKTHROUGH.sub(
        "Hamada β: Yahoo βL 0.86; unlever at mkt D/E ~0.19 → βu ~0.76; "
        "relever at lease-adjusted D/E ~0.16 → β used ~0.84.",
        out,
    )
    return out


def humanize_text(text: str) -> str:
    if not text or text.startswith("="):
        return text
    out = text.replace("**", "")
    out = _apply_phrase_replacements(out)
    out = _humanize_ctrlf_lines(out)
    out = _shorten_working_capital(out)
    out = _trim_over_explanation(out)
    out = _collapse_whitespace(out)
    out = re.sub(r"\s+\.", ".", out)
    # Final pass: strip any remaining Ctrl+F mentions in labels or prose.
    out = re.sub(r"Ctrl\+F", "Source", out, flags=re.I)
    return out.strip()


def humanize_workbook(src: Path = INPUT, dst: Path = OUTPUT) -> Path:
    if not src.exists():
        raise FileNotFoundError(src)
    tmp = dst.with_suffix(".humanizing.xlsx")
    shutil.copy2(src, tmp)
    wb = openpyxl.load_workbook(tmp)

    changed = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if not isinstance(val, str) or val.startswith("="):
                    continue
                new_val = humanize_text(val)
                if new_val != val:
                    cell.value = new_val if new_val else None
                    changed += 1

    restore_outline_groups(wb)
    wb.save(tmp)
    tmp.replace(dst)
    print(f"Saved {dst} ({changed} text cells humanized)")
    return dst


def verify(src: Path = INPUT, dst: Path = OUTPUT) -> None:
    src_wb = openpyxl.load_workbook(src, data_only=False)
    dst_wb = openpyxl.load_workbook(dst, data_only=False)

    for sheet in src_wb.sheetnames:
        s_ws, d_ws = src_wb[sheet], dst_wb[sheet]
        for row in range(1, s_ws.max_row + 1):
            for col in range(1, s_ws.max_column + 1):
                sv = s_ws.cell(row, col).value
                if isinstance(sv, str) and sv.startswith("="):
                    dv = d_ws.cell(row, col).value
                    if sv != dv:
                        raise AssertionError(
                            f"Formula changed: {sheet}!"
                            f"{openpyxl.utils.get_column_letter(col)}{row}"
                        )
                elif isinstance(sv, (int, float)) or sv is None:
                    if sv != d_ws.cell(row, col).value:
                        raise AssertionError(
                            f"Number changed: {sheet}!"
                            f"{openpyxl.utils.get_column_letter(col)}{row}"
                        )

    residual = {"ctrlf": 0, "firecrawl": 0, "agent": 0, "not_pct": 0}
    for ws in dst_wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not isinstance(v, str) or v.startswith("="):
                    continue
                if "ctrl+f" in v.lower():
                    residual["ctrlf"] += 1
                if "firecrawl" in v.lower():
                    residual["firecrawl"] += 1
                if re.search(r"agent workflow|\bfirecrawl\b", v, re.I):
                    residual["agent"] += 1
                if v.lower().startswith("not % of revenue"):
                    residual["not_pct"] += 1

    print("Residual flags:", residual)
    if any(residual.values()):
        raise AssertionError(f"Humanization incomplete: {residual}")

    wacc = dst_wb["WACC"]
    assert str(wacc["E25"].value).startswith("=")
    assert wacc["E17"].value == 0.86
    print("Verify OK: formulas/numbers intact, AI tells removed")


if __name__ == "__main__":
    if not INPUT.exists():
        shutil.copy2(ROOT / "model18unaltered.xlsx", INPUT)
    out = humanize_workbook()
    verify(INPUT, out)
