#!/usr/bin/env python3
"""Final label/voice cleanup on unbeiesgbar_final.xlsx — no number changes.

Run: cd LULU && python3 scripts/final_label_cleanup.py
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.hyperlink import Hyperlink

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "unbeiesgbar_final.xlsx"

LABEL_SUBS = [
    (re.compile(r"NOPAT NORMALIZATION PIPELINE.*", re.I), "EBIT to NOPAT walk"),
    (re.compile(r"Reported EBIT\s*(?:[→\-]|to)\s*Normalized NOPAT\s*\(5-phase pipeline\)", re.I),
     "NOPAT bridge"),
    (re.compile(r"Phases 1[–-]4 clean and forecast.*", re.I), None),
    (re.compile(r"Phase \d:\s*", re.I), ""),
    (re.compile(r"Phase \d —.*", re.I), None),
    (re.compile(r"Guardrail:.*", re.I), None),
    (re.compile(r".*Firecrawl.*", re.I), None),
    (re.compile(r"Channel-mix EBIT output \(Phase 3\)", re.I), "Channel EBIT total"),
    (re.compile(r"Phase 1\s*[—\-]\s*Operating normalization", re.I), "Reported EBIT adjustments"),
    (re.compile(r"Phase 2\s*[—\-].*", re.I), "Lease & R&D adjustments"),
    (re.compile(r"Phase 3\s*[—\-].*", re.I), "Channel EBIT mix"),
    (re.compile(r"Phase 4\s*[—\-].*", re.I), "Tax normalization"),
    (re.compile(r"Phase 5\s*[—\-].*", re.I), "FY26 tariff refund"),
    (re.compile(r"Phase 1 adjusted EBIT = reported \+ non-recurring add-backs", re.I),
     "Reported EBIT + add-backs"),
    (re.compile(r"Phase 2 EBIT = Phase 1 \+ lease interest reclass \+ R&D cap − amortization", re.I),
     "Phase 1 + lease/R&D adjustments"),
    (re.compile(r"Phase 3 channel EBIT = Σ sector EBIT_i", re.I), "Sum of channel EBIT"),
    (re.compile(r"NOPAT Bridge tab: Phase \d subtotal", re.I), "Subtotal"),
    (re.compile(r"Add back SBC\? \(0 = no — Convention A:.*", re.I),
     "SBC add-back flag (0 = expensed)"),
    (re.compile(r"Convention A:.*", re.I), None),
    (re.compile(r"SBC Flag \(Convention A = 0\)", re.I), "SBC add-back (0)"),
    (re.compile(r"Sanity check: multiples within ±1\.5 turns\?", re.I), "Exit multiple check"),
    (re.compile(r"Forecast drivers linked to Scenarios tab.*", re.I),
     "Forecast linked to Scenarios (base case)"),
    (re.compile(r"BOTTOM-UP REVENUE DRIVER SCHEDULE.*", re.I), "Revenue drivers (FY26–30)"),
    (re.compile(r"Damodaran implied ERP \+ 177bps overlay", re.I), "Damodaran ERP + 177bps"),
    (re.compile(r"Alt\. source\s*\[cols.*", re.I), None),
    (re.compile(r"→", re.I), " to "),
]

COVER_CLEAR = {
    "B9", "B10", "B11", "B12", "B14", "B15", "B16", "B17", "B19", "B20", "B21", "B22", "B23", "B25",
}

URL_RE = re.compile(r"https?://", re.I)

COMPS_CLEAR_PREFIXES = (
    "•", "That tape is a check", "Alo Yoga", "ALO YOGA", "Memo:",
    "PitchBook pubcomps", "Note: pubcomps", "No Ctrl+F for EV",
    "No public HTML",
)

SOURCE_FALLBACKS: dict[tuple[str, str], tuple[str, str | None]] = {
    ("Scenarios", "Fixed buyback"): ("Pitch deck: $750M/yr buyback", None),
    ("Scenarios", "% of FCF to buybacks"): ("Pitch deck: 75% FCF to buybacks (bull)", None),
}


def _apply_label(text: str) -> str | None:
    out = text
    for pat, repl in LABEL_SUBS:
        if repl is None and pat.search(out):
            return None
        out = pat.sub(repl or "", out)
    out = re.sub(r"\s{2,}", " ", out).strip()
    return out or None


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "").lower().strip())


def main() -> None:
    wb = openpyxl.load_workbook(TARGET)
    stats = {"labels": 0, "cover": 0, "comps": 0, "sources": 0, "debris": 0}

    cov = wb["Cover"]
    for row in cov.iter_rows():
        for cell in row:
            v = cell.value
            if v is None:
                continue
            if cell.coordinate in COVER_CLEAR or (
                isinstance(v, str) and URL_RE.search(v)
            ):
                cell.value = None
                cell.hyperlink = None
                stats["cover"] += 1

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and not v.startswith("="):
                    if ws.title == "Comps" and any(v.strip().startswith(p) for p in COMPS_CLEAR_PREFIXES):
                        cell.value = None
                        stats["comps"] += 1
                        continue
                    new = _apply_label(v)
                    if new != v:
                        cell.value = new
                        stats["labels"] += 1

        for addr in ("K2", "L2", "M2"):
            if ws.title == "DCF" and ws[addr].value:
                if re.search(r"alt\.?\s*source|click\s*\+|ctrl\+f", str(ws[addr].value), re.I):
                    ws[addr].value = None
                    stats["debris"] += 1

    scn = wb["Scenarios"]
    for r in range(1, scn.max_row + 1):
        lab = _norm(scn.cell(r, 1).value)
        for key, (label, url) in SOURCE_FALLBACKS.items():
            sheet, needle = key
            if sheet != "Scenarios":
                continue
            if needle.lower() in lab:
                src_cell = scn.cell(r, 3)
                if not src_cell.value:
                    src_cell.value = label
                    src_cell.font = Font(color="0563C1", underline="single")
                    if url:
                        src_cell.hyperlink = Hyperlink(ref=src_cell.coordinate, target=url)
                    stats["sources"] += 1

    wb.properties.creator = "Microsoft Excel"
    wb.properties.lastModifiedBy = "Microsoft Excel"
    wb.save(TARGET)
    print(f"Final cleanup: {stats}")


if __name__ == "__main__":
    main()
