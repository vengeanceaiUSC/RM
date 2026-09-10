#!/usr/bin/env python3
"""Insert the GIS cover slide as slide 1 and renumber the deck.

Matches the official GIS title layout: NYC skyline background, GIS branding
top-left, centered ticker / recommendation / prices, division footer.

Run:  cd LULU && python3 scripts/add_gis_cover_slide.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

SCRIPTS = Path(__file__).resolve().parent
ROOT = SCRIPTS.parent
DECK = ROOT / "LULU_Investment_Pitch_Deck.pptx"
SKYLINE_SRC = ROOT / "assets" / "nyc_skyline.jpg"
SKYLINE_TINTED = ROOT / "assets" / "nyc_skyline_gis.jpg"

sys.path.insert(0, str(ROOT.parent / "GIS"))
from gis_pitch import FONT, GREY, NAVY, WHITE, _set_font  # noqa: E402

COMPANY = "lululemon athletica inc."
TICKER = "NASDAQ: LULU"
RECOMMENDATION = "LONG / OVERWEIGHT"
CURRENT_PRICE = "$100.61"
TARGET_PRICE = "$140"
UPSIDE = "+39% Upside"
DATE_LINE = "September 2026"
DIVISION = "Investment Research Division"
DISCLAIMER = (
    "Content slides can be added / removed depending on the "
    "investment opportunity (PM discretion)."
)


def cover_offset(prs: Presentation) -> int:
    """Return 1 when slide 1 is the GIS cover slide."""
    return 1 if _has_cover(prs) else 0


def _has_cover(prs: Presentation) -> bool:
    if not prs.slides:
        return False
    slide = prs.slides[0]
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text
        if "Table of Contents" in text:
            return False
        if "GLOBAL INVESTMENT SOCIETY" in text or (
            TICKER in text and RECOMMENDATION.split()[0] in text
        ):
            return True
    return False


def _tint_skyline() -> Path:
    SKYLINE_TINTED.parent.mkdir(parents=True, exist_ok=True)
    if not SKYLINE_SRC.exists():
        raise FileNotFoundError(f"Missing skyline image: {SKYLINE_SRC}")
    img = Image.open(SKYLINE_SRC).convert("RGB")
    img = img.resize((1920, 1080))
    overlay = Image.new("RGB", img.size, (31, 42, 68))
    blended = Image.blend(img, overlay, alpha=0.62)
    blended.save(SKYLINE_TINTED, quality=92)
    return SKYLINE_TINTED


def _insert_blank_at(prs: Presentation, index: int):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    sld_id_lst = prs.slides._sldIdLst
    ids = list(sld_id_lst)
    new_id = ids[-1]
    sld_id_lst.remove(new_id)
    sld_id_lst.insert(index, new_id)
    return slide


def _add_run(paragraph, text: str, size: float, color, *, bold: bool = False, italic: bool = False):
    run = paragraph.add_run()
    run.text = text
    _set_font(run, size, color, bold=bold, italic=italic)
    return run


def _build_cover_slide(slide, prs: Presentation) -> None:
    sw, sh = prs.slide_width, prs.slide_height
    bg_path = _tint_skyline()
    slide.shapes.add_picture(str(bg_path), 0, 0, width=sw, height=sh)

    # Top-left GIS branding (matches official GIS cover template)
    logo = slide.shapes.add_textbox(Inches(0.55), Inches(0.40), Inches(1.35), Inches(0.70))
    ltf = logo.text_frame
    ltf.clear()
    lp = ltf.paragraphs[0]
    _add_run(lp, "GIS", 36, WHITE, bold=True)

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.05), Inches(0.48), Inches(0.02), Inches(0.55))
    line.fill.solid()
    line.fill.fore_color.rgb = WHITE
    line.line.fill.background()

    name = slide.shapes.add_textbox(Inches(2.20), Inches(0.44), Inches(2.35), Inches(0.70))
    ntf = name.text_frame
    ntf.clear()
    ntf.word_wrap = True
    np1 = ntf.paragraphs[0]
    _add_run(np1, "Global", 11, WHITE)
    np2 = ntf.add_paragraph()
    _add_run(np2, "Investment", 11, WHITE)
    np3 = ntf.add_paragraph()
    _add_run(np3, "Society", 11, WHITE)

    # Center block
    cx = slide.shapes.add_textbox(Inches(1.0), Inches(2.55), Inches(11.33), Inches(2.35))
    ctf = cx.text_frame
    ctf.clear()
    ctf.word_wrap = True
    p1 = ctf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    _add_run(p1, f"{COMPANY} ({TICKER})", 30, WHITE, bold=True)

    p2 = ctf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(10)
    _add_run(p2, RECOMMENDATION, 22, WHITE)

    p3 = ctf.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(14)
    _add_run(p3, f"Current Price: {CURRENT_PRICE}", 17, WHITE)

    p4 = ctf.add_paragraph()
    p4.alignment = PP_ALIGN.CENTER
    p4.space_before = Pt(4)
    _add_run(p4, f"Price Target: {TARGET_PRICE} ({UPSIDE})", 17, WHITE)

    # Bottom-left division + disclaimer
    lb = slide.shapes.add_textbox(Inches(0.55), Inches(6.35), Inches(8.5), Inches(0.95))
    ltf = lb.text_frame
    ltf.clear()
    ltf.word_wrap = True
    lp1 = ltf.paragraphs[0]
    _add_run(lp1, DIVISION, 12, WHITE, bold=True)
    lp2 = ltf.add_paragraph()
    lp2.space_before = Pt(4)
    _add_run(lp2, DISCLAIMER, 9.5, WHITE, italic=True)

    # Bottom-right date
    rb = slide.shapes.add_textbox(Inches(10.2), Inches(6.72), Inches(2.6), Inches(0.35))
    rtf = rb.text_frame
    rtf.clear()
    rp = rtf.paragraphs[0]
    rp.alignment = PP_ALIGN.RIGHT
    _add_run(rp, DATE_LINE, 12, WHITE)


def _renumber_footers(prs: Presentation) -> None:
    """Page numbers start at 2 on the TOC slide; cover has no page number."""
    for idx, slide in enumerate(prs.slides):
        page = idx + 1
        footer_nums = []
        footer_srcs = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            top = shape.top
            if top < Inches(7.0):
                continue
            if text.isdigit():
                footer_nums.append(shape)
            elif text.startswith("Source:") or text == " ":
                footer_srcs.append(shape)
        if page == 1:
            for shape in footer_nums:
                shape.text_frame.clear()
            continue
        for shape in footer_nums:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT
            _add_run(p, str(page), 9, GREY)


def _update_toc_numbers(prs: Presentation) -> None:
    """Shift TOC # column by +1 after cover insertion (TOC is slide 2)."""
    if len(prs.slides) < 2:
        return
    slide = prs.slides[1]
    for shape in slide.shapes:
        if not shape.has_table:
            continue
        table = shape.table
        if table.rows[0].cells[0].text.strip() != "#":
            continue
        for ri in range(1, len(table.rows)):
            num_cell = table.rows[ri].cells[0]
            try:
                n = int(num_cell.text.strip())
            except ValueError:
                continue
            num_cell.text = str(n + 1)


def add_cover(path: Path = DECK) -> bool:
    prs = Presentation(str(path))
    if _has_cover(prs):
        slide = prs.slides[0]
        for shape in list(slide.shapes):
            shape.element.getparent().remove(shape.element)
        _build_cover_slide(slide, prs)
        prs.save(str(path))
        return False

    slide = _insert_blank_at(prs, 0)
    _build_cover_slide(slide, prs)
    _renumber_footers(prs)
    _update_toc_numbers(prs)
    prs.save(str(path))
    return True


if __name__ == "__main__":
    added = add_cover()
    print(f"Cover slide {'added' if added else 'already present'} → {DECK}")
