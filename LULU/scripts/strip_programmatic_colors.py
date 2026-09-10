#!/usr/bin/env python3
"""Reset all cell fonts to default black — no programmatic blue/green hex colors.

Human analysts apply GIS Cell Styles (blue inputs, black calcs, green links) manually
via the Excel ribbon. Run this after humanize_workbook_authentic.py if colors return.

Run:  cd LULU && python3 scripts/strip_programmatic_colors.py
"""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

SCRIPTS = Path(__file__).resolve().parent
import sys

sys.path.insert(0, str(SCRIPTS))
from styles import BLACK, FONT_NAME  # noqa: E402

ROOT = SCRIPTS.parent
TARGET = ROOT / "unbeiesgbar_final.xlsx"

MODEL_SHEETS = (
    "Cover",
    "WACC",
    "Scenarios",
    "Revenue Drivers",
    "NOPAT Bridge",
    "DCF",
    "Comps",
    "Scratch",
)


def _default_font(existing: Font | None) -> Font:
    bold = existing.bold if existing else False
    italic = existing.italic if existing else False
    size = existing.size if existing and existing.size else 10
    underline = existing.underline if existing else None
    return Font(
        name=FONT_NAME,
        color=BLACK,
        bold=bold,
        italic=italic,
        size=size,
        underline=underline,
    )


def strip_colors(path: Path = TARGET) -> dict[str, int]:
    tmp = path.with_suffix(".strip_colors.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)
    stats = {"cells_reset": 0, "sheets": 0}

    for name in MODEL_SHEETS:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        stats["sheets"] += 1
        for row in ws.iter_rows():
            for cell in row:
                if cell.font is None:
                    continue
                rgb = cell.font.color.rgb if cell.font.color else None
                if rgb and str(rgb).upper() not in ("000000", "00000000", "FF000000"):
                    cell.font = _default_font(cell.font)
                    stats["cells_reset"] += 1
                elif cell.font.name and cell.font.name != FONT_NAME:
                    cell.font = _default_font(cell.font)
                    stats["cells_reset"] += 1

    wb.save(tmp)
    tmp.replace(path)
    return stats


if __name__ == "__main__":
    result = strip_colors()
    print(f"Stripped programmatic colors: {result}")
