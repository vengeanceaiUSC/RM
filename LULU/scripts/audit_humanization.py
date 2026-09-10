#!/usr/bin/env python3
"""Comprehensive humanization audit: unbeiesgbar_final vs model18unaltered (12).

Read-only. Run: cd LULU && python3 scripts/audit_humanization.py
"""
from __future__ import annotations

import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
FINAL = ROOT / "unbeiesgbar_final.xlsx"
UNALTERED = ROOT / "model18unaltered (12).xlsx"

# Sheets to compare for assumptions / structure
KEY_SHEETS = ("WACC", "Scenarios", "NOPAT Bridge", "DCF", "Comps", "Revenue Drivers")

# Numeric literal inside formula (not cell ref, not function name)
FORMULA_NUM = re.compile(
    r"(?<![A-Z$])(?<![A-Z]\$)(?<!\$)(?<![\w.])(-?\d+\.?\d*|-?\.\d+)(?![\w.])(?!\$)",
)
# Skip common Excel constants / year-like in refs context
SKIP_NUMS = {"0", "1", "-1", "2", "3", "4", "5", "6", "7", "8", "9", "10", "12", "30", "100"}

ADD_CHAIN = re.compile(r"^=[\w\s'!$]+(?:\+[\w\s'!$]+){4,}", re.I)
DIV_FORM = re.compile(r"/", re.I)

AI_PHRASES = [
    (re.compile(r"justification|~20\s*word|click\s*\+|ctrl\+f|source\s*\(click\)|equation", re.I), "prompt/formula debris"),
    (re.compile(r"convention a|phase [1-5]|5-phase|agent workflow|pipeline", re.I), "pipeline jargon"),
    (re.compile(r"high-conviction|read-through|margin-of-safety|model cross-reference|on reported baseline", re.I), "AI boilerplate"),
    (re.compile(r"underwrite|hangs together|valuation anchor|sanity check|pitch deck", re.I), "consulting tone"),
    (re.compile(r"\banchor\.|\boverlay\b|operating assumption with", re.I), "template voice"),
    (re.compile(r"=\s*[A-Z]+\d|formula|cell ref|hardcod", re.I), "formula/meta language"),
]

SCENARIO_ASSUMP_ROWS = range(4, 21)  # typical assumption block
SCENARIO_COLS = (3, 4, 5)  # C-E bear/base/bull


@dataclass
class Hit:
    category: str
    sheet: str
    cell: str
    detail: str
    severity: str = "medium"
    example: str = ""


@dataclass
class AuditReport:
    hardcoded_in_formulas: list[Hit] = field(default_factory=list)
    raw_floats: list[Hit] = field(default_factory=list)
    notes_ai: list[Hit] = field(default_factory=list)
    add_chains: list[Hit] = field(default_factory=list)
    div_sources: list[Hit] = field(default_factory=list)
    sheet_diffs: dict = field(default_factory=dict)
    counts: dict = field(default_factory=dict)


def _is_formula(v) -> bool:
    return isinstance(v, str) and v.startswith("=")


def _long_float(v) -> bool:
    if not isinstance(v, (int, float)) or isinstance(v, bool):
        return False
    if isinstance(v, int):
        return False
    s = repr(v)
    if "e" in s.lower():
        return True
    dec = s.split(".")[-1].rstrip(")")
    return len(dec) > 4


def _extract_formula_nums(formula: str) -> list[str]:
    """Find numeric literals embedded in formula string."""
    # Strip string literals in quotes
    stripped = re.sub(r'"[^"]*"', '""', formula)
    nums = []
    for m in FORMULA_NUM.finditer(stripped):
        n = m.group(1)
        if n in SKIP_NUMS:
            continue
        # Skip if part of cell ref like C12 -> already excluded by lookbehind mostly
        nums.append(n)
    return nums


def _scan_hardcoded_formulas(wb) -> list[Hit]:
    hits: list[Hit] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not _is_formula(v):
                    continue
                nums = _extract_formula_nums(v)
                if not nums:
                    continue
                # Flag meaningful hardcodes (not tiny ints already skipped)
                meaningful = [n for n in nums if not (n.lstrip("-").isdigit() and abs(float(n)) <= 10)]
                if not meaningful and len(nums) <= 2:
                    # Still flag large ints like FY25 revenue literals
                    big = [n for n in nums if n.lstrip("-").isdigit() and len(n.replace(".", "").replace("-", "")) >= 5]
                    if not big:
                        continue
                    meaningful = big
                elif not meaningful:
                    meaningful = nums
                hits.append(
                    Hit(
                        "hardcoded_in_formula",
                        ws.title,
                        cell.coordinate,
                        f"nums={meaningful}",
                        "high" if any(len(n) >= 5 for n in meaningful) else "medium",
                        str(v)[:120],
                    )
                )
    return hits


def _scan_raw_floats(wb, sheet_filter: tuple[str, ...] = ("Scenarios", "WACC")) -> list[Hit]:
    hits: list[Hit] = []
    for name in sheet_filter:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        for row in ws.iter_rows():
            for cell in row:
                if _is_formula(cell.value):
                    continue
                if _long_float(cell.value):
                    label = str(ws.cell(cell.row, 1).value or "")[:40]
                    hits.append(
                        Hit(
                            "raw_float",
                            name,
                            cell.coordinate,
                            f"val={cell.value} fmt={cell.number_format} label={label!r}",
                            "high" if name == "Scenarios" else "medium",
                        )
                    )
    return hits


def _scan_notes_comments(wb) -> list[Hit]:
    hits: list[Hit] = []
    note_cols = {"WACC": 2, "Scenarios": 2, "NOPAT Bridge": 2, "Comps": 2, "Revenue Drivers": 9}
    for ws in wb.worksheets:
        ncol = note_cols.get(ws.title)
        for row in ws.iter_rows():
            for cell in row:
                texts: list[tuple[str, str]] = []
                if ncol and cell.column == ncol and isinstance(cell.value, str):
                    texts.append(("notes", cell.value))
                if cell.column == 1 and isinstance(cell.value, str):
                    texts.append(("label", cell.value))
                if cell.comment and cell.comment.text:
                    texts.append(("comment", cell.comment.text))
                for kind, text in texts:
                    if not text or _is_formula(text):
                        continue
                    for pat, cat in AI_PHRASES:
                        if pat.search(text):
                            hits.append(
                                Hit(
                                    "notes_ai",
                                    ws.title,
                                    cell.coordinate,
                                    f"{kind}/{cat}: {text[:80]}",
                                    "high" if cat in ("prompt/formula debris", "formula/meta language") else "medium",
                                )
                            )
                            break
    return hits


def _scan_add_chains(wb) -> list[Hit]:
    hits: list[Hit] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not _is_formula(v):
                    continue
                if "SUM(" in v.upper():
                    continue
                plus_count = v.count("+")
                if plus_count >= 5 and ADD_CHAIN.match(v.replace(" ", "")):
                    hits.append(
                        Hit(
                            "add_chain",
                            ws.title,
                            cell.coordinate,
                            f"{plus_count} terms",
                            "medium",
                            str(v)[:150],
                        )
                    )
                elif plus_count >= 8:
                    hits.append(
                        Hit(
                            "add_chain",
                            ws.title,
                            cell.coordinate,
                            f"{plus_count} terms (no SUM)",
                            "high",
                            str(v)[:150],
                        )
                    )
    return hits


def _scan_div_sources(wb) -> list[Hit]:
    """Heuristic DIV/0 sources: division formulas with risky denominators."""
    hits: list[Hit] = []
    risky_patterns = [
        (re.compile(r"/\s*0\b"), "divide by literal 0"),
        (re.compile(r"/\s*[A-Z]+\$\d+"), "divide by fixed ref (may be zero)"),
        (re.compile(r"/\([^)]*-[)]"), "divide by difference (can be zero)"),
        (re.compile(r"'Revenue Drivers'![FGHIJ]\d+"), "NOPAT stale RD col offset"),
        (re.compile(r"DCF!\$B\$[56]"), "stale DCF FY25 anchor ref"),
        (re.compile(r"Scenarios!\$H\$"), "stale Scenarios H col lock"),
        (re.compile(r"/\s*[A-Z]+\d+\s*$"), "trailing division (denominator may be 0)"),
    ]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not _is_formula(v):
                    continue
                for pat, reason in risky_patterns:
                    if pat.search(v):
                        hits.append(
                            Hit(
                                "div_source",
                                ws.title,
                                cell.coordinate,
                                reason,
                                "high" if "stale" in reason or "literal 0" in reason else "medium",
                                str(v)[:120],
                            )
                        )
                        break
    return hits


def _compare_sheets(wb_final, wb_unalt) -> dict:
    diffs = {}
    for name in KEY_SHEETS:
        if name not in wb_final.sheetnames or name not in wb_unalt.sheetnames:
            continue
        wf, wu = wb_final[name], wb_unalt[name]
        formula_diff = 0
        value_diff = 0
        samples = []
        max_r = max(wf.max_row, wu.max_row)
        max_c = max(wf.max_column, wu.max_column)
        for r in range(1, min(max_r + 1, 200)):
            for c in range(1, min(max_c + 1, 15)):
                cf, cu = wf.cell(r, c), wu.cell(r, c)
                vf, vu = cf.value, cu.value
                if vf == vu:
                    continue
                if _is_formula(vf) or _is_formula(vu):
                    formula_diff += 1
                    if len(samples) < 5:
                        samples.append({"cell": cf.coordinate, "final": str(vf)[:80], "unaltered": str(vu)[:80]})
                elif vf != vu:
                    value_diff += 1
                    if len(samples) < 8 and name == "Scenarios" and c in SCENARIO_COLS:
                        samples.append({"cell": cf.coordinate, "final": vf, "unaltered": vu})
        diffs[name] = {"formula_diffs": formula_diff, "value_diffs": value_diff, "samples": samples}
    return diffs


def run_audit() -> AuditReport:
    wb_f = openpyxl.load_workbook(FINAL, data_only=False)
    wb_u = openpyxl.load_workbook(UNALTERED, data_only=False) if UNALTERED.exists() else None

    report = AuditReport()
    report.hardcoded_in_formulas = _scan_hardcoded_formulas(wb_f)
    report.raw_floats = _scan_raw_floats(wb_f)
    report.notes_ai = _scan_notes_comments(wb_f)
    report.add_chains = _scan_add_chains(wb_f)
    report.div_sources = _scan_div_sources(wb_f)
    if wb_u:
        report.sheet_diffs = _compare_sheets(wb_f, wb_u)

    report.counts = {
        "hardcoded_in_formulas": len(report.hardcoded_in_formulas),
        "raw_floats": len(report.raw_floats),
        "raw_floats_scenarios": sum(1 for h in report.raw_floats if h.sheet == "Scenarios"),
        "notes_ai": len(report.notes_ai),
        "add_chains": len(report.add_chains),
        "div_sources": len(report.div_sources),
        "div_stale_refs": sum(1 for h in report.div_sources if "stale" in h.detail),
    }
    return report


def _print_section(title: str, hits: list[Hit], limit: int = 15):
    print(f"\n{'='*60}")
    print(f"{title}: {len(hits)} total")
    by_sheet = Counter(h.sheet for h in hits)
    for sh, n in by_sheet.most_common():
        print(f"  {sh}: {n}")
    print("\nExamples:")
    for h in hits[:limit]:
        print(f"  [{h.severity}] {h.sheet}!{h.cell} — {h.detail}")
        if h.example:
            print(f"       {h.example}")


def main():
    if not FINAL.exists():
        print(f"Missing {FINAL}", file=sys.stderr)
        return 1
    report = run_audit()
    print(f"Audit: {FINAL.name} vs {UNALTERED.name}")
    print(f"Counts: {json.dumps(report.counts, indent=2)}")

    _print_section("1. HARDCODED NUMBERS IN FORMULAS", report.hardcoded_in_formulas)
    _print_section("2. RAW DECIMAL FLOATS (Scenarios/WACC)", report.raw_floats)
    _print_section("3. NOTES/COMMENTS (AI/formula language)", report.notes_ai)
    _print_section("4. LONG ADD CHAINS (should be SUM)", report.add_chains)
    _print_section("5. DIV ERROR SOURCES", report.div_sources)

    print(f"\n{'='*60}")
    print("SHEET DIFFS vs unaltered:")
    for sh, d in report.sheet_diffs.items():
        print(f"  {sh}: {d['formula_diffs']} formula diffs, {d['value_diffs']} value diffs")
        for s in d.get("samples", [])[:3]:
            print(f"    {s}")

    out = ROOT / "scripts" / "audit_humanization_report.json"
    with open(out, "w") as f:
        json.dump(
            {
                "counts": report.counts,
                "hardcoded_in_formulas": [asdict(h) for h in report.hardcoded_in_formulas[:100]],
                "raw_floats": [asdict(h) for h in report.raw_floats[:100]],
                "notes_ai": [asdict(h) for h in report.notes_ai[:100]],
                "add_chains": [asdict(h) for h in report.add_chains[:50]],
                "div_sources": [asdict(h) for h in report.div_sources[:100]],
                "sheet_diffs": report.sheet_diffs,
            },
            f,
            indent=2,
            default=str,
        )
    print(f"\nFull report: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
