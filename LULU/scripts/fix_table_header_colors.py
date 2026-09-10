#!/usr/bin/env python3
"""Fix invisible navy table headers across the pitch deck.

GIS Garamond pass can strip explicit white font color from header runs,
leaving dark/invisible text on navy (1F2A44) header cells in PowerPoint.

Run:  cd LULU && python3 scripts/fix_table_header_colors.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "LULU_Investment_Pitch_Deck.pptx"
NAVY = "1F2A44"

sys.path.insert(0, str(ROOT.parent / "GIS"))
from gis_pitch import WHITE, _set_font  # noqa: E402


def _is_navy_header(cell) -> bool:
    try:
        return str(cell.fill.fore_color.rgb).upper() == NAVY
    except AttributeError:
        return False


def _run_is_white(run) -> bool:
    try:
        return str(run.font.color.rgb).upper() in ("FFFFFF", "FFFFFFFF")
    except AttributeError:
        return False


def fix(path: Path = DECK) -> int:
    prs = Presentation(str(path))
    fixed = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            table = shape.table
            for ci in range(len(table.columns)):
                cell = table.cell(0, ci)
                if not _is_navy_header(cell):
                    continue
                for para in cell.text_frame.paragraphs:
                    for run in para.runs:
                        if not run.text.strip() or _run_is_white(run):
                            continue
                        size = run.font.size.pt if run.font.size else 9
                        _set_font(run, size, WHITE, bold=bool(run.font.bold))
                        fixed += 1
    prs.save(str(path))
    return fixed


if __name__ == "__main__":
    n = fix()
    print(f"Fixed {n} header run(s) → {DECK}")
