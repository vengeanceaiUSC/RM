#!/usr/bin/env python3
"""Club-submission humanization for unbeiesgbar_final.xlsx.

- Rename template / AI section headers to analyst-style labels
- Fix bull-case share repurchase price path ($B$198:$B$202)
- Add visible Source column citations on hardcoded inputs
- Simplify NOPAT reconciliation step labels
- Set file metadata to Microsoft Excel (not openpyxl)

Run:  cd LULU && python3 scripts/humanize_for_club_submission.py
"""
from __future__ import annotations

import re
import shutil
import sys
from difflib import get_close_matches
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.worksheet.hyperlink import Hyperlink

SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPTS))

from restore_source_columns import (  # noqa: E402
    _apply_source_cell,
    _build_label_map,
    _clean_source_label,
    _match_row,
    _truncate_to_words,
)
from rewrite_model_notes import NOTES_MAP  # noqa: E402

ROOT = SCRIPTS.parent
ORIGINAL = ROOT / "model18unaltered (12).xlsx"
TARGET = ROOT / "unbeiesgbar_final.xlsx"

BLUE = "0563C1"
SOURCE_HEADER = "Source"

# Per-sheet column for visible Source (no value-column shift)
SOURCE_COL: dict[str, int] = {
    "WACC": 3,
    "Scenarios": 2,
    "NOPAT Bridge": 8,
    "DCF": 8,
    "Comps": 8,
    "Revenue Drivers": 10,
}

VALUE_COLS: dict[str, tuple[int, ...]] = {
    "WACC": (2,),
    "Scenarios": (3, 4, 5),
    "NOPAT Bridge": (2, 3, 4, 5, 6, 7),
    "DCF": (2, 3, 4, 5, 6, 7),
    "Comps": (2, 3, 4, 5, 6, 7),
    "Revenue Drivers": tuple(range(2, 9)),
}

# Exact label replacements (column A and selected headers)
LABEL_RENAMES: dict[str, str] = {
    "BOTTOM-UP REVENUE DRIVER SCHEDULE (FY2026E-FY2030E)": "Revenue drivers (FY26–30)",
    "V. CONSOLIDATED REVENUE & RECONCILIATION": "Consolidated revenue check",
    "Bottom-up total revenue ($000)": "Total revenue ($000)",
    "Variance (bottom-up - scenarios) ($000)": "Variance vs scenarios ($000)",
    "NOPAT RECONCILIATION (BASE CASE)": "NOPAT bridge",
    "Reconciliation (execution steps)": "Build checklist",
    "VALUATION - GORDON GROWTH (PERPETUITY) METHOD": "Terminal value — Gordon growth",
    "TERMINAL VALUE RECONCILIATION (GORDON vs EXIT MULTIPLE)": "Terminal value check",
    "Sanity check: multiples within ±1.5 turns?": "Exit multiple check (±1.5x)",
    "Pitch CFF: fixed buyback ($k/yr) bear/base": "Fixed buyback ($k/yr)",
    "Pitch CFF: % of FCF to buybacks (bull)": "% of FCF to buybacks (bull)",
    "Terminal value (Gordon growth)": "Terminal value",
    "Implied exit EV/EBITDA = Gordon TV / FY30 EBITDA": "Implied exit EV/EBITDA",
    "Selected exit EV/EBITDA (Gordon implied)": "Selected exit EV/EBITDA",
    "Terminal value = Terminal EBITDA x exit multiple": "Terminal value (exit multiple)",
    "Gordon Growth implied exit EV/EBITDA": "Gordon-implied exit EV/EBITDA",
    "Spread (selected - Gordon implied)": "Spread vs Gordon-implied",
    "Selected exit multiple (Gordon growth)": "Selected exit multiple",
    "DCF exit method (Gordon-implied on FY2030E EBITDA)": "DCF exit method",
    "% of EV from terminal value": "Terminal value % of EV",
}

NOPAT_STEP_RENAMES: dict[str, str] = {
    "1. Clean base": "Reported EBIT adjustments",
    "2. Capitalize R&D": "R&D capitalization",
    "3. Lease shift": "Lease interest reclass",
    "4. Segment mix": "Channel EBIT mix",
    "5. NOPAT bridge": "Tax-normalized NOPAT",
    "6. Tariff refund (Scenarios)": "FY26 tariff refund",
}

FALLBACK_SOURCES: dict[tuple[str, str], tuple[str, str | None]] = {
    ("Scenarios", "% of FCF to buybacks (bull)"): (
        "Pitch deck: 75% of FCF to buybacks (bull case)",
        None,
    ),
    ("Scenarios", "Fixed buyback ($k/yr)"): (
        "Pitch deck: $750M/yr buyback (bear/base)",
        None,
    ),
    ("NOPAT Bridge", "SBC add-back flag (0 = expensed)"): (
        "LULU CF statement (10-K): stock-based compensation",
        "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187",
    ),
    ("NOPAT Bridge", "+ Impairment / intangible amortization (add-back)"): (
        "LULU FY2025 10-K: amortization of intangible assets",
        "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187",
    ),
    ("NOPAT Bridge", "+ Restructuring costs (add-back)"): (
        "LULU FY2025 10-K: restructuring charges",
        "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187",
    ),
    ("NOPAT Bridge", "+ Legal / M&A one-offs (add-back)"): (
        "LULU FY2025 10-K: SG&A one-offs",
        "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187",
    ),
}


SCENARIO_ROW_SOURCES: dict[int, tuple[str, str | None]] = {
    4: ("LULU Q2 FY2026 earnings release", "https://corporate.lululemon.com/investors/news-events/press-releases"),
    5: ("StockAnalysis: LULU 3Y revenue forecast", "https://stockanalysis.com/stocks/lulu/forecast/"),
    6: ("Q2 FY2026 release: run-rate operating margin", "https://corporate.lululemon.com/investors/news-events/press-releases"),
    7: ("Q2 FY2026 release: $134.5M IEEPA tariff refunds", "https://corporate.lululemon.com/investors/news-events/press-releases"),
    8: ("FY2025 10-K: operating margin benchmark", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    9: ("WACC tab: CAPM build", None),
    10: ("FRED: Real GDP (GDPC1)", "https://fred.stlouisfed.org/series/GDPC1"),
    11: ("Q2 FY2026 outlook: tax rate ≈ 30%", "https://corporate.lululemon.com/investors/news-events/press-releases"),
    12: ("LULU CF statement (10-K)", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    13: ("LULU 10-K: capex guide", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    14: ("LULU FY2025 10-K: Gross profit & COGS", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    15: ("LULU FY2025 10-K: Accounts receivable, net", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    16: ("LULU FY2025 10-K: Inventories & COGS", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    17: ("LULU FY2025 10-K: inventory outlook", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    18: ("LULU FY2025 10-K: Accounts payable & COGS", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    19: ("LULU FY2025 10-K: other current assets", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    20: ("LULU FY2025 10-K: Accrued liabilities", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    21: ("Pitch deck: $750M/yr buyback (bear/base)", None),
    22: ("Pitch deck: 75% of FCF to buybacks (bull case)", None),
    138: ("LULU FY2025 10-K: other income", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    139: ("LULU FY2025 10-K: other income", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    140: ("LULU FY2025 10-K: other income", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    141: ("LULU FY2025 10-K: other income", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    142: ("LULU FY2025 10-K: other income", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    194: ("LULU FY2025 10-K: total assets", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    195: ("LULU FY2025 10-K: total liabilities", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    196: ("LULU FY2025 10-K: total equity", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    197: ("LULU FY2025 10-K: diluted share count", "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187"),
    198: ("Pitch deck: repurchase price path Y1", None),
    199: ("Pitch deck: repurchase price path Y2", None),
    200: ("Pitch deck: repurchase price path Y3", None),
    201: ("Pitch deck: repurchase price path Y4", None),
    202: ("Pitch deck: repurchase price path Y5", None),
}

DCF_ROW_SOURCES: dict[int, tuple[str, str | None]] = {
    10: ("No tariff refund post-FY26 (forecast years 2–5)", None),
    37: ("Mid-year discount convention (yrs 1–5)", None),
    95: ("Model check: Gordon vs exit-multiple spread", None),
}

NOPAT_ROW_SOURCES: dict[int, tuple[str, str | None]] = {
    7: (
        "LULU CF statement (10-K): stock-based compensation",
        "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187",
    ),
    17: (
        "LULU FY2025 10-K: amortization of intangible assets",
        "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187",
    ),
    18: (
        "LULU FY2025 10-K: restructuring charges",
        "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187",
    ),
    19: (
        "LULU FY2025 10-K: SG&A one-offs",
        "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0001397187",
    ),
}

COMPS_ROW_SOURCES: dict[int, tuple[str, str | None]] = {
    56: ("PitchBook EV/EBITDA comps (peer set)", "https://pitchbook.com/"),
}

ROW_SOURCE_MAP: dict[str, dict[int, tuple[str, str | None]]] = {
    "Scenarios": SCENARIO_ROW_SOURCES,
    "DCF": DCF_ROW_SOURCES,
    "NOPAT Bridge": NOPAT_ROW_SOURCES,
    "Comps": COMPS_ROW_SOURCES,
}

STALE_SOURCE_RE = re.compile(r"convention a|source:\s*source\b", re.I)

PROMPT_DEBRIS_RE = re.compile(
    r"justification|source:\s*source\b|ctrl\+f|click\s*\+|\[cols",
    re.I,
)


def _clean_prompt_comments(wb) -> int:
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not cell.comment or not cell.comment.text:
                    continue
                if PROMPT_DEBRIS_RE.search(cell.comment.text):
                    lines = [
                        ln.strip()
                        for ln in cell.comment.text.splitlines()
                        if ln.strip() and not PROMPT_DEBRIS_RE.search(ln)
                    ]
                    if lines:
                        cell.comment = Comment("\n".join(lines)[:32000], "Analyst")
                    else:
                        cell.comment = None
                    n += 1
    return n


def _norm(text) -> str:
    if text is None:
        return ""
    s = str(text).lower().strip().replace("–", "-").replace("—", "-")
    s = re.sub(r"phase [1-5][:\s-]*", "", s)
    s = re.sub(r"pitch cff:\s*", "", s)
    s = re.sub(r"\s+", " ", s)
    return s


def _rename_labels(wb) -> int:
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_col=1, max_col=1):
            cell = row[0]
            if not isinstance(cell.value, str) or cell.value.startswith("="):
                continue
            original = cell.value.strip()
            new = LABEL_RENAMES.get(original, original)
            if new == original and "Phase" in original:
                new = re.sub(r"Phase [1-5]:\s*", "", original).strip()
            if new == original and original in NOPAT_STEP_RENAMES:
                new = NOPAT_STEP_RENAMES[original]
            if new != original:
                cell.value = new
                n += 1
    return n


def _strip_phase_prefixes(wb) -> int:
    n = 0
    if "NOPAT Bridge" not in wb.sheetnames:
        return 0
    ws = wb["NOPAT Bridge"]
    for r in range(1, ws.max_row + 1):
        val = ws.cell(r, 1).value
        if not isinstance(val, str):
            continue
        new = re.sub(r"Phase [1-5]:\s*", "", val).strip()
        new = re.sub(r"\s*\(Phase [1-5]\)\s*", " ", new).strip()
        new = re.sub(r"\s{2,}", " ", new)
        if new != val:
            ws.cell(r, 1).value = new
            n += 1
    return n


def _fix_bull_share_path(wb) -> int:
    scn = wb["Scenarios"]
    n = 0
    for i, row in enumerate(range(183, 188)):
        price = f"$B${198 + i}"
        fcf_row = 153 + i
        if i == 0:
            want = f"=$B$197-ROUND(E22*E{fcf_row},0)/{price}"
        else:
            want = f"=E{row - 1}-ROUND(E22*E{fcf_row},0)/{price}"
        if scn[f"E{row}"].value != want:
            scn[f"E{row}"] = want
            n += 1
    return n


def _is_hardcode(val) -> bool:
    return isinstance(val, (int, float)) and not isinstance(val, bool)


def _collect_original_source(orig_ws, row: int) -> tuple[str | None, str | None, str | None]:
    c_cell = orig_ws.cell(row, 3)
    label = _clean_source_label(c_cell.value)
    url = None
    location = None
    if c_cell.hyperlink:
        url = c_cell.hyperlink.target
        location = c_cell.hyperlink.location
    if not label and url:
        label = "Source"
    return label, url, location


def _note_for_row(sheet: str, label: str) -> str | None:
    smap = NOTES_MAP.get(sheet, {})
    if label in smap:
        return smap[label]
    for lk, note in smap.items():
        if _norm(lk) == _norm(label):
            return note
    return None


def _set_source_headers(wb) -> None:
    for sheet, col in SOURCE_COL.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        header_row = 2 if sheet != "Revenue Drivers" else 5
        ws.cell(header_row, col).value = SOURCE_HEADER
        ws.cell(header_row, col).font = Font(bold=True, size=10)


def _add_visible_sources(wb, orig_wb) -> dict[str, int]:
    stats = {"filled": 0, "skipped": 0, "fallback": 0}

    for sheet, src_col in SOURCE_COL.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        orig_ws = orig_wb[sheet] if sheet in orig_wb.sheetnames else None
        orig_map = _build_label_map(orig_ws) if orig_ws else {}
        vcols = VALUE_COLS.get(sheet, (2,))

        for r in range(1, ws.max_row + 1):
            label_raw = ws.cell(r, 1).value
            if not label_raw:
                continue
            label = str(label_raw).strip()
            if not label or label.lower() in {"source", "notes", "driver", "company"}:
                continue

            has_hardcode = any(_is_hardcode(ws.cell(r, c).value) for c in vcols)
            row_map = ROW_SOURCE_MAP.get(sheet, {})
            if not has_hardcode and r not in row_map:
                continue

            src_cell = ws.cell(r, src_col)
            stale = isinstance(src_cell.value, str) and STALE_SOURCE_RE.search(src_cell.value)
            if (src_cell.value or src_cell.hyperlink) and not stale:
                stats["skipped"] += 1
                continue

            src_label, url, location = (None, None, None)
            if orig_ws:
                orig_row = _match_row(label, orig_map)
                if orig_row:
                    src_label, url, location = _collect_original_source(orig_ws, orig_row)

            row_map = ROW_SOURCE_MAP.get(sheet, {})
            if not src_label and r in row_map:
                src_label, url = row_map[r]
                stats["fallback"] += 1

            if not src_label:
                for (sh, key), (fb_label, fb_url) in FALLBACK_SOURCES.items():
                    if sh == sheet and key.lower() in label.lower():
                        src_label, url = fb_label, fb_url
                        stats["fallback"] += 1
                        break

            if not src_label:
                note = _note_for_row(sheet, label)
                if note:
                    src_label = note.split(".")[0][:60]
                    stats["fallback"] += 1

            if not src_label:
                continue

            if _apply_source_cell(src_cell, src_label, url, location):
                stats["filled"] += 1
            else:
                src_cell.value = src_label
                src_cell.font = Font(color=BLUE, underline="single", italic=True, size=9)
                stats["filled"] += 1

    return stats


def _integrate_assumption_block(wb) -> int:
    """Rename buried assumption rows into a readable block."""
    if "Scenarios" not in wb.sheetnames:
        return 0
    scn = wb["Scenarios"]
    n = 0
    if scn["A193"].value != "Assumptions (FY25 anchors & repurchase prices)":
        scn["A193"] = "Assumptions (FY25 anchors & repurchase prices)"
        n += 1
    price_labels = [
        (198, "Repurchase avg price Y1 ($/sh)"),
        (199, "Repurchase avg price Y2 ($/sh)"),
        (200, "Repurchase avg price Y3 ($/sh)"),
        (201, "Repurchase avg price Y4 ($/sh)"),
        (202, "Repurchase avg price Y5 ($/sh)"),
    ]
    for row, label in price_labels:
        if scn[f"A{row}"].value != label:
            scn[f"A{row}"] = label
            n += 1
    return n


def humanize(path: Path = TARGET, original: Path = ORIGINAL) -> dict:
    tmp = path.with_suffix(".club.xlsx")
    shutil.copy2(path, tmp)
    wb = openpyxl.load_workbook(tmp)
    orig_wb = openpyxl.load_workbook(original, data_only=False) if original.exists() else None

    stats = {
        "labels": _rename_labels(wb),
        "phase_strip": _strip_phase_prefixes(wb),
        "bull_shares": _fix_bull_share_path(wb),
        "assumption_block": _integrate_assumption_block(wb),
        "prompt_comments": _clean_prompt_comments(wb),
    }

    _set_source_headers(wb)
    if orig_wb:
        stats["sources"] = _add_visible_sources(wb, orig_wb)

    wb.save(tmp)

    from add_revenue_driver_sources import add_sources as add_rd_sources

    stats["rd_sources"] = add_rd_sources(path=tmp)

    from fix_file_metadata import fix as fix_meta

    stats["metadata"] = fix_meta(path=tmp)

    tmp.replace(path)
    return stats


if __name__ == "__main__":
    print(f"Club submission humanization: {humanize()}")
