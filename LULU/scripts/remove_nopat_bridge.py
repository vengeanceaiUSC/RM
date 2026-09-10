#!/usr/bin/env python3
"""Remove NOPAT Bridge tab and rewire DCF/Scenarios to flat tax on Scenarios EBIT.

The bridge is a reconciliation memo only; forecast NOPAT = EBIT × (1 − cash tax rate).
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPTS))
from restore_outline_groups import restore_outline_groups  # noqa: E402

ROOT = SCRIPTS.parent
TARGETS = [
    ROOT / "model18_wsp_formulas.xlsx",
    ROOT / "model18_humanized.xlsx",
    ROOT / "model18_humanized (1).xlsx",
]

EBIT_ROWS = [45, 46, 47, 48, 49]
NOPAT_ROWS = [50, 51, 52, 53, 54]
SCEN_COLS = {"F": "F", "G": "G", "H": "H"}


def remove_nopat_bridge(path: Path) -> list[str]:
    wb = openpyxl.load_workbook(path)
    changes: list[str] = []

    if "NOPAT Bridge" not in wb.sheetnames:
        return ["already removed"]

    scn = wb["Scenarios"]
    scn_tax = "$G$11"
    dcf_tax = "Scenarios!$G$11"
    for ebit_row, nopat_row in zip(EBIT_ROWS, NOPAT_ROWS):
        for col in SCEN_COLS:
            ebit_col = col
            new = f"={ebit_col}{ebit_row}*(1-{scn_tax})"
            scn[f"{col}{nopat_row}"].value = new
            changes.append(f"Scenarios {col}{nopat_row} -> flat-tax NOPAT")

    dcf = wb["DCF"]
    dcf["A13"].value = "Less: unlevered tax (cash tax rate)"
    for col in "FGHIJ":
        dcf[f"{col}13"].value = f"=-{col}11*{dcf_tax}"
        dcf[f"{col}77"].value = f"={col}14"
        changes.append(f"DCF {col}13/{col}77 rewired")

    if dcf["C4"].value == "NOPAT Bridge tab":
        dcf["C4"].value = None
        dcf["C4"].hyperlink = None
        changes.append("DCF C4 NOPAT link cleared")

    if "Cover" in wb.sheetnames:
        cov = wb["Cover"]
        tabs = cov["B14"].value
        if tabs and "NOPAT Bridge" in str(tabs):
            cov["B14"].value = (
                str(tabs)
                .replace("NOPAT Bridge • ", "")
                .replace(" • NOPAT Bridge", "")
                .replace("NOPAT Bridge", "")
                .strip()
            )
            changes.append("Cover tabs list updated")

    del wb["NOPAT Bridge"]
    changes.append("deleted NOPAT Bridge sheet")

    restore_outline_groups(wb)
    wb.save(path)
    return changes


if __name__ == "__main__":
    for path in TARGETS:
        if not path.exists():
            print(f"Skip missing: {path.name}")
            continue
        changes = remove_nopat_bridge(path)
        print(f"{path.name}: {len(changes)} changes")
        for c in changes:
            print(f"  {c}")
