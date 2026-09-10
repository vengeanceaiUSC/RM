#!/usr/bin/env python3
"""Remove NOPAT terminology from Excel model labels and notes."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
TARGETS = [
    ROOT / "model18_wsp_formulas.xlsx",
    ROOT / "LULU_DCF_Valuation_Model.xlsx",
]

_REPLACEMENTS = (
    ("NOPAT: year 1", "EBIT after tax: year 1"),
    ("NOPAT: year 2", "EBIT after tax: year 2"),
    ("NOPAT: year 3", "EBIT after tax: year 3"),
    ("NOPAT: year 4", "EBIT after tax: year 4"),
    ("NOPAT: year 5", "EBIT after tax: year 5"),
    ("NOPAT plus depreciation minus capex plus working capital.", "After-tax EBIT plus depreciation minus capex plus working capital."),
    ("NOPAT", "EBIT after tax"),
)


def scrub(path: Path) -> int:
    wb = load_workbook(path)
    changes = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if not isinstance(val, str) or "nopat" not in val.lower():
                    continue
                new = val
                for old, repl in _REPLACEMENTS:
                    new = new.replace(old, repl)
                if new != val:
                    cell.value = new
                    changes += 1
    wb.save(path)
    return changes


def main() -> None:
    for path in TARGETS:
        if not path.exists():
            continue
        n = scrub(path)
        print(f"{path.name}: {n} label(s) updated")


if __name__ == "__main__":
    main()
