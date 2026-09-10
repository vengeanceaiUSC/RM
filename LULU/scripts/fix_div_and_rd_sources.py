#!/usr/bin/env python3
"""Fix Scenarios #DIV/0! and restore visible Revenue Drivers sources.

Run:  cd LULU && python3 scripts/fix_div_and_rd_sources.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPTS))
from restore_outline_groups import restore_outline_groups  # noqa: E402
from restore_source_hyperlinks import (  # noqa: E402
    ORIGINAL,
    _apply_source_cell,
    _build_label_map,
    _collect_original_source,
    _match_row,
    _resolve_internal_link,
    _style_source_cell,
    RD_FY25_LINK_ROWS,
)

ROOT = SCRIPTS.parent
TARGETS = [
    ROOT / "model18_wsp_formulas.xlsx",
    ROOT / "model18_humanized.xlsx",
    ROOT / "model18_humanized (1).xlsx",
]
FY25_REV = 11102600


# Naive DCF!$E$5 → $B$198 replace corrupted DCF!$E$50/51/55 into $B$1980/81/85.
_CORRUPT_RE = re.compile(r"\$B\$198([0-9])")
_CORRUPT_MAP = {"0": "DCF!$E$50", "1": "DCF!$E$51", "5": "DCF!$E$55"}

IMPLIED_PX = {
    "F": "=(F131+DCF!$E$50+DCF!$E$51)/DCF!$E$55",
    "G": "=(G131+DCF!$E$50+DCF!$E$51)/DCF!$E$55",
    "H": "=(H131+DCF!$E$50+DCF!$E$51)/DCF!$E$55",
}


def _repair_corrupt_dcf_refs(formula: str) -> str:
    return _CORRUPT_RE.sub(lambda m: _CORRUPT_MAP[m.group(1)], formula)


def fix_scenarios_div(wb) -> int:
    if "Scenarios" not in wb.sheetnames:
        return 0
    ws = wb["Scenarios"]
    ws["A198"] = "FY25 net revenue ($000)"
    ws["B198"] = FY25_REV
    n = 0

    for col, formula in IMPLIED_PX.items():
        cell = ws[f"{col}132"]
        if cell.value != formula:
            cell.value = formula
            n += 1

    for col in ("F", "G", "H"):
        cell = ws[f"{col}163"]
        want = f"=DCF!$E$50+{col}153+{col}158"
        if cell.value != want:
            cell.value = want
            n += 1

    for row in ws.iter_rows():
        for cell in row:
            val = cell.value
            if not isinstance(val, str) or not val.startswith("="):
                continue
            new = val
            new = new.replace("Scenarios!$B$", "$B$")
            # Only replace exact DCF!$E$5 (revenue), not $E$50/$E$51/$E$55.
            new = re.sub(r"DCF!\$E\$5(?!\d)", "$B$198", new)
            new = _repair_corrupt_dcf_refs(new)
            if new != val:
                cell.value = new
                n += 1
    return n


def restore_rd_sources(wb, orig_wb) -> int:
    if "Revenue Drivers" not in wb.sheetnames:
        return 0
    ws = wb["Revenue Drivers"]
    orig_ws = orig_wb["Revenue Drivers"]
    orig_map = _build_label_map(orig_ws)
    n = 0
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if not label:
            continue
        orig_row = _match_row(str(label).strip(), orig_map)
        if orig_row is None:
            continue
        orig_label, orig_url, orig_loc = _collect_original_source(orig_ws, orig_row, source_col=11)
        if not orig_label and not orig_url:
            continue
        cell = ws.cell(r, 10)  # J
        if _apply_source_cell(cell, orig_label, orig_url, orig_loc):
            n += 1
        elif orig_label and not cell.value:
            cell.value = orig_label
            n += 1
        if cell.hyperlink:
            _style_source_cell(cell)

    for row, url in RD_FY25_LINK_ROWS.items():
        cell = ws.cell(row, 10)
        if cell.value and not cell.hyperlink:
            cell.hyperlink = url
            _style_source_cell(cell)
            n += 1

    ws.cell(5, 10).value = "Source"
    ws.cell(5, 10).font = Font(bold=True, size=10)
    return n


def unhide_rd_sources(wb) -> None:
    if "Revenue Drivers" not in wb.sheetnames:
        return
    ws = wb["Revenue Drivers"]
    for col in ("J", "K"):
        dim = ws.column_dimensions[col]
        dim.hidden = False
        dim.outline_level = 0
        dim.collapsed = False


def fix_workbook(path: Path) -> dict[str, int]:
    wb = openpyxl.load_workbook(path)
    orig = openpyxl.load_workbook(ORIGINAL, data_only=False)
    stats = {
        "scenarios_formulas": fix_scenarios_div(wb),
        "rd_sources": restore_rd_sources(wb, orig),
    }
    unhide_rd_sources(wb)
    restore_outline_groups(wb, col_hidden=False)
    unhide_rd_sources(wb)
    wb.save(path)
    return stats


if __name__ == "__main__":
    for path in TARGETS:
        if not path.exists():
            print(f"Skip missing: {path.name}")
            continue
        stats = fix_workbook(path)
        print(f"{path.name}: {stats}")
