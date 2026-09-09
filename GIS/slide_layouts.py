"""GIS pitch slide layouts — charts, tables, and structured content blocks."""
from __future__ import annotations

from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData

from gis_pitch import (
    PitchDeck,
    add_para,
    textbox,
    rect,
    _set_font,
    NAVY,
    CARD,
    GOLD,
    INK,
    GREY,
    LGREY,
    WHITE,
    GREEN,
)


def _kpi_box(slide, l, t, w, h, title, lines, title_color=CARD):
    box = rect(slide, l, t, w, h, fill=LGREY)
    btf = box.text_frame
    btf.word_wrap = True
    add_para(btf, title, 12, title_color, bold=True, first=True, space_after=2)
    for i, (text, size, color, bold) in enumerate(lines):
        add_para(btf, text, size, color, bold=bold, first=False, space_after=0 if i == len(lines) - 1 else 2)
    return box


def thesis_summary(
    deck: PitchDeck,
    descriptor: str,
    target: str,
    target_sub: str,
    rating: str,
    why_now: str,
    pillars: list[str],
    payoff: str,
):
    s = deck.slide_base("Investment Thesis Summary", descriptor, page=deck.pg())
    _kpi_box(s, Inches(0.5), Inches(1.2), Inches(3.9), Inches(1.5), "PRICE TARGET", [
        (target, 30, NAVY, True),
        (target_sub, 12, GREEN, True),
    ])
    _kpi_box(s, Inches(4.6), Inches(1.2), Inches(3.9), Inches(1.5), "RATING", [
        (rating, 24, NAVY, True),
        ("12-month horizon", 12, GREY, False),
    ])
    _kpi_box(s, Inches(8.7), Inches(1.2), Inches(4.15), Inches(1.5), "WHY NOW", [
        (why_now, 12.5, INK, False),
    ])
    tb, tf = textbox(s, Inches(0.5), Inches(2.95), Inches(12.35), Inches(4.0))
    add_para(tf, "Three reasons to be long", 15, CARD, bold=True, first=True, space_after=6)
    for i, p in enumerate(pillars, 1):
        add_para(tf, f"{i}.  {p}", 13.5, INK, space_after=7)
    add_para(tf, payoff, 13, NAVY, bold=True, italic=True, space_after=0)
    return s


def bullet_sections(deck: PitchDeck, title: str, descriptor: str, sections: list[tuple[str, list[str]]], sources=None):
    s = deck.slide_base(title, descriptor, sources=sources, page=deck.pg())
    tb, tf = deck.body_box(s)
    for i, (heading, bullets) in enumerate(sections):
        add_para(tf, heading, 14.5, CARD, bold=True, first=(i == 0), space_after=5)
        for b in bullets:
            add_para(tf, b, 13, INK, bullet=True, space_after=5)
    return s


def company_overview(
    deck: PitchDeck,
    descriptor: str,
    business_bullets: list[str],
    rev_lines: list[str],
    arch_blocks: list[tuple[str, str]],
    rev_chart_labels: list[str],
    rev_chart_values: list[float],
):
    s = deck.slide_base("Company Overview", descriptor, page=deck.pg())
    tb, tf = textbox(s, Inches(0.5), Inches(1.15), Inches(6.8), Inches(3.2))
    add_para(tf, "Business", 14.5, CARD, bold=True, first=True, space_after=5)
    for b in business_bullets:
        add_para(tf, b, 12.8, INK, bullet=True, space_after=5)
    add_para(tf, "Revenue trajectory (US$ M)", 14.5, CARD, bold=True, space_after=5)
    for line in rev_lines:
        add_para(tf, line, 12.8, INK, bullet=True, space_after=4)

    # Revenue bar chart
    chart_data = CategoryChartData()
    chart_data.categories = rev_chart_labels
    chart_data.add_series("Net revenue", rev_chart_values)
    chart = s.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.55), Inches(4.35), Inches(6.5), Inches(2.35),
        chart_data,
    ).chart
    chart.has_legend = False
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.tick_labels.font.size = Pt(9)
    chart.category_axis.tick_labels.font.size = Pt(9)
    if chart.series:
        chart.series[0].format.fill.solid()
        chart.series[0].format.fill.fore_color.rgb = NAVY

    box = rect(s, Inches(7.55), Inches(1.15), Inches(5.3), Inches(5.55), fill=LGREY)
    btf = box.text_frame
    btf.word_wrap = True
    add_para(btf, "GROWTH ARCHITECTURE", 12.5, CARD, bold=True, first=True, space_after=6)
    for name, desc in arch_blocks:
        add_para(btf, name, 13.5, NAVY, bold=True, space_after=1)
        add_para(btf, desc, 12, INK, space_after=8)
    return s


def business_model_slide(
    deck: PitchDeck,
    descriptor: str,
    unit_bullets: list[str],
    moat_bullets: list[str],
    margin_lines: list[str],
    om_chart_labels: list[str],
    om_chart_values: list[float],
):
    s = deck.slide_base("Business Model & Unit Economics", descriptor, page=deck.pg())
    tb, tf = textbox(s, Inches(0.5), Inches(1.15), Inches(6.8), Inches(3.5))
    add_para(tf, "Unit economics", 14.5, CARD, bold=True, first=True, space_after=5)
    for b in unit_bullets:
        add_para(tf, b, 12.8, INK, bullet=True, space_after=5)
    add_para(tf, "Competitive moats", 14.5, CARD, bold=True, space_after=5)
    for b in moat_bullets:
        add_para(tf, b, 12.8, INK, bullet=True, space_after=5)

    chart_data = CategoryChartData()
    chart_data.categories = om_chart_labels
    chart_data.add_series("Operating margin %", om_chart_values)
    chart = s.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(0.55), Inches(4.75), Inches(6.5), Inches(2.0),
        chart_data,
    ).chart
    chart.has_legend = False
    chart.value_axis.maximum_scale = max(om_chart_values) + 4
    chart.value_axis.minimum_scale = max(0, min(om_chart_values) - 4)
    chart.value_axis.tick_labels.font.size = Pt(9)
    chart.category_axis.tick_labels.font.size = Pt(8)
    if chart.series:
        chart.series[0].format.line.color.rgb = CARD
        chart.series[0].format.line.width = Pt(2.5)

    box = rect(s, Inches(7.55), Inches(1.15), Inches(5.3), Inches(3.2), fill=NAVY)
    btf = box.text_frame
    btf.word_wrap = True
    add_para(btf, "MARGIN PROFILE (FY2025)", 12.5, GOLD, bold=True, first=True, space_after=8)
    for line in margin_lines:
        add_para(btf, line, 14.5, WHITE, space_after=6)
    return s


def industry_slide(deck: PitchDeck, descriptor: str, category_bullets: list[str], fit_bullets: list[str], tam_segments: list[tuple[str, float]]):
    s = deck.slide_base("Industry Overview", descriptor, sources="Source: company filings; industry estimates (analyst)", page=deck.pg())
    tb, tf = textbox(s, Inches(0.5), Inches(1.15), Inches(6.5), Inches(5.7))
    add_para(tf, "Category dynamics", 14.5, CARD, bold=True, first=True, space_after=5)
    for b in category_bullets:
        add_para(tf, b, 13, INK, bullet=True, space_after=6)
    add_para(tf, "Where we fit", 14.5, CARD, bold=True, space_after=5)
    for b in fit_bullets:
        add_para(tf, b, 13, INK, bullet=True, space_after=6)

    chart_data = CategoryChartData()
    chart_data.categories = [x[0] for x in tam_segments]
    chart_data.add_series("Category share %", [x[1] for x in tam_segments])
    chart = s.shapes.add_chart(
        XL_CHART_TYPE.PIE,
        Inches(7.4), Inches(1.35), Inches(5.4), Inches(4.8),
        chart_data,
    ).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(10)
    tb2, tf2 = textbox(s, Inches(7.4), Inches(6.15), Inches(5.4), Inches(0.5))
    add_para(tf2, "Illustrative category mix — replace with your segment data", 10, GREY, italic=True, align=PP_ALIGN.CENTER, first=True)
    return s


def thesis_slide(
    deck: PitchDeck,
    num: str,
    descriptor: str,
    headline: str,
    bullets: list[str],
    metric_title: str,
    metric_lines: list[str],
):
    s = deck.slide_base(f"Investment Thesis {num}", descriptor, page=deck.pg())
    tb, tf = textbox(s, Inches(0.5), Inches(1.15), Inches(7.4), Inches(5.9))
    add_para(tf, headline, 15, CARD, bold=True, first=True, space_after=7)
    for t in bullets:
        add_para(tf, t, 13.3, INK, bullet=True, space_after=7)
    box = rect(s, Inches(8.15), Inches(1.15), Inches(4.7), Inches(len(metric_lines) * 0.55 + 0.9), fill=LGREY)
    btf = box.text_frame
    btf.word_wrap = True
    add_para(btf, metric_title, 12.5, CARD, bold=True, first=True, space_after=8)
    for ml in metric_lines:
        add_para(btf, ml, 13.5, NAVY, bold=True, space_after=6)
    return s


def risks_mitigants(deck: PitchDeck, descriptor: str, pairs: list[tuple[str, str]]):
    s = deck.slide_base("Risks & Mitigants", descriptor, page=deck.pg())
    tb, tf = textbox(s, Inches(0.5), Inches(1.15), Inches(12.35), Inches(0.3))
    add_para(tf, "Risk  \u2192  Mitigant", 14, CARD, bold=True, first=True, space_after=0)
    top = 1.55
    for rk, mg in pairs:
        b = rect(s, Inches(0.5), Inches(top), Inches(6.0), Inches(0.95), fill=LGREY)
        bt = b.text_frame
        bt.word_wrap = True
        bt.vertical_anchor = MSO_ANCHOR.MIDDLE
        add_para(bt, rk, 12.5, CARD, bold=True, first=True, space_after=0)
        b2 = rect(s, Inches(6.7), Inches(top), Inches(6.15), Inches(0.95), fill=WHITE, line=GREY)
        bt2 = b2.text_frame
        bt2.word_wrap = True
        bt2.vertical_anchor = MSO_ANCHOR.MIDDLE
        add_para(bt2, mg, 12.5, INK, first=True, space_after=0)
        top += 1.06
    return s


def catalyst_timeline(deck: PitchDeck, descriptor: str, items: list[tuple[str, str]]):
    s = deck.slide_base("Catalyst Timeline", descriptor, page=deck.pg())
    top = 1.5
    for when, what in items:
        rect(s, Inches(0.6), Inches(top + 0.05), Inches(0.22), Inches(0.22), fill=CARD)
        b = rect(s, Inches(1.1), Inches(top), Inches(2.9), Inches(1.05), fill=NAVY)
        bt = b.text_frame
        bt.word_wrap = True
        bt.vertical_anchor = MSO_ANCHOR.MIDDLE
        add_para(bt, when, 13, GOLD, bold=True, first=True, space_after=0)
        b2 = rect(s, Inches(4.2), Inches(top), Inches(8.6), Inches(1.05), fill=LGREY)
        bt2 = b2.text_frame
        bt2.word_wrap = True
        bt2.vertical_anchor = MSO_ANCHOR.MIDDLE
        add_para(bt2, what, 13, INK, first=True, space_after=0)
        top += 1.25
    return s


def financials_slide(deck: PitchDeck, title: str, descriptor: str, headers: list[str], rows: list[list[str]], footnote: str, bold_rows=()):
    s = deck.slide_base(title, descriptor, sources="Source: company 10-K filings; projections per GIS operating model", page=deck.pg())
    deck.stmt_table(s, rows, headers, col0w=3.2, bold_rows=bold_rows)
    tb, tf = textbox(s, Inches(0.5), Inches(6.75), Inches(12.35), Inches(0.35))
    add_para(tf, footnote, 11, GREY, italic=True, first=True, space_after=0)
    return s


def wacc_slide(deck: PitchDeck, descriptor: str, cap_bullets: list[str], wacc_rows: list[tuple[str, str]], wacc_result: str):
    s = deck.slide_base("Capital Structure & WACC", descriptor, page=deck.pg())
    tb, tf = textbox(s, Inches(0.5), Inches(1.2), Inches(6.2), Inches(5.6))
    add_para(tf, "Capital structure", 14.5, CARD, bold=True, first=True, space_after=5)
    for t in cap_bullets:
        add_para(tf, t, 13, INK, bullet=True, space_after=6)
    box = rect(s, Inches(7.0), Inches(1.2), Inches(5.85), Inches(4.6), fill=LGREY)
    btf = box.text_frame
    btf.word_wrap = True
    add_para(btf, "WACC BUILD (CAPM)", 13, CARD, bold=True, first=True, space_after=8)
    for label, val in wacc_rows:
        p = btf.add_paragraph()
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = label
        _set_font(r, 13, INK, bold=(label.startswith("Cost of equity")))
        r2 = p.add_run()
        r2.text = f"      {val}"
        _set_font(r2, 13, NAVY, bold=True)
    p = btf.add_paragraph()
    p.space_before = Pt(4)
    r = p.add_run()
    r.text = wacc_result
    _set_font(r, 16, GREEN, bold=True)
    return s


def football_field(
    deck: PitchDeck,
    descriptor: str,
    methods: list[tuple[str, float, float]],
    current: float,
    target: float,
    summary: str,
    subtext: str,
    vmin: float = 20,
    vmax: float = 100,
):
    s = deck.slide_base("Valuation Summary", descriptor, sources="Source: GIS DCF and comps models; multiples are analyst ranges", page=deck.pg())
    chart_l, chart_r = 3.2, 12.6

    def xpos(v):
        return chart_l + (v - vmin) / (vmax - vmin) * (chart_r - chart_l)

    top = 1.7
    for name, lo, hi in methods:
        tb, tf = textbox(s, Inches(0.5), Inches(top - 0.02), Inches(2.6), Inches(0.5), anchor=MSO_ANCHOR.MIDDLE)
        add_para(tf, name, 12, NAVY, bold=True, first=True, space_after=0)
        bar = rect(s, Inches(xpos(lo)), Inches(top), Inches(xpos(hi) - xpos(lo)), Inches(0.45), fill=GOLD)
        bt = bar.text_frame
        bt.vertical_anchor = MSO_ANCHOR.MIDDLE
        bt.margin_left = Pt(4)
        p = bt.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = f"${lo:.0f}"
        _set_font(r, 10.5, NAVY, bold=True)
        tb2, tf2 = textbox(s, Inches(xpos(hi) + 0.02), Inches(top), Inches(0.9), Inches(0.45), anchor=MSO_ANCHOR.MIDDLE)
        add_para(tf2, f"${hi:.0f}", 10.5, NAVY, bold=True, first=True, space_after=0)
        top += 0.75

    cp_x = xpos(current)
    pt_x = xpos(target)
    rect(s, Inches(cp_x), Inches(1.55), Pt(2), Inches(3.4), fill=INK)
    rect(s, Inches(pt_x), Inches(1.55), Pt(2), Inches(3.4), fill=CARD)
    tb, tf = textbox(s, Inches(cp_x - 0.7), Inches(4.95), Inches(1.6), Inches(0.3))
    add_para(tf, f"Current ${current:.0f}", 10.5, INK, bold=True, align=PP_ALIGN.CENTER, first=True, space_after=0)
    tb, tf = textbox(s, Inches(pt_x - 0.7), Inches(5.2), Inches(1.6), Inches(0.3))
    add_para(tf, f"Target ${target:.0f}", 10.5, CARD, bold=True, align=PP_ALIGN.CENTER, first=True, space_after=0)
    tb, tf = textbox(s, Inches(0.5), Inches(5.7), Inches(12.35), Inches(1.2))
    add_para(tf, summary, 13, NAVY, bold=True, first=True, space_after=5)
    add_para(tf, subtext, 12.5, INK, italic=True, space_after=0)
    return s


def dcf_slide(
    deck: PitchDeck,
    descriptor: str,
    assumptions: list[list[str]],
    output_lines: list[tuple[str, str]],
    implied: str,
    sensitivity: list[list[str]],
    hot_cell: tuple[int, int] = (2, 3),
):
    s = deck.slide_base("DCF Valuation", descriptor, sources="Source: GIS DCF model (from scratch); cash & share count per 10-K", page=deck.pg())
    tb, tf = textbox(s, Inches(0.5), Inches(1.15), Inches(6.1), Inches(0.4))
    add_para(tf, "Base-case assumptions", 14, CARD, bold=True, first=True, space_after=0)
    t = s.shapes.add_table(len(assumptions), 2, Inches(0.5), Inches(1.6), Inches(6.0), Inches(3.0)).table
    t.columns[0].width = Inches(4.0)
    t.columns[1].width = Inches(2.0)
    for ri, row in enumerate(assumptions):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if ri == 0 else (WHITE if ri % 2 else LGREY)
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.RIGHT
            r = p.add_run()
            r.text = val
            _set_font(r, 11.5, WHITE if ri == 0 else (CARD if ci == 1 and ri > 0 else INK), bold=(ri == 0))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    box = rect(s, Inches(7.0), Inches(1.6), Inches(5.85), Inches(3.0), fill=LGREY)
    btf = box.text_frame
    btf.word_wrap = True
    add_para(btf, "VALUATION OUTPUT (US$ M)", 12.5, CARD, bold=True, first=True, space_after=6)
    for label, val in output_lines:
        p = btf.add_paragraph()
        p.space_after = Pt(5)
        r = p.add_run()
        r.text = label
        _set_font(r, 12.5, INK, bold=("Enterprise" in label or "Equity" in label))
        r2 = p.add_run()
        r2.text = f"      {val}"
        _set_font(r2, 12.5, NAVY, bold=True)
    p = btf.add_paragraph()
    p.space_before = Pt(6)
    r = p.add_run()
    r.text = implied
    _set_font(r, 16, GREEN, bold=True)

    tb, tf = textbox(s, Inches(0.5), Inches(4.85), Inches(12.35), Inches(2.0))
    add_para(tf, "Sensitivity \u2014 implied share price (WACC vs terminal growth)", 13, CARD, bold=True, first=True, space_after=4)
    st = s.shapes.add_table(len(sensitivity), len(sensitivity[0]), Inches(0.5), Inches(5.35), Inches(7.6), Inches(1.5)).table
    for ri, row in enumerate(sensitivity):
        for ci, val in enumerate(row):
            cell = st.cell(ri, ci)
            cell.fill.solid()
            hot = (ri == hot_cell[0] and ci == hot_cell[1])
            cell.fill.fore_color.rgb = NAVY if (ri == 0 or ci == 0) else (GOLD if hot else (WHITE if ri % 2 else LGREY))
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = val
            _set_font(r, 10.5, WHITE if (ri == 0 or ci == 0) else (NAVY if hot else INK), bold=(ri == 0 or ci == 0 or hot))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    return s


def comps_slide(deck: PitchDeck, descriptor: str, headers: list[str], rows: list[list[str]], read_through: list[str]):
    s = deck.slide_base("Comparable Companies", descriptor, sources="Source: GIS estimates; peer multiples for illustration", page=deck.pg())
    deck.stmt_table(s, rows, headers, col0w=3.4, top=1.20, height=3.90, bold_rows=(0,))
    tb, tf = textbox(s, Inches(0.5), Inches(5.20), Inches(12.35), Inches(1.75))
    add_para(tf, "Read-through", 13.5, CARD, bold=True, first=True, space_after=3)
    for rt in read_through:
        add_para(tf, rt, 12.5, INK, bullet=True, space_after=3)
    return s


def scenario_columns(deck: PitchDeck, descriptor: str, cols: list[tuple[str, str, str, object, list[str]]], footer: str):
    s = deck.slide_base("Appendix \u2014 Bull / Bear Scenarios", descriptor, sources="Source: GIS DCF model (scenario tab)", page=deck.pg())
    x = 0.5
    for name, px, up, color, bullets in cols:
        head = rect(s, Inches(x), Inches(1.35), Inches(4.05), Inches(0.95), fill=color)
        ht = head.text_frame
        ht.word_wrap = True
        ht.vertical_anchor = MSO_ANCHOR.MIDDLE
        add_para(ht, f"{name}   {px}", 20, WHITE, bold=True, first=True, space_after=0, align=PP_ALIGN.CENTER)
        add_para(ht, up, 12.5, WHITE, align=PP_ALIGN.CENTER, space_after=0)
        b = rect(s, Inches(x), Inches(2.4), Inches(4.05), Inches(3.6), fill=LGREY)
        bt = b.text_frame
        bt.word_wrap = True
        for i, blt in enumerate(bullets):
            add_para(bt, blt, 12.5, INK, bullet=True, first=(i == 0), space_after=8)
        x += 4.25
    tb, tf = textbox(s, Inches(0.5), Inches(6.2), Inches(12.35), Inches(0.8))
    add_para(tf, footer, 13.5, NAVY, bold=True, italic=True, first=True, space_after=0)
    return s
